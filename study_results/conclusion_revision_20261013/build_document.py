"""Build and verify a chapter-only conclusion using the completed study's existing evidence."""
from pathlib import Path
from datetime import datetime, timezone
import csv
import hashlib
import json
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
STUDY = ROOT / 'study_results/ollama_clean_study_20261009'
LITERATURE = ROOT / 'study_results/literature_review_20261013/chapter_2.md'
TARGET = OUT / 'Conclusion_and_Future_Research_Rewritten.docx'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def field(paragraph, instruction):
    element = OxmlElement('w:fldSimple')
    element.set(qn('w:instr'), instruction)
    paragraph._p.append(element)


def build():
    manuscript = (OUT / 'chapter_7.md').read_text()
    effects = json.loads((STUDY / 'analysis/overall_effects.json').read_text())
    diagnostics = list(csv.DictReader((STUDY / 'analysis/llm_diagnostics.csv').open()))
    means = list(csv.DictReader((STUDY / 'analysis/policy_means.csv').open()))
    days = list(csv.DictReader((STUDY / 'analysis/daily_results.csv').open()))
    evidence = {'RQ1': {}, 'RQ2': {}, 'RQ3': {}, 'RQ4': {
        'assessment': 'Descriptive and exploratory; no formal cross-dataset effect test performed',
        'restriction': 'No inference that perishability alone caused the difference',
    }}
    for effect in effects:
        dataset = effect['dataset']
        key = 'RQ1' if effect['comparison'] == 'conventional' else 'RQ2'
        evidence[key][dataset] = effect
        fill = f"{abs(effect['fill_effect_pp']['mean']):.2f}"
        cost = f"{abs(effect['cost_effect_pct']['mean']):.2f}"
        if key == 'RQ1' or dataset == 'RetailNet':
            assert fill in manuscript and cost in manuscript
    for row in diagnostics:
        dataset = row['dataset']
        active = [d for d in days if d['arm'] == 'llm_agents' and d['dataset'] == dataset and d['scenario'] == 'promotion' and int(d['day']) >= 2]
        correct = sum(int(d['priority_correct']) for d in active)
        active_percent = correct / len(active) * 100
        evidence['RQ3'][dataset] = {
            'supplier_exact_extraction_percent': float(row['exact_extraction_rate']) * 100,
            'priority_accuracy_percent': float(row['priority_accuracy']) * 100,
            'active_promotion_correct': correct, 'active_promotion_decisions': len(active),
            'active_promotion_accuracy_percent': active_percent,
            'committed_constraint_violations': int(row['actual_violations']),
        }
        assert f"{float(row['priority_accuracy'])*100:.2f}%" in manuscript
        if dataset == 'RetailNet':
            assert f'{active_percent:.2f}%' in manuscript
    for dataset in ['M5', 'RetailNet']:
        for arm in ['conventional', 'llm_agents']:
            selected = [r for r in means if r['dataset'] == dataset and r['arm'] == arm]
            fill = sum(float(r['fill_rate']) for r in selected) / len(selected) * 100
            assert f'{fill:.2f}%' in manuscript
    evidence['hypothesis_assessments'] = {
        'H1a': 'Supported in both simulated datasets',
        'H1b': 'Supported in both simulated datasets',
        'H2a': 'Not supported; M5 superiority inconclusive, RetailNet service lower',
        'H2b': 'Not supported; M5 superiority inconclusive, RetailNet cost higher',
        'H3': 'Supported descriptively for the controlled tasks',
        'H4': 'Exploratory; formal support not established',
    }
    inputs = [LITERATURE, STUDY / 'analysis/overall_effects.json', STUDY / 'analysis/policy_means.csv',
              STUDY / 'analysis/llm_diagnostics.csv', STUDY / 'analysis/daily_results.csv',
              STUDY / 'analysis/audit_receipt.json', STUDY / 'protocol.json',
              STUDY / 'logs/runtime_environment.json']
    evidence['inputs'] = [{'path': str(p.relative_to(ROOT)), 'sha256': sha(p)} for p in inputs]
    evidence['analysis_scope'] = 'Reuses completed-study results; no new experiment, effect test or model calls'
    (OUT / 'evidence_map.json').write_text(json.dumps(evidence, indent=2) + '\n')

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.45)
    for style in doc.styles:
        if style.type in (1, 2):
            style.font.name = 'Times New Roman'
    for fonts in doc.styles.element.iter(qn('w:rFonts')):
        for attribute in list(fonts.attrib):
            if 'theme' in attribute.lower():
                del fonts.attrib[attribute]
        fonts.set(qn('w:eastAsia'), 'Times New Roman')
        fonts.set(qn('w:cs'), 'Times New Roman')
    normal = doc.styles['Normal']
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control = True
    for name, size in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 12)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string('20343D')
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
    header = section.header.paragraphs[0]
    header.text = 'CONCLUSION | LLM AGENTS FOR RETAIL REPLENISHMENT'
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Page ')
    field(footer, 'PAGE')
    footer.add_run(' of ')
    field(footer, 'NUMPAGES')
    for run in footer.runs:
        run.font.size = Pt(10)

    lines = manuscript.splitlines()
    table_count = 0
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                cells = [c.strip() for c in lines[index].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in cells):
                    rows.append(cells)
                index += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Table Grid'
            table.autofit = False
            widths = [1.90, 1.45, 2.91]
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
            for j, row in enumerate(rows):
                cells = table.add_row().cells
                split = OxmlElement('w:cantSplit')
                table.rows[j]._tr.get_or_add_trPr().append(split)
                if j == 0:
                    repeat = OxmlElement('w:tblHeader')
                    table.rows[j]._tr.get_or_add_trPr().append(repeat)
                for k, value in enumerate(row):
                    cells[k].width = Inches(widths[k])
                    cells[k].text = value
                    if j == 0:
                        shade = OxmlElement('w:shd')
                        shade.set(qn('w:fill'), 'E8EEF2')
                        cells[k]._tc.get_or_add_tcPr().append(shade)
                    for paragraph in cells[k].paragraphs:
                        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        paragraph.paragraph_format.line_spacing = 1.1
                        paragraph.paragraph_format.space_after = Pt(4)
                        for run in paragraph.runs:
                            run.font.size = Pt(10.5)
                            run.bold = j == 0
            table_count += 1
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        heading = re.match(r'^(#{1,3}) (.*)', line)
        if heading:
            paragraph = doc.add_heading(heading[2], level=len(heading[1]))
            if heading[2].startswith('7.3 Assessment'):
                paragraph.paragraph_format.page_break_before = True
        else:
            doc.add_paragraph(line)
        index += 1

    doc.core_properties.title = 'Conclusion and Future Research — LLM Agents for Retail Replenishment'
    doc.core_properties.subject = 'Rewritten Chapter 7 answering RQ1–RQ4 and hypotheses H1–H4'
    doc.core_properties.author = ''
    doc.core_properties.modified = datetime.now(timezone.utc).replace(tzinfo=None)
    doc.save(TARGET)
    prose = '\n'.join(l for l in lines if not l.startswith(('#', '|')))
    manifest = {
        'built_at_utc': datetime.now(timezone.utc).isoformat(),
        'scope': 'Only rewritten Chapter 7',
        'research_question_framework': str(LITERATURE.relative_to(ROOT)),
        'prose_words_excluding_headings_and_table': len(re.findall(r"\b[\w’'-]+\b", prose)),
        'tables': table_count, 'model': 'Pretrained Qwen2.5 1.5B via Ollama',
        'llm_retrained': False, 'new_experiments': False,
        'docx_sha256': sha(TARGET), 'manuscript_sha256': sha(OUT / 'chapter_7.md'),
        'evidence_map_sha256': sha(OUT / 'evidence_map.json'), 'builder_sha256': sha(__file__),
        'format': {'paper': 'A4', 'body_font': 'Times New Roman', 'body_points': 12,
                   'line_spacing': 1.5, 'margins_inches': 1, 'alignment': 'justified'},
    }
    (OUT / 'build_manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(json.dumps(manifest, indent=2))


if __name__ == '__main__':
    build()
