"""Check document completeness, scope, required questions and formatting; render review pages."""
from pathlib import Path
import hashlib
import json
import re

from docx import Document
from PIL import Image, ImageDraw
import fitz

OUT = Path(__file__).resolve().parent
docx = OUT / 'Chapters_1_and_2_Literature_Review.docx'
pdf = OUT / 'Chapters_1_and_2_Literature_Review.pdf'
document = Document(docx)
expected = []
for name in ['chapter_1.md', 'chapter_2.md']:
    expected.extend(re.sub(r'^#{1,3} ', '', line) for line in (OUT / name).read_text().splitlines() if line.strip())
actual = [p.text for p in document.paragraphs if p.text.strip()]
assert actual == expected, 'Word text differs from the complete chapter drafts'
heads = [p.text for p in document.paragraphs if p.style.name == 'Heading 1']
assert len(heads) == 2 and heads[0].startswith('1.') and heads[1].startswith('2.')
assert not any(p.text == 'References' for p in document.paragraphs)
normal = document.styles['Normal']
assert normal.font.name == 'Times New Roman' and normal.font.size.pt == 12
assert normal.paragraph_format.line_spacing == 1.5
assert all(abs(getattr(document.sections[0], k).inches - 1) < .001 for k in ['top_margin', 'bottom_margin', 'left_margin', 'right_margin'])
for tag in ['RQ1.', 'RQ2.', 'RQ3.', 'RQ4.', 'H1a:', 'H1b:', 'H2a:', 'H2b:', 'H3.', 'H4.']:
    assert any(tag in p.text for p in document.paragraphs), tag

source_text = '\n'.join(expected)
checks = {
    'Chen': r'Chen et al\. \(2000\)', 'Clark': r'Clark and Scarf \(1960\)',
    'Croston': r'Croston \(1972\)', 'DeHoratius': r'DeHoratius and Raman \(2008\)',
    'Dingdong': r'Dingdong-Inc\., 2025', 'Gneiting': r'Gneiting and Raftery \(2007\)',
    'Hevner': r'Hevner et al\. \(2004\)', 'Kang': r'Kang and Gershwin \(2005\)',
    'Ke': r'Ke et al\. \(2017\)', 'Lee': r'Lee et al\. \(1997\)',
    'Li': r'Li et al\.,? \(?2023\)', 'Makridakis': r'Makridakis et al\.,? \(?2022\)',
    'Nahmias 1982': r'Nahmias,? \(?1982\)', 'Nahmias 1994': r'Nahmias \(1994\)',
    'Ollama': r'Ollama, 2026', 'Parasuraman 1997': r'Parasuraman and Riley \(1997\)',
    'Parasuraman 2000': r'Parasuraman et al\. \(2000\)', 'Qwen': r'Qwen Team, 2024',
    'Syntetos': r'Syntetos and Boylan \(2005\)', 'Turpin': r'Turpin et al\.,? \(?2023\)',
    'van Donselaar': r'Van Donselaar et al\. \(2010\)', 'Wang 1996': r'Wang (?:and|&) Strong[, ]+\(?1996\)',
    'Wang 2025': r'Wang et al\.,? \(?2025\)', 'Wu': r'Wu et al\. \(2023\)',
    'Yao': r'Yao et al\.,? \(?2023\)', 'Zipkin': r'Zipkin,? \(?2000\)',
}
for name, pattern in checks.items():
    assert re.search(pattern, source_text), name
references = [x for x in (OUT / 'references.md').read_text().split('\n\n') if x.strip()]
assert len(references) == len(checks) == 26

reader = fitz.open(pdf)
texts = [re.sub(r'\s+', ' ', page.get_text()).strip() for page in reader]
assert all(len(t) > 30 for t in texts)
def page_of(text):
    return next(i + 1 for i, page in enumerate(texts) if text in page)
for tag in ['Main research question', 'RQ1.', 'RQ2.', 'RQ3.', 'RQ4.', 'H1a:', 'H1b:', 'H2a:', 'H2b:', 'H3.', 'H4.']:
    assert tag in ' '.join(texts), tag
for i, page in enumerate(reader, 1):
    for x0, y0, x1, y1, *_ in page.get_text('blocks'):
        assert x0 >= 20 and y0 >= 15 and x1 <= page.rect.width - 20 and y1 <= page.rect.height - 15, (i, x0, y0, x1, y1)

sheet = Image.new('RGB', (260 * 4, 365 * ((len(reader) + 3) // 4)), 'white')
for i, page in enumerate(reader):
    pix = page.get_pixmap(matrix=fitz.Matrix(.4, .4))
    image = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
    tile = Image.new('RGB', (260, 365), '#E9ECEF')
    tile.paste(image, ((260 - image.width) // 2, 8))
    ImageDraw.Draw(tile).text((10, 347), f'Page {i + 1}', fill='black')
    sheet.paste(tile, ((i % 4) * 260, (i // 4) * 365))
sheet.save(OUT / 'review_contact_sheet.png')
selected = sorted(set([1, page_of('2.9 Research aim'), page_of('H2. The LLM'), len(reader)]))
for number in selected:
    reader[number - 1].get_pixmap(matrix=fitz.Matrix(1.2, 1.2)).save(OUT / f'review_page_{number:02d}.png')
receipt = {
    'status': 'PASS', 'scope': 'Exactly two chapters, no additional sections',
    'docx_exact_text_match': True, 'citation_sources_cross_checked': 26,
    'format_verified': True, 'pdf_pages': len(reader), 'pdf_nonempty_pages': True,
    'pdf_text_within_page_bounds': True, 'visual_review_pages': selected,
    'document_sha256': hashlib.sha256(docx.read_bytes()).hexdigest(),
    'pdf_sha256': hashlib.sha256(pdf.read_bytes()).hexdigest(),
    'chapter_2_start_page': page_of('2. Literature review:'),
    'question_section_page': page_of('2.9 Research aim'),
    'hypothesis_section_page': page_of('2.10 Research hypotheses'),
}
(OUT / 'validation.json').write_text(json.dumps(receipt, indent=2) + '\n')
print(json.dumps(receipt, indent=2))
