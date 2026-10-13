"""Build a standalone literature-review part without modifying or publishing the repository."""
from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent
SOURCE = Path('/workspace/MasterDissertation/study_results/ollama_dissertation_20261012/references.md')
TARGET = OUT / 'Chapters_1_and_2_Literature_Review.docx'
CHAPTERS = [OUT / 'chapter_1.md', OUT / 'chapter_2.md']


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def add_field(paragraph, instruction):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    paragraph._p.append(field)


def add_text(paragraph, text):
    """Keep URLs clickable, while preserving the exact manuscript text."""
    position = 0
    for match in re.finditer(r'https?://\S+', text):
        paragraph.add_run(text[position:match.start()])
        url = match.group().rstrip('.,;')
        relationship = paragraph.part.relate_to(
            url,
            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
            is_external=True,
        )
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), relationship)
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        fonts = OxmlElement('w:rFonts')
        fonts.set(qn('w:ascii'), 'Times New Roman')
        fonts.set(qn('w:hAnsi'), 'Times New Roman')
        props.append(fonts)
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '24548A')
        props.append(color)
        run.append(props)
        node = OxmlElement('w:t')
        node.text = url
        run.append(node)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        paragraph.add_run(match.group()[len(url):])
        position = match.end()
    paragraph.add_run(text[position:])


def build():
    manuscripts = [p.read_text().strip() for p in CHAPTERS]
    references = [x.strip() for x in SOURCE.read_text().split('\n\n') if x.strip()]
    references = [x for x in references if not x.startswith('Gilbert,')]
    references.append(
        'Wu, Q., Bansal, G., Zhang, J., Wu, Y., Zhang, S., Zhu, E., Li, B., Jiang, L., '
        'Zhang, X., & Wang, C. (2023). AutoGen: Enabling next-gen LLM applications via '
        'multi-agent conversation framework [Preprint]. arXiv. https://arxiv.org/abs/2308.08155'
    )
    references.sort(key=lambda x: x.lower())
    (OUT / 'references.md').write_text('\n\n'.join(references) + '\n')
    combined = '\n\n'.join(manuscripts) + '\n'
    (OUT / 'Chapters_1_and_2_Literature_Review.md').write_text(combined)

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.45)

    for style in document.styles:
        if style.type in (1, 2):
            style.font.name = 'Times New Roman'
    for fonts in document.styles.element.iter(qn('w:rFonts')):
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():
                del fonts.attrib[attr]
        fonts.set(qn('w:eastAsia'), 'Times New Roman')
        fonts.set(qn('w:cs'), 'Times New Roman')
    normal = document.styles['Normal']
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string('17212B')
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control = True
    for name, size in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 12)]:
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string('20343D')
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
    ref_style = document.styles.add_style('Reference', 1)
    ref_style.base_style = normal
    ref_style.font.name = 'Times New Roman'
    ref_style.font.size = Pt(12)
    ref_style.paragraph_format.line_spacing = 1.15
    ref_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ref_style.paragraph_format.left_indent = Inches(.25)
    ref_style.paragraph_format.first_line_indent = Inches(-.25)
    ref_style.paragraph_format.space_after = Pt(8)
    ref_style.paragraph_format.keep_together = True

    header = section.header.paragraphs[0]
    header.text = 'LITERATURE REVIEW | LLM AGENTS FOR RETAIL REPLENISHMENT'
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Page ')
    add_field(footer, 'PAGE')
    footer.add_run(' of ')
    add_field(footer, 'NUMPAGES')
    for run in footer.runs:
        run.font.size = Pt(10)

    for index, manuscript in enumerate(manuscripts):
        for line in manuscript.splitlines():
            if not line.strip():
                continue
            heading = re.match(r'^(#{1,3}) (.*)', line)
            if heading:
                paragraph = document.add_heading(heading[2], level=len(heading[1]))
                if len(heading[1]) == 1:
                    paragraph.paragraph_format.page_break_before = index > 0
            else:
                paragraph = document.add_paragraph()
                add_text(paragraph, line)
                # Highlight the question and separately assessable hypotheses without adding content.
                if re.match(r'^H[12][ab]:', line) or line.startswith('To what extent can an LLM-agent'):
                    for run in paragraph.runs:
                        run.bold = True
    document.core_properties.title = 'Chapters 1 and 2: Literature Review — LLM Agents for Retail Replenishment'
    document.core_properties.subject = 'Thematic literature review, research questions and hypotheses'
    document.core_properties.author = ''
    document.core_properties.keywords = 'retail; replenishment; literature review; LLM agents; M5; FreshRetailNet; hypotheses'
    document.core_properties.modified = datetime.now(timezone.utc).replace(tzinfo=None)
    document.save(TARGET)

    counts = {}
    for path, manuscript in zip(CHAPTERS, manuscripts):
        prose = '\n'.join(line for line in manuscript.splitlines() if not line.startswith('#'))
        counts[path.name] = len(re.findall(r"\b[\w’'-]+\b", prose))
    manifest = {
        'created_at_utc': datetime.now(timezone.utc).isoformat(),
        'scope': 'Only rewritten Chapters 1 and 2; research questions and hypotheses are within Chapter 2',
        'publication': 'Local files only; no Git push',
        'chapter_words_excluding_headings': counts,
        'cited_sources': len(references),
        'bibliography_verification_file_not_in_document': 'references.md',
        'research_questions': ['Main question', 'RQ1', 'RQ2', 'RQ3', 'RQ4'],
        'hypotheses': ['H1a', 'H1b', 'H2a', 'H2b', 'H3', 'H4'],
        'hypothesis_status': 'Analytical framework; no claim of preregistration',
        'format': {
            'paper': 'A4', 'font': 'Times New Roman', 'body_size_pt': 12,
            'body_line_spacing': 1.5, 'margins_inch': 1, 'alignment': 'justified',
        },
        'docx_sha256': digest(TARGET),
        'manuscript_sha256': digest(OUT / 'Chapters_1_and_2_Literature_Review.md'),
        'reference_source': str(SOURCE),
        'reference_source_sha256': digest(SOURCE),
        'additional_source': {
            'citation': 'Wu et al. (2023), AutoGen',
            'bibliography_verified_from': 'https://raw.githubusercontent.com/microsoft/autogen/v0.2.0/README.md',
            'local_snapshot': '/tmp/literature_review_autogen_readme.md',
            'snapshot_sha256': digest('/tmp/literature_review_autogen_readme.md'),
        },
        'builder_sha256': digest(__file__),
    }
    (OUT / 'build_manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(json.dumps({'output': str(TARGET), 'chapter_words': counts, 'references': len(references)}, indent=2))


if __name__ == '__main__':
    build()
