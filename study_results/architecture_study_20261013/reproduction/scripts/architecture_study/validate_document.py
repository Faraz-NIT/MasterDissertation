"""Check that the single illustrated chapter contains the full audited text and figures."""
from pathlib import Path
import hashlib
import json
import re

from docx import Document
from PIL import Image, ImageDraw
import fitz

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'study_results/architecture_study_20261013'
docx = OUT / 'Selective_Agent_Architecture_Chapter.docx'
pdf_path = OUT / 'Selective_Agent_Architecture_Chapter.pdf'
document = Document(docx)
manuscript = (OUT / 'chapter.md').read_text()
expected = [re.sub(r'^#{1,3} ', '', line) for line in manuscript.splitlines()
            if line.strip() and not line.startswith('[[')]
actual = [p.text for p in document.paragraphs if p.text.strip() and p.style.name != 'Caption']
assert actual == expected, 'Chapter text differs from the complete manuscript'
assert len(document.inline_shapes) == 6 and len(document.tables) == 2
assert len([p for p in document.paragraphs if p.style.name == 'Heading 1']) == 1
tables = json.loads((OUT / 'chapter_tables.json').read_text())
for table, key in zip(document.tables, ['headline','hypotheses']):
    assert [[c.text for c in row.cells] for row in table.rows] == [tables[key]['headers'], *tables[key]['rows']]
normal = document.styles['Normal']
assert normal.font.name == 'Times New Roman' and normal.font.size.pt == 12
assert normal.paragraph_format.line_spacing == 1.5
assert all(abs(getattr(document.sections[0], name).inches-1) < .001 for name in ['top_margin','bottom_margin','left_margin','right_margin'])
assert json.loads((OUT / 'analysis/audit_receipt.json').read_text())['status'] == 'VERIFIED'
assert json.loads((OUT / 'execution_receipt.json').read_text())['status'] == 'COMPLETE'
pdf = fitz.open(pdf_path)
texts = [re.sub(r'\s+', ' ', page.get_text()).strip() for page in pdf]
assert all(len(text) > 40 for text in texts)
assert all(term in ' '.join(texts) for term in ['M5','RetailNet','LangGraph','A1:','A2:','A3:','A4:'])
for number, page in enumerate(pdf,1):
    for x0,y0,x1,y1,*_ in page.get_text('blocks'):
        assert x0 >= 18 and y0 >= 14 and x1 <= page.rect.width-18 and y1 <= page.rect.height-14, number
sheet = Image.new('RGB',(270*4,385*((len(pdf)+3)//4)),'white')
for i,page in enumerate(pdf):
    pix = page.get_pixmap(matrix=fitz.Matrix(.42,.42))
    image = Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    tile = Image.new('RGB',(270,385),'#ECEFF1')
    tile.paste(image,((270-image.width)//2,8))
    ImageDraw.Draw(tile).text((10,364),f'Page {i+1}',fill='black')
    sheet.paste(tile,((i%4)*270,(i//4)*385))
sheet.save(OUT/'review_contact_sheet.png')
for term,name in [('Figure 1.','architecture'),('Table 1.','headline'),('Table 2.','hypotheses')]:
    index = next(i for i,text in enumerate(texts) if term in text)
    pdf[index].get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(OUT/f'review_{name}.png')
record = {'status':'PASS','scope':'Exactly one additional chapter','docx_exact_text_match':True,
          'figures':6,'tables':2,'pdf_pages':len(pdf),'format_verified':True,
          'full_study_and_audit_required':True,'pdf_nonempty_pages':True,'page_bounds_verified':True,
          'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),
          'pdf_sha256':hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
          'visual_review':'Pending human-visible model review of rendered pages'}
(OUT/'document_validation.json').write_text(json.dumps(record,indent=2)+'\n')
print(json.dumps(record,indent=2))
