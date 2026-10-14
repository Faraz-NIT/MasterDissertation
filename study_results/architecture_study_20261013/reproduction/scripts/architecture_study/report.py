"""Create one illustrated dissertation chapter from the complete, audited architecture results."""
from __future__ import annotations
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'study_results/architecture_study_20261013'
DOCX = OUT / 'Selective_Agent_Architecture_Chapter.docx'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def field(paragraph, value):
    node = OxmlElement('w:fldSimple')
    node.set(qn('w:instr'), value)
    paragraph._p.append(node)


def draft(summary):
    diagnostics = {(r['dataset'], r['arm']): r for r in summary['diagnostics']}
    effects = {(r['dataset'], r['comparison'], r['metric']): r for r in summary['paired_effects'] if r['scenario'] == 'overall'}
    overall = pd.read_csv(OUT / 'analysis/overall_seed_results.csv')
    means = overall.groupby(['dataset','arm']).mean(numeric_only=True)
    verdicts = {r['dataset']: r for r in summary['registered_assessments']}
    def e(dataset, metric, comparison='fixed_agents'):
        return effects[(dataset, comparison, metric)]
    def interpretation(dataset):
        result = verdicts[dataset]
        service = ('The paired comparison established a service improvement within the simulated weeks.'
                   if result['service_superiority_supported'] else
                   'The paired comparison did not establish a service improvement within the simulated weeks.')
        cost = ('The paired comparison also established a lower primary cost index.'
                if result['cost_superiority_supported'] else
                'The paired comparison did not establish a lower primary cost index.')
        return service + ' ' + cost
    rows = []
    for dataset in ['M5','RetailNet']:
        fixed, selective = diagnostics[(dataset,'fixed_agents')], diagnostics[(dataset,'selective_agents')]
        rows.append([dataset,
                     f"{100*means.loc[(dataset,'fixed_agents'),'fill_rate']:.2f}% / {100*means.loc[(dataset,'selective_agents'),'fill_rate']:.2f}%",
                     f"{e(dataset,'cost_per_demand')['mean']:+.2f}%",
                     f"{fixed['fresh_model_calls']/fixed['portfolio_days']:.2f} / {selective['fresh_model_calls']/selective['portfolio_days']:.2f}",
                     f"{fixed['mean_decision_seconds']:.2f} / {selective['mean_decision_seconds']:.2f} s"])
    headline = {'headers': ['Dataset','Fill rate: fixed / selective','Cost change vs fixed','Calls: fixed / selective','Time: fixed / selective'], 'rows': rows}
    hypotheses = {'headers': ['Architecture hypothesis','M5','RetailNet'], 'rows': [
        ['A1: At least 50% fewer model calls', *['Supported' if verdicts[ds]['A1_supported'] else 'Not supported' for ds in ['M5','RetailNet']]],
        ['A2: Lower mean decision time', *['Supported' if verdicts[ds]['A2_supported'] else 'Not supported' for ds in ['M5','RetailNet']]],
        ['A3: Preserved or improved priority accuracy; no committed violations', *['Supported' if verdicts[ds]['A3_supported'] else 'Not supported' for ds in ['M5','RetailNet']]],
        ['A4: Service and cost preserved within the stated tolerances', *['Supported' if verdicts[ds]['A4_supported'] else 'Not supported' for ds in ['M5','RetailNet']]],
    ]}
    outcome = ('The architecture met all four registered success conditions in both datasets.'
               if summary['all_registered_success_conditions_met'] else
               'The architecture did not meet every registered success condition; the hypothesis table shows the limits of the result.')
    m5f, m5s = diagnostics[('M5','fixed_agents')], diagnostics[('M5','selective_agents')]
    frf, frs = diagnostics[('RetailNet','fixed_agents')], diagnostics[('RetailNet','selective_agents')]
    matches = pd.read_csv(OUT / 'analysis/matched_actions.csv')
    identical = bool(matches.orders_identical_to_structured.all())
    attribution = ('Every selective-workflow order matched the structured-information planner in the main study. '
                   'The new architecture therefore connected the tested language inputs to the reference planning behaviour without an additional numerical-performance advantage over that reference.'
                   if identical else
                   f"The selective workflow matched the structured-information planner on {100*matches.orders_identical_to_structured.mean():.2f}% of recorded orders. Differences and source failures remain part of the reported results; the architecture did not always reproduce the reference planning behaviour.")
    fixed_total = m5f['fresh_model_calls']+frf['fresh_model_calls']
    selective_total = m5s['fresh_model_calls']+frs['fresh_model_calls']
    latency_m5 = 100*(1-m5s['mean_decision_seconds']/m5f['mean_decision_seconds'])
    latency_fr = 100*(1-frs['mean_decision_seconds']/frf['mean_decision_seconds'])
    text = f'''# A selective agent architecture for retail replenishment

## 1. Purpose and architectural contribution

The earlier study showed that a small language model could read supplier information reliably, while its business-priority decisions were less consistent. This finding raises a practical design question: should every replenishment decision pass through the same sequence of language-model roles? A retailer needs a correct and timely order, but some parts of the process already have an explicit rule or a structured input.

This chapter develops and evaluates selective delegation. The architecture uses the language model where supplier information needs interpretation, applies declared business policies through fixed rules, and checks the proposed order before release. Its contribution is the design of these responsibilities and the evidence about their consequences. LangGraph provides the software framework; using the framework itself is not presented as the original research contribution.

The additional research question is: can selective delegation improve the reliability and computational efficiency of an LLM-agent replenishment system while preserving service and simulated cost performance? The comparison changes the architecture while keeping the model, forecasts and numerical planning tool constant.

## 2. How the architecture works

The supplier-intake component first examines the format of the incoming message. A complete table in the declared format is parsed directly. Supplier prose is passed to the language-model reader. A malformed declared table requests review. The parser reads the message itself; it does not fill missing values from the authenticated reference used for verification.

The extracted fields are then checked against the current supplier record. This record contains delivery and ordering terms, rather than future customer demand. A policy router applies the declared business instruction: an active promotion prioritises availability, while routine trading and supplier disruption use the balanced policy. Unknown promotion status prevents automatic release. The router does not infer a priority from the experiment's future sales.

The shared numerical tool calculates the order. Deterministic checks verify its quantities against the budget, supplier capacity, carton size and minimum-order rules. Routine valid decisions proceed without a separate LLM critic call. Failed or unknown conditions invoke the exception-review path and hold the action. A model approval cannot override a failed check. Review requests are recorded; the experiment does not simulate a human resolving them.

[[FIGURE:selective_architecture]]

LangGraph implements this process as an explicit state graph with conditional edges. The main nodes cover intake, extraction, source verification, priority routing, constrained planning, action validation, review and release. Each decision record includes the nodes actually visited, the reason for delegation, the extracted information and the released quantity. The fixed comparison also runs through LangGraph, so the study does not confound the new architecture with the mere introduction of the framework.

## 3. Experiment design

The experiment includes thirty selected M5 series and thirty selected FreshRetailNet-50K series, labelled RetailNet. Each dataset is evaluated in two seven-day windows under routine trading, supplier disruption and promotion. Thirty new simulation seeds produce a total of 1,080 policy runs. Every matched case uses the same starting-stock rule, demand realisation and supplier opportunities across workflows.

Three workflows are compared. The fixed architecture calls the supplier reader, LLM coordinator and LLM critic on every decision. The selective architecture uses the conditional process described above. The structured-information reference uses the same numerical planner with already prepared current supplier fields and no LLM. This reference helps identify whether an improvement comes from planning, interpretation or the decision about which component performs a task.

All workflows use the same previously fitted, past-only LightGBM forecasts and the same constrained-order tool. The language model remains pretrained Qwen2.5 1.5B, running locally through Ollama on the four-CPU Linux cloud instance with 32 GiB of memory and no GPU. It was not trained or fine-tuned for this experiment. LangGraph 1.0.10 coordinates the workflow. Generation settings, role prompts and worked examples are unchanged from the fixed system.

Development checks used invented panels and separate development seeds. Eleven automated checks and live LangGraph integrations passed before the main protocol was frozen. Every main model response is fresh and recorded; no previous answer is loaded and no response cache or retry is used. The workflows' execution order is counterbalanced to reduce systematic timing advantages.

The registered assessment asks for at least 50% fewer model calls, lower mean decision time, preserved or improved priority accuracy without committed-action violations, and preservation of inventory outcomes. For the last criterion, the lower service-change bound must be greater than -0.10 percentage points and the upper cost-change bound must be below +0.5%. These are study tolerances, not retailer-approved targets. Paired uncertainty intervals describe thirty simulated seeds conditional on the chosen weeks; they do not estimate a retailer population effect.

## 4. Main results

{outcome} The table summarises the measured outcomes. Fixed and selective values are shown together to make the comparison easy to follow. A negative cost change favours the selective workflow. Costs remain synthetic inventory operating-cost indices, not realised financial savings.

[[TABLE:headline]]

### 4.1 Fewer model calls and faster decisions

The fixed workflow made {fixed_total:,} primary model calls, compared with {selective_total:,} for selective delegation, a reduction of {100*(1-selective_total/fixed_total):.2f}%. The reduction follows from skipping routine coordinator and critic calls and using the declared parser for supplier tables. It is an architectural saving, rather than evidence that the model generated answers more quickly or became more capable.

Mean decision time changed from {m5f['mean_decision_seconds']:.2f} to {m5s['mean_decision_seconds']:.2f} seconds in M5, and from {frf['mean_decision_seconds']:.2f} to {frs['mean_decision_seconds']:.2f} seconds in RetailNet. These represent measured reductions of {latency_m5:.2f}% and {latency_fr:.2f}%, respectively. The measures include graph orchestration and numerical planning on this instance. They do not measure employee time or production economics.

[[FIGURE:model_calls]]

[[FIGURE:decision_time]]

### 4.2 Business-priority reliability

Priority accuracy changed from {100*m5f['priority_accuracy']:.2f}% to {100*m5s['priority_accuracy']:.2f}% in M5, and from {100*frf['priority_accuracy']:.2f}% to {100*frs['priority_accuracy']:.2f}% in RetailNet. On active promotion days, the corresponding values were {100*m5f['active_promotion_accuracy']:.2f}% and {100*m5s['active_promotion_accuracy']:.2f}% for M5, and {100*frf['active_promotion_accuracy']:.2f}% and {100*frs['active_promotion_accuracy']:.2f}% for RetailNet.

This measure evaluates agreement with the declared business policy. It does not prove that the policy is commercially optimal. The improvement, where present, comes from implementing that instruction explicitly instead of asking the model to choose it repeatedly. Supplier extraction accuracy is reported separately in the diagnostics, so a successful priority decision cannot conceal an incorrect supplier field.

Exact supplier-information accuracy was {100*m5f['exact_source_accuracy']:.2f}% for the fixed workflow and {100*m5s['exact_source_accuracy']:.2f}% for selective delegation in M5. The corresponding RetailNet values were {100*frf['exact_source_accuracy']:.2f}% and {100*frs['exact_source_accuracy']:.2f}%. These figures include the current budget and all supplier terms and concern the controlled message corpus.

[[FIGURE:priority_accuracy]]

### 4.3 Service, cost and fresh-stock outcomes

Relative to the fixed architecture, the selective workflow changed M5 fill rate by {e('M5','fill_rate')['mean']:+.2f} percentage points and its operating-cost index by {e('M5','cost_per_demand')['mean']:+.2f}%. {interpretation('M5')}

For RetailNet, fill rate changed by {e('RetailNet','fill_rate')['mean']:+.2f} percentage points and the operating-cost index by {e('RetailNet','cost_per_demand')['mean']:+.2f}%. {interpretation('RetailNet')} These effects concern a change between two workflows already using the same strong planner, and should be kept distinct from the earlier comparison with conventional ordering.

[[FIGURE:inventory_effects]]

RetailNet also requires attention to waste. Its hypothetical three-day FIFO process means that a stronger availability priority can change the amount of stock that expires. The expiry graph reports all three situations, rather than assuming that a better service or cost outcome must also mean less waste. These expiry results depend on the simulated stock-life assumption; they are not observed disposal quantities from the dataset publisher.

The selective workflow's overall RetailNet expiry share was {e('RetailNet','expiry_share')['mean']:+.2f} percentage points different from the fixed workflow. In this study, the increase was concentrated in promotion handling. The lower primary cost index therefore coexisted with a small increase in waste, making the service–waste trade-off explicit.

[[FIGURE:retailnet_expiry]]

The selective workflow recorded {m5s['committed_violations']+frs['committed_violations']} committed-action constraint violations and {m5s['review_requests']+frs['review_requests']} review requests in the main grid. Failed-condition behaviour was checked separately during development. No unresolved operational exceptions arose in the primary study, so those tests support the checks' implementation rather than establishing performance during live exception handling.

## 5. Assessment and attribution

The architecture hypotheses are separate from the dissertation's earlier hypotheses about the complete replenishment package. Their results are stated below using the criteria frozen before this main evaluation.

[[TABLE:hypotheses]]

{attribution} The structured reference therefore remains essential to interpretation. Any gain over the fixed workflow can be explained by how known policies are applied and how unnecessary calls are avoided; it should not be attributed to a more powerful forecast, a different optimisation method or training of the language model.

The engineering contribution is a tested allocation of decision responsibilities: language interpretation for the tested prose, direct parsing for the declared tables, explicit priority rules and verifiable release conditions. Local graph traces make that allocation inspectable. The study evaluates this design against a matched fixed workflow rather than treating more agent conversations as a measure of success.

## 6. Business relevance and limitations

For managers, the design offers a way to use language assistance where communication creates work while retaining established numerical controls. Fewer model calls and shorter processing time can improve technical feasibility on modest infrastructure. Whether this translates into a useful commercial return requires a separate assessment of integration, maintenance and employee checking effort.

The architecture was motivated by previously inspected diagnostics. New simulation seeds and fresh model calls do not make the historical weeks unseen research holdouts. The evidence covers thirty selected series per dataset, controlled supplier messages, fixed declared policies and short simulated horizons. It does not establish performance across all M5 products or all 50,000 RetailNet series. RetailNet quantities are normalised sales proxies, and the inventory costs and stock-life process are hypothetical.

The source-verification reference provides current authenticated supplier fields to all workflows. Real deployment would need equivalent access to reliable operational records. The primary study does not measure human review or the handling of arbitrary conflicting messages. Correct routing under a known event flag is useful engineering evidence, but it is not evidence that the model independently understands every promotion or supplier exception.

## 7. Conclusion and next step

{outcome} The specific verdicts and uncertainty intervals determine where stronger claims about service or cost are justified. The main architectural contribution is selective delegation: deciding when an LLM is needed, when a declared rule is sufficient and when an action must be held.

The next step is a pilot with real supplier communications and measured staff review. It should retain the same planner, include unfamiliar message formats and incomplete instructions, and count the time needed to produce a correct input. Longer periods and realistic product lives would then test whether the architecture's efficiency and reliability persist under broader operating conditions.
'''
    return text, {'headline': headline, 'hypotheses': hypotheses}


def build():
    assert json.loads((OUT / 'execution_receipt.json').read_text())['status'] == 'COMPLETE'
    audit = json.loads((OUT / 'analysis/audit_receipt.json').read_text())
    assert audit['status'] == 'VERIFIED'
    summary = json.loads((OUT / 'analysis/analysis_summary.json').read_text())
    manuscript, tables = draft(summary)
    (OUT / 'chapter.md').write_text(manuscript.strip()+'\n')
    (OUT / 'chapter_tables.json').write_text(json.dumps(tables, indent=2)+'\n')
    register = json.loads((OUT / 'figure_register.json').read_text())['figures']
    figures = {r['key']: r for r in register}
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.45)
    for style in doc.styles:
        if style.type in [1,2]:
            style.font.name = 'Times New Roman'
    for fonts in doc.styles.element.iter(qn('w:rFonts')):
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():
                del fonts.attrib[attr]
    normal = doc.styles['Normal']
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control = True
    for name, size in [('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string('20343D')
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
    caption = doc.styles['Caption']
    caption.font.name = 'Times New Roman'
    caption.font.size = Pt(10.5)
    caption.paragraph_format.line_spacing = 1.1
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header = section.header.paragraphs[0]
    header.text = 'SELECTIVE AGENT ARCHITECTURE | M5 AND RETAILNET'
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Page ')
    field(footer,'PAGE')
    footer.add_run(' of ')
    field(footer,'NUMPAGES')
    for run in footer.runs:
        run.font.size = Pt(10)
    figure_number = table_number = 0
    for line in manuscript.splitlines():
        if not line.strip():
            continue
        heading = re.match(r'^(#{1,3}) (.*)', line)
        if heading:
            doc.add_heading(heading[2], level=len(heading[1]))
        elif line.startswith('[[FIGURE:'):
            key = line[9:-2]
            figure = figures[key]
            figure_number += 1
            p = doc.add_paragraph(style='Caption')
            p.paragraph_format.keep_with_next = True
            p.add_run(f"Figure {figure_number}. {figure['scope']}: {figure['title']}").bold = True
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_with_next = True
            width = 4.4 if key == 'selective_architecture' else 6.22
            p.add_run().add_picture(str(OUT / figure['path']), width=Inches(width))
            p = doc.add_paragraph(figure['caption'], style='Caption')
            p.paragraph_format.keep_together = True
        elif line.startswith('[[TABLE:'):
            key = line[8:-2]
            data = tables[key]
            table_number += 1
            p = doc.add_paragraph(style='Caption')
            p.paragraph_format.keep_with_next = True
            p.add_run(f"Table {table_number}. {'Overall comparison' if key=='headline' else 'Registered architecture hypotheses'}").bold = True
            table = doc.add_table(rows=0, cols=len(data['headers']))
            table.style = 'Table Grid'
            for i, row in enumerate([data['headers'], *data['rows']]):
                cells = table.add_row().cells
                split = OxmlElement('w:cantSplit')
                table.rows[i]._tr.get_or_add_trPr().append(split)
                if i == 0:
                    repeat = OxmlElement('w:tblHeader')
                    table.rows[i]._tr.get_or_add_trPr().append(repeat)
                for cell, value in zip(cells,row):
                    cell.text = value
                    for p in cell.paragraphs:
                        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p.paragraph_format.line_spacing = 1.1
                        p.paragraph_format.space_after = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(10)
                            run.bold = i == 0
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        else:
            doc.add_paragraph(line)
    doc.core_properties.title = 'A Selective Agent Architecture for Retail Replenishment'
    doc.core_properties.subject = 'One chapter covering LangGraph design and a fresh matched M5/RetailNet experiment'
    doc.core_properties.author = ''
    doc.core_properties.modified = datetime.now(timezone.utc).replace(tzinfo=None)
    doc.save(DOCX)
    prose = '\n'.join(l for l in manuscript.splitlines() if not l.startswith(('#','[[')))
    manifest = {'status':'BUILT_PENDING_RENDER','at_utc':datetime.now(timezone.utc).isoformat(),
                'scope':'One additional chapter only','prose_words':len(re.findall(r"\b[\w’'-]+\b",prose)),
                'figures':figure_number,'tables':table_number,'docx_sha256':sha(DOCX),
                'chapter_sha256':sha(OUT/'chapter.md'),'protocol_sha256':sha(OUT/'protocol.json'),
                'audit_sha256':sha(OUT/'analysis/audit_receipt.json'),'analysis_sha256':sha(OUT/'analysis/analysis_summary.json'),
                'format':{'font':'Times New Roman','body_points':12,'line_spacing':1.5,'margins_inches':1},
                'all_registered_success_conditions_met':summary['all_registered_success_conditions_met']}
    (OUT / 'build_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(json.dumps(manifest,indent=2))


if __name__ == '__main__':
    build()
