"""Build a new business-school Word dissertation from the verified evidence."""
from __future__ import annotations
import hashlib,json,re
from datetime import datetime,timezone
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor
import pandas as pd

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'study_results/business_dissertation'
TARGET=OUT/'LLM_Replenishment_Business_Dissertation.docx'
INK='20343D';BLUE='24548A';GREEN='27775D'

EXECUTIVE=[
 'Retail replenishment depends on more than predicting sales. A useful order also needs a reliable stock balance, the correct supplier rules and clear permission to spend. This dissertation examines a workflow in which a small local language model works with forecasting, planning and independent evidence checks. It asks when that arrangement helps replenishment and whether the language model adds an advantage over a capable rule-based reader.',
 'The study uses thirty selected store–product series from M5 and thirty from FreshRetailNet-50K, shortened to RetailNet in the graph labels. The complete RetailNet source release was checked, but its performance findings concern the selected panel. Historical observations inform simulations; inventory, supply and costs are disclosed experimental assumptions rather than live retailer outcomes.',
 'Verified inventory repair was the clearest useful mechanism. In the improved M5 pilot, it reduced mean simulated cost by 10.93% and increased fulfilled demand by seven percentage points. In RetailNet’s exploratory calibrated follow-up, the changes were 11.71% and 12.25 points. The original RetailNet spending cap had first held routine orders, showing why business thresholds need to fit a new operating context.',
 'The language-model and rule-based readers produced the same matched actions, costs and service outcomes. The shared recovery workflow therefore explains the benefit; an additional operational advantage from the LLM was not established. Forecast recovery also had limits, and separate fresh-product calculations showed that improved service could coexist with more expiry.',
 'The experiments ran on a Linux CPU cloud computer with four allocated cores, 32 GiB of memory and no GPU. Qwen2.5 1.5B was already pretrained and ran locally through Ollama; the forecasting models were fitted to the retail histories. The managerial recommendation is to begin with verified information, strong planning and clear authority, then test any additional language-model value in a controlled operational pilot.'
]

def sha(p):
 with Path(p).open('rb') as s:return hashlib.file_digest(s,'sha256').hexdigest()
def read(p):return json.loads(Path(p).read_text())
def main_words(text):
 prose='\n'.join(x for x in text.splitlines() if x.strip() and not x.startswith('#') and not x.startswith('[['))
 return len(re.findall(r"\b[\w’'-]+\b",prose))
def field(p,instruction):
 e=OxmlElement('w:fldSimple');e.set(qn('w:instr'),instruction);p._p.append(e)
def link(p,text,target,internal=False):
 e=OxmlElement('w:hyperlink')
 if internal:e.set(qn('w:anchor'),target)
 else:
  rel=p.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);e.set(qn('r:id'),rel)
 r=OxmlElement('w:r');props=OxmlElement('w:rPr')
 color=OxmlElement('w:color');color.set(qn('w:val'),BLUE);props.append(color)
 fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');props.append(fonts)
 r.append(props);t=OxmlElement('w:t');t.text=text;r.append(t);e.append(r);p._p.append(e)
def add_text(p,text):
 pos=0
 for m in re.finditer(r'https?://\S+',text):
  p.add_run(text[pos:m.start()]);url=m.group().rstrip('.;,');link(p,url,url);p.add_run(m.group()[len(url):]);pos=m.end()
 p.add_run(text[pos:])

def get_refs():
 source=ROOT/'study_results/freshretailnet/Dual_Benchmark_Dissertation_Final.docx'
 refs=[p.text for p in Document(source).paragraphs if p.style.name=='Reference']
 prefixes=['Chen, F.','Clark, A. J.','Croston, J. D.','DeHoratius, N., & Raman','DeHoratius, N., Mersereau','Dingdong-Inc.','Gneiting, T.','Hevner, A. R.','Kang, Y.','Lee, H. L.','Li, B., Mellou','Makridakis, S., Spiliotis, E., & Assimakopoulos','Nahmias, S. (1982)','Nahmias, S. (1994)','Parasuraman, R., & Riley','Parasuraman, R., Sheridan','Salinas, D.','Syntetos, A. A.','Turpin, M.','van Donselaar','Wang, R. Y.','Wang, Y., Gu','Yao, S.','Zipkin, P.']
 selected=[]
 for prefix in prefixes:
  matches=[r for r in refs if r.startswith(prefix)];assert len(matches)==1,(prefix,matches)
  selected.append(matches[0].replace('(2022a)','(2022)').replace('(arXiv:2505.16319; updated 2026)','(arXiv:2505.16319)'))
 selected += ['Qwen Team. (2024). Qwen2.5-1.5B-Instruct [Model card]. Hugging Face. https://huggingface.co/Qwen/Qwen2.5-1.5B-Instruct',
  'Gilbert, M. (2025). Predicting temperature variations across Europe using machine learning and geospatial data for climate change impact analysis [Master’s thesis, emlyon business school; user-supplied structural reference].']
 return sorted(selected,key=str.casefold)

def tables():
 num=pd.read_csv(ROOT/'study_results/freshretailnet/FreshRetailNet_numerical_means.csv')
 old=pd.read_csv(ROOT/'study_results/v1/M5_six_hour_results.csv');old=old[old.comparison_group.eq('numerical_7_days')].groupby(['policy','scenario']).agg(cost=('cost','mean'),fill_rate=('fill_rate','mean'),hard_violations=('hard_violations','sum')).reset_index()
 names={'B1':'Seasonal rule','B2':'ML reorder rule','B3':'Constrained optimizer','B4':'Optimizer + approval'}
 scenarios={'normal':'Normal','derived_field_collapse':'Inventory-field fault','feed_gap':'Stale information'}
 def numeric(frame):return [[names[r.policy],scenarios[r.scenario],f'{r.cost:,.2f}',f'{100*r.fill_rate:.2f}%',str(int(r.hard_violations))] for r in frame.itertuples()]
 return {
 'dataset_profiles':('Dataset profiles and scope',['Feature','M5','RetailNet'],[
  ['Source setting','Walmart retail sales','Dingdong fresh retail'],['Study panel','30 series; 3 products; 10 stores','30 series; 26 products; 29 stores'],['Prepared history','1,941 daily positions','90 training + 7 evaluation days'],['Availability information','No released hourly stockout labels used','Hourly sales and stockout indicators'],['Outcome economics','Modelled study monetary conventions','Normalized quantities; simulated cost indices'],['Scope reminder','Reduced panel, not full competition','Full 50,000-series release validated; 30-series performance panel']]),
 'workflow_comparisons':('The matched reader configurations',['Reader','Recovery permission','Shared tools'],[
  ['Rule-based reader','No repair','Same forecast, planner and approval checks'],['Rule-based reader','Verified repair allowed','Same forecast, planner and approval checks'],['Qwen language model','No repair','Same forecast, planner and approval checks'],['Qwen language model','Verified repair allowed','Same forecast, planner and approval checks']]),
 'study_scope':('Completed experimental scope',['Study','Runs','Days per run','Simulation seeds'],[
  ['M5 original numerical study','360','7','30'],['M5 initial live / structured feasibility','8 / 4','2','1'],['M5 improved matched pilot','32','14','2'],['RetailNet primary numerical controls','360','7','30'],['RetailNet primary matched readers','32','7','2'],['RetailNet exploratory calibrated readers','480','7','30'],['RetailNet separate policy calculations','1,140','7','30 per cell']]),
 'model_environment':('Model and cloud environment in plain terms',['Component','Actual study setting'],[
  ['Cloud computer','Linux x86_64; four CPU cores allocated; 32 GiB memory; no GPU'],['Language model','Pretrained Qwen2.5 1.5B, run locally through Ollama'],['LLM training in this project','None: no retail-data fine-tuning'],['Forecast models','Small GRU neural model, LightGBM and reference methods'],['What was fitted','Numerical forecasting models using historical retail sales'],['Main software','Python; CPU-only PyTorch 2.10.0; LightGBM 4.7.0; PyArrow 21.0.0'],['Costs not measured','Cloud, integration, oversight and production operating cost']]),
 'artifact_guide':('Where to inspect the evidence',['Evidence','Repository location'],[
  ['M5 original report and run table','study_results/v1/'],['M5 improved matched results and audit','study_results/v2/'],['RetailNet primary, exploratory and sensitivity tables','study_results/freshretailnet/'],['RetailNet source and independent audit receipts','study_results/freshretailnet/receipts/'],['Complete archived RetailNet records and public source files','study_results/freshretailnet/archive_parts/'],['Restore every original path with hash verification','scripts/freshretailnet_delivery.py'],['Current manuscript, figures and validation','study_results/business_dissertation/']]),
 'm5_numerical_summary':('M5: original numerical means and actual violated checks',['Policy','Scenario','Mean cost','Fill','Violated checks'],numeric(old)),
 'retailnet_numerical_summary':('RetailNet: primary numerical means and actual violated checks',['Policy','Scenario','Cost index','Fill','Violated checks'],numeric(num)),
 'technical_settings':('Selected implementation settings',['Item','Setting and interpretation'],[
  ['Improved-study model tag','ega-qwen2.5:1.5b-v2'],['Working context / CPU threads','8,192 tokens / three threads'],['Generation settings','Temperature 0; fixed seed 42; structured responses'],['Language-model representation','Q4_K_M; pretrained weights, no project fine-tuning'],['Main neural forecaster','GRU; hidden size 32; 10 epochs; lookback 56 days'],['Primary RetailNet spending cap','1,500 in main cost-index conventions'],['Exploratory calibrated cap','1,850; training-only formula, reused evaluation period'],['Main physical shelf life','Simulated 3 days; older stock used first'],['Separate sensitivity lives','1, 3 and 7 days; hypothetical'],['Separate demand multipliers','1, 1.25 and 1.5; assumptions, not observed lost demand'],['Matched main horizon','M5 improved: 14 days; RetailNet: 7 days']])}

def build():
 body=(OUT/'manuscript.md').read_text();appendix=(OUT/'appendices.md').read_text();assert main_words(body)>=15000
 register=read(OUT/'figure_register.json');figs={x['key']:x for x in register['figures']};assert len(figs)==22
 doc=Document();section=doc.sections[0]
 section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
 section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
 section.header_distance=section.footer_distance=Inches(.45);section.different_first_page_header_footer=True
 for s in doc.styles:
  if s.type in [1,2]:
   s.font.name='Times New Roman'
   if s.element.rPr is not None:
    fonts=s.element.rPr.find(qn('w:rFonts'))
    if fonts is not None:
     fonts.set(qn('w:eastAsia'),'Times New Roman');fonts.set(qn('w:cs'),'Times New Roman')
 normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 normal.paragraph_format.widow_control=True
 for name,size in [('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
  s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor.from_string(INK);s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.15
  if name=='Heading 1':s.paragraph_format.page_break_before=True
 for name,size in [('Caption',10.5),('Compact',11),('Reference',12),('Contents Entry',11)]:
  if name not in doc.styles:doc.styles.add_style(name,1)
  s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.paragraph_format.line_spacing=1.1;s.paragraph_format.space_after=Pt(5)
  s.font.bold=False;s.font.italic=False;s.font.color.rgb=RGBColor.from_string(INK)
  s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
 # Remove theme font aliases so office renderers honour the explicit family.
 for fonts in doc.styles.element.iter(qn('w:rFonts')):
  for attr in list(fonts.attrib):
   if 'theme' in attr.lower():del fonts.attrib[attr]
 doc.styles['Reference'].paragraph_format.left_indent=Inches(.22);doc.styles['Reference'].paragraph_format.first_line_indent=Inches(-.22)
 header=section.header.paragraphs[0];header.text='LLM AGENTS FOR RETAIL REPLENISHMENT | M5 & RETAILNET';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 for r in header.runs:r.font.size=Pt(9)
 footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('Page ');field(footer,'PAGE');footer.add_run(' of ');field(footer,'NUMPAGES')
 for r in footer.runs:r.font.size=Pt(10)
 title=doc.add_paragraph();title.paragraph_format.space_before=Pt(100);title.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=title.add_run('LLM Agents for\nRetail Replenishment');r.bold=True;r.font.size=Pt(26);r.font.color.rgb=RGBColor.from_string(INK)
 p=doc.add_paragraph('Evidence from M5 and FreshRetailNet-50K');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(18);p.runs[0].font.size=Pt(17)
 p=doc.add_paragraph('A comparative study of data quality, verified recovery and business controls');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(12)
 p=doc.add_paragraph('Master’s dissertation\nemlyon business school\nOctober 2026');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(65)
 doc.add_heading('Executive summary',level=1)
 for text in EXECUTIVE:doc.add_paragraph(text)
 # Contents are clickable and filled with checked PDF page numbers after first render.
 records=[]
 for text in [body,appendix]:
  for line in text.splitlines():
   if re.match(r'^#{1,2} ',line):records.append({'title':line.split(' ',1)[1],'level':len(line.split(' ',1)[0]),'bookmark':f'h{len(records):03d}'})
 body_n=sum(1 for line in body.splitlines() if re.match(r'^#{1,2} ',line))
 records.insert(body_n,{'title':'References','level':1,'bookmark':'references'})
 record_by_title={r['title']:r for r in records}
 pages=read(OUT/'contents_pages.json') if (OUT/'contents_pages.json').exists() else {}
 doc.add_heading('Contents',level=1)
 for r in records:
  p=doc.add_paragraph(style='Contents Entry');p.paragraph_format.tab_stops.add_tab_stop(Inches(6.22),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
  if r['level']==2:p.paragraph_format.left_indent=Inches(.18)
  link(p,r['title'],r['bookmark'],internal=True);p.add_run('\t'+str(pages.get(r['title'],'—')))
 doc.add_heading('List of figures',level=1)
 ordered_keys=re.findall(r'\[\[FIGURE:(\w+)\]\]',body+appendix)
 for i,key in enumerate(ordered_keys,1):
  f=figs[key];p=doc.add_paragraph(style='Compact');link(p,f'Figure {i}. {f["dataset"]}: {f["title"]}',f'figure{i}',internal=True)
 figure_number=0;table_number=0;heading_id=1;inserted=[]
 table_defs=tables()
 def bookmark(p,name):
  nonlocal heading_id
  start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(heading_id));start.set(qn('w:name'),name)
  end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(heading_id));p._p.insert(0,start);p._p.append(end);heading_id+=1
 def figure(key):
  nonlocal figure_number
  f=figs[key];assert sha(OUT/f['path'])==f['sha256'];figure_number+=1
  p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True;run=p.add_run(f'Figure {figure_number}. {f["dataset"]}: {f["title"]}');run.bold=True;run.font.color.rgb=RGBColor.from_string(GREEN if f['dataset']=='RETAILNET' else BLUE if f['dataset']=='M5' else INK);bookmark(p,f'figure{figure_number}')
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(str(OUT/f['path']),width=Inches(6.22))
  p=doc.add_paragraph('Source and scope: '+f['caption'],style='Caption');p.paragraph_format.keep_together=True
  inserted.append({'number':figure_number,**f})
 def table(key):
  nonlocal table_number
  title,headers,rows=table_defs[key];table_number+=1
  p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True;p.add_run(f'Table {table_number}. {title}').bold=True
  t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=True
  for i,v in enumerate(headers):t.rows[0].cells[i].text=v
  prop=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(prop)
  for row in rows:
   cells=t.add_row().cells
   for cell,val in zip(cells,row):cell.text=str(val)
  for j,row in enumerate(t.rows):
   prop=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(prop)
   for cell in row.cells:
    for p in cell.paragraphs:
     p.style='Compact';p.paragraph_format.space_after=Pt(3)
     for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(11);run.bold=(j==0)
    if j==0:
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF2');cell._tc.get_or_add_tcPr().append(shade)
  doc.add_paragraph().paragraph_format.space_after=Pt(2)
 def manuscript(text):
  for line in text.splitlines():
   if not line.strip():continue
   h=re.match(r'^(#{1,3}) (.*)',line)
   if h:
    p=doc.add_heading(h[2],level=len(h[1]));r=record_by_title.get(h[2])
    if r:bookmark(p,r['bookmark'])
   elif line.startswith('[[FIGURE:'):figure(line[9:-2])
   elif line.startswith('[[TABLE:'):table(line[8:-2])
   else:add_text(doc.add_paragraph(),line)
 manuscript(body)
 p=doc.add_heading('References',level=1);bookmark(p,'references')
 refs=get_refs()
 for text in refs:add_text(doc.add_paragraph(style='Reference'),text)
 (OUT/'references.md').write_text('\n\n'.join(refs)+'\n')
 manuscript(appendix)
 doc.core_properties.title='LLM Agents for Retail Replenishment: Evidence from M5 and FreshRetailNet-50K'
 doc.core_properties.subject='Business-school rewrite of completed studies; concise findings and dataset-labelled graphs'
 doc.core_properties.author='';doc.core_properties.keywords='M5; FreshRetailNet; replenishment; Qwen; business; verified recovery'
 doc.core_properties.modified=datetime.now(timezone.utc).replace(tzinfo=None)
 doc.save(TARGET)
 (OUT/'heading_register.json').write_text(json.dumps(records,indent=2)+'\n')
 counts={}
 for section in re.split(r'(?m)^# ',body)[1:]:counts[section.splitlines()[0]]=main_words('\n'.join(section.splitlines()[1:]))
 manifest={'status':'BUILT_PENDING_RENDER','main_prose_words_excluding_front_matter_references_appendices':main_words(body),'chapter_words':counts,'figures':inserted,'tables':table_number,'references':len(refs),'format':{'font':'Times New Roman','body_points':12,'line_spacing':1.5,'margins_inches':1,'alignment':'justified'},'docx_sha256':sha(TARGET),'inputs':register['inputs'],'manuscript_sha256':sha(OUT/'manuscript.md'),'appendices_sha256':sha(OUT/'appendices.md'),'figure_register_sha256':sha(OUT/'figure_register.json'),'builder_sha256':sha(__file__),'contents_has_verified_pages':bool(pages),'built_at_utc':datetime.now(timezone.utc).isoformat()}
 (OUT/'build_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
 print(json.dumps({k:manifest[k] for k in ['status','main_prose_words_excluding_front_matter_references_appendices','chapter_words','tables','references']},indent=2))

if __name__=='__main__':build()
