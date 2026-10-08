"""Render the completed bounded FreshRetailNet study without changing its evidence.

Consumes immutable CSVs and audit receipts. Produces standalone PDF/Markdown,
scientific figures and an updated eight-chapter business-school dissertation.
A complete build requires the registered core, forecast and sensitivity runs.
"""
from __future__ import annotations
import argparse,hashlib,json,math,re,shutil,textwrap
from copy import deepcopy
from datetime import datetime,timezone
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,PageBreak,KeepTogether
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
DEFAULT_SOURCE=Path('/workspace/attachments/05bb558e-50b1-4067-b21f-af48d109fc7e/Evidence_Gated_Autonomy_M5_FreshRetailNet_Full_Dissertation.docx')
ARMS=['parser_no_recovery','parser_recovery','llm_no_recovery','llm_recovery']
ARM_LABEL={'parser_no_recovery':'Parser','parser_recovery':'Parser + repair','llm_no_recovery':'LLM','llm_recovery':'LLM + repair'}
SCENARIOS=['normal','derived_field_collapse','feed_gap','capacity_cut']
SCENARIO_LABEL={'normal':'Normal','derived_field_collapse':'Inventory-field fault','feed_gap':'Stale feed','capacity_cut':'Capacity cut'}
METHOD_LABEL={'raw_zero':'No recovery','profile':'Hourly profile','gbm':'Hourly LightGBM','seasonal_naive_raw':'Seasonal naive','lightgbm_raw':'LightGBM, raw','lightgbm_profile_recovered':'LightGBM, profile','lightgbm_gbm_recovered':'LightGBM, learned recovery'}
AUDIT_LABEL={'runs':'Completed runs','decisions':'Decision days','numerical_runs':'Conventional core runs','agent_runs':'Matched controller runs','physical_API_attempts':'Fresh model attempts','physical_API_response_records':'Fresh model responses','tokens':'Recorded model tokens','client_or_schema_error_events':'Client or schema errors','physical_extraction_responses':'Fresh numerical-extraction responses','physical_selector_responses':'Fresh recovery selections','committed_true_constraint_violations':'Violated active-rule checks in requested plans','held_decisions':'Held decision days','selector_instruction_compliant':'Instruction-compliant recovery selections','selector_semantic_error_responses':'Incorrect recovery selections','parser_llm_matched_pairs':'Matched parser–LLM pairs','parser_llm_pairs_primary_metrics_equal':'Pairs with identical primary outcomes','parser_llm_pairs_all_action_days_equal':'Pairs with identical daily actions','objects_sha256_verified':'Verified decision objects','sqlite_events_hash_verified':'Verified linked event records','physical_mass_balance_cases_passed':'Verified physical stock balances','forbidden_model_context_findings':'Forbidden model-context findings'}
PALETTE=['#245B68','#56A391','#CC8B3B','#A75651','#738298','#9573A6']
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':12,'axes.spines.top':False,'axes.spines.right':False,'axes.titleweight':'bold','axes.labelcolor':'#253A41','figure.facecolor':'white','axes.facecolor':'white','savefig.facecolor':'white','axes.prop_cycle':matplotlib.cycler(color=PALETTE)})


def sha(path):
    with Path(path).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def json_read(path,default=None):
    return json.loads(Path(path).read_text()) if Path(path).exists() else default

def number(x,d=2):
    return 'not available' if pd.isna(x) else f'{float(x):,.{d}f}'
def percent(x,d=1):return number(100*float(x),d)+'%'
def label(method):return METHOD_LABEL.get(str(method),str(method).replace('_',' '))
def csv(path):return pd.read_csv(path)
def require(path):
    if not Path(path).exists():raise FileNotFoundError(f'Authoritative input missing: {path}')
    return Path(path)
def scalar_frame(d):
    rows=[]
    def recurse(value,prefix=''):
        if isinstance(value,dict):
            for k,v in value.items():recurse(v,f'{prefix}.{k}' if prefix else k)
        elif isinstance(value,(str,int,float,bool)) or value is None:rows.append([prefix,str(value)])
    recurse(d);return rows


class Study:
    def __init__(self,root,source,allow_partial=False):
        self.root=Path(root);self.source=Path(source);self.out=self.root/'report';self.out.mkdir(parents=True,exist_ok=True)
        self.figdir=self.out/'figures';self.figdir.mkdir(exist_ok=True)
        self.protocol=json_read(require(self.root/'protocol.json'))
        self.selection=json_read(require(self.root/'data/selection.json'))
        self.population=json_read(require(self.root/'data/train_validation.json'))
        self.eval_population=json_read(self.root/'data/eval_validation.json',{})
        self.panel=csv(require(self.root/'data/selected_series.csv'))
        self.forecast=csv(require(self.root/'forecast/forecast_metrics.csv'))
        self.hourly=csv(require(self.root/'forecast/recovery_hour_metrics.csv'))
        self.daily_recovery=csv(require(self.root/'forecast/recovery_daily_metrics.csv'))
        self.forecasts=csv(require(self.root/'forecast/forecasts_eval.csv'))
        self.forecast_summary=json_read(require(self.root/'forecast/summary.json'))
        self.forecast_selection=json_read(require(self.root/'forecast/selection_receipt.json'))
        self.leakage=json_read(require(self.root/'forecast/leakage_audit.json'))
        self.numeric=csv(require(self.root/'numerical/summary.csv'))
        self.agent=pd.concat([csv(require(self.root/f'agents/{arm}/summary.csv')).assign(arm=arm) for arm in ARMS],ignore_index=True)
        self.sensitivity=csv(require(self.root/'policy_sensitivity/runs.csv'))
        self.sensitivity_summary=json_read(self.root/'policy_sensitivity/summary.json',{})
        self.audit=json_read(require(self.root/'audit/audit.json'))
        self.forecast_independent_audit=json_read(require(self.root/'forecast/independent_audit.json'))
        self.sensitivity_audit=json_read(require(self.root/'policy_sensitivity/audit.json'))
        self.sensitivity_independent_audit=json_read(require(self.root/'development/policy_sensitivity_independent_audit.json'))
        if not allow_partial:
            assert self.audit.get('status')=='VERIFIED_COMPLETE', 'Primary audit is not fully verified'
            assert str(self.forecast_independent_audit.get('status','')).lower()=='verified', 'Forecast independent audit is incomplete'
            assert self.sensitivity_audit.get('status')=='VERIFIED', 'Sensitivity audit is incomplete'
            assert self.sensitivity_independent_audit.get('status')=='VERIFIED', 'Sensitivity independent audit is incomplete'
        self.model_calls=csv(self.root/'audit/model_calls.csv') if (self.root/'audit/model_calls.csv').exists() else pd.DataFrame()
        self.decisions=csv(self.root/'audit/decision_audit.csv') if (self.root/'audit/decision_audit.csv').exists() else pd.DataFrame()
        self.grounding=csv(self.root/'audit/grounding_aggregates.csv') if (self.root/'audit/grounding_aggregates.csv').exists() else pd.DataFrame()
        self.m5=csv(ROOT/'study_results/v2/M5_v2_results.csv').replace({'llm_v2_no_recovery':'llm_no_recovery','llm_v2_recovery':'llm_recovery'})
        if not allow_partial:
            assert len(self.numeric)==360,f'Expected360 numerical runs; found{len(self.numeric)}'
            assert len(self.agent)==32,f'Expected32 agent runs; found{len(self.agent)}'
            assert self.agent.groupby('arm').size().eq(8).all()
            assert int(self.numeric.trace_count.sum())==2520
            assert int(self.agent.trace_count.sum())==224
            assert len(self.sensitivity)==1140,f'Expected1140 policy runs; found{len(self.sensitivity)}'
            assert self.forecast_summary['status']=='completed'
            for p in ['numerical/execution_complete.json','agents/execution_complete.json','forecast/completion_receipt.json']:
                receipt=json_read(require(self.root/p));assert receipt.get('status','').lower() in ['complete','completed','verified','pass']
        self.numeric_means=self.numeric.groupby(['scenario','policy'],sort=False).agg(
            runs=('cost','size'),cost=('cost','mean'),fill_rate=('fill_rate','mean'),stockout_rate=('stockout_rate','mean'),
            mean_inventory=('mean_inventory','mean'),bullwhip=('bullwhip','mean'),hard_violations=('hard_violations','sum'),held_decisions=('held_decisions','sum'),reference_deviations=('reference_deviations','sum'),executed_actions=('executed_actions','sum')).reset_index()
        self.agent_means=self.agent.groupby(['arm','scenario'],sort=False).agg(
            runs=('cost','size'),cost=('cost','mean'),fill_rate=('fill_rate','mean'),stockout_rate=('stockout_rate','mean'),
            mean_inventory=('mean_inventory','mean'),bullwhip=('bullwhip','mean'),hard_violations=('hard_violations','sum'),held_decisions=('held_decisions','sum'),
            llm_calls=('llm_calls','sum'),tokens=('tokens','sum')).reset_index()
        self.pairs=self.pair_controls(self.agent)
        self.comparison=self.cross_dataset()
        self.figures=[];self.sections=[];self.inputs={}
        for p in [self.source,self.root/'protocol.json',self.root/'data/selected_series.csv',self.root/'forecast/forecast_metrics.csv',self.root/'forecast/recovery_hour_metrics.csv',self.root/'forecast/selection_receipt.json',self.root/'forecast/independent_audit.json',self.root/'numerical/summary.csv',self.root/'policy_sensitivity/runs.csv',self.root/'policy_sensitivity/manifest.json',self.root/'policy_sensitivity/audit.json',self.root/'development/policy_sensitivity_independent_audit.json',self.root/'audit/audit.json',ROOT/'study_results/v2/M5_v2_results.csv']+[self.root/f'agents/{arm}/summary.csv' for arm in ARMS]:
            self.inputs[str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p)]={'sha256':sha(p),'bytes':p.stat().st_size}
        for p in [self.root/name for name in ['data/selection.json','data/train_validation.json','data/eval_validation.json','forecast/summary.json','forecast/recovery_daily_metrics.csv','forecast/forecasts_eval.csv','forecast/leakage_audit.json','audit/model_calls.csv','audit/decision_audit.csv','audit/grounding_aggregates.csv','audit/recovery_cases.json','audit/violation_families.json']]+[ROOT/'study_results/dissertation/dissertation_build_manifest.json',ROOT/'study_results/dissertation/Evidence_Gated_Autonomy_Dissertation_Final.docx']:
            self.inputs[str(p.relative_to(ROOT))]={'sha256':sha(require(p)),'bytes':p.stat().st_size}
        load_followup(self,allow_partial)
        self.numeric.to_csv(self.out/'FreshRetailNet_numerical_all_runs.csv',index=False)
        self.agent.to_csv(self.out/'FreshRetailNet_agent_all_runs.csv',index=False)
        self.sensitivity.to_csv(self.out/'FreshRetailNet_policy_sensitivity_all_runs.csv',index=False)
        self.numeric_means.to_csv(self.out/'FreshRetailNet_numerical_means.csv',index=False)
        self.agent_means.to_csv(self.out/'FreshRetailNet_agent_means.csv',index=False)
        self.pairs.to_csv(self.out/'FreshRetailNet_parser_llm_pairs.csv',index=False)
        self.comparison.to_csv(self.out/'M5_FreshRetailNet_descriptive_comparison.csv',index=False)

    @staticmethod
    def pair_controls(frame):
        pairs=[]
        for suffix in ['no_recovery','recovery']:
            p=frame[frame.arm.eq('parser_'+suffix)].set_index(['scenario','seed'])
            l=frame[frame.arm.eq('llm_'+suffix)].set_index(['scenario','seed'])
            assert set(p.index)==set(l.index)
            for key in p.index:
                a,b=p.loc[key],l.loc[key]
                pairs.append({'repair':suffix,'scenario':key[0],'seed':int(key[1]),'parser_cost':float(a.cost),'llm_cost':float(b.cost),
                    'cost_difference':float(b.cost-a.cost),'fill_difference_pp':float(100*(b.fill_rate-a.fill_rate)),
                    'violations_difference':float(b.hard_violations-a.hard_violations),'holds_difference':float(b.held_decisions-a.held_decisions),
                    'cost_fill_identical':bool(a.cost==b.cost and a.fill_rate==b.fill_rate)})
        return pd.DataFrame(pairs)

    def cross_dataset(self):
        rows=[]
        for name,frame,days in [('M5',self.m5,14),('FreshRetailNet',self.agent,7)]:
            means=frame.groupby(['arm','scenario'])[['cost','fill_rate','held_decisions','hard_violations']].mean()
            pairs=self.pair_controls(frame)
            for sc in SCENARIOS:
                before=means.loc[('llm_no_recovery',sc)];after=means.loc[('llm_recovery',sc)]
                matched=pairs[pairs.scenario.eq(sc)]
                rows.append({'dataset':name,'scenario':sc,'series':30,'seeds':2,'days':days,
                    'repair_cost_change_pct':100*(after.cost/before.cost-1) if before.cost else np.nan,
                    'repair_fill_change_pp':100*(after.fill_rate-before.fill_rate),'repair_holds_change_per_run':after.held_decisions-before.held_decisions,
                    'parser_llm_pairs':len(matched),'cost_fill_identical_pairs':int(matched.cost_fill_identical.sum()),
                    'mean_llm_minus_parser_cost_pct':float(np.mean(100*(matched.llm_cost/matched.parser_cost-1))),
                    'comparison_type':'Within-dataset relative effect; cross-dataset descriptive; economic units and horizons differ'})
        return pd.DataFrame(rows)

    def savefig(self,key,title,caption,fig):
        fig.suptitle(title,fontsize=14,fontweight='bold',color='#173D43',y=1.01)
        fig.tight_layout(rect=getattr(fig,'_report_layout_rect',(0,0,1,1)))
        p=self.figdir/f'{key}.png';fig.savefig(p,dpi=220,bbox_inches='tight');plt.close(fig)
        self.figures.append({'key':key,'title':title,'caption':caption,'path':str(p),'sha256':sha(p)})
        return p

    def plots(self):
        fig,ax=plt.subplots(figsize=(8,3.8))
        values=[self.population['operating_hour_stockout_fraction'],self.population['daily_any_stockout_fraction'],float(self.panel.censor_rate.mean())]
        bars=ax.bar(['All training hours\nflagged out of stock','All training days\nwith any stockout','Selected panel hours\n(first 62 days)'],values,color=PALETTE[:3]);ax.yaxis.set_major_formatter(PercentFormatter(1));ax.set_ylim(0,max(values)*1.25);ax.set_ylabel('Share of stated denominator')
        for b,v in zip(bars,values):ax.text(b.get_x()+b.get_width()/2,v+.01,percent(v),ha='center')
        self.savefig('01_source_availability','Historical availability is observable; lost demand is not',
            'Training population rates use 72 million operating-hour flags and 4.5 million daily rows. The selected-panel rate uses its first 62 training days. Daily and hourly percentages have different denominators. None measures the quantity of naturally lost demand.',fig)
        h=self.hourly[(self.hourly.series_id=='ALL')&(self.hourly.mask_duration.astype(str)!='all')].copy();durations=[2,4,8];methods=sorted(h.method.unique())
        fig,ax=plt.subplots(figsize=(8,4))
        for i,m in enumerate(methods):
            a=h[h.method.eq(m)].set_index(h[h.method.eq(m)].mask_duration.astype(int));vs=[a.loc[k,'wape'] for k in durations];ax.bar(np.arange(3)+(i-(len(methods)-1)/2)*.24,vs,.24,label=label(m))
        ax.set_xticks(range(3),['2 hours','4 hours','8 hours']);ax.set_ylabel('WAPE on hidden operating-hour sales');ax.yaxis.set_major_formatter(PercentFormatter(1));ax.legend(fontsize=9)
        self.savefig('02_artificial_mask_recovery','Recovery accuracy where a target can actually be checked',
            'Originally observed sales are hidden in 2-, 4- or 8-hour validation blocks. Models fit only earlier data. The target is deliberately withheld sales, rather than unobserved natural stockout demand. Mask cases overlap and are not independent retailers.',fig)
        f=self.forecast[(self.forecast.partition=='eval')&(self.forecast.fold.astype(str)=='pooled')&(self.forecast.series_id=='ALL')]
        methods=list(f.method.unique());fig,axes=plt.subplots(1,2,figsize=(11,4.4))
        for ax,target,title in zip(axes,['observed_daily_proxy','fully_in_stock_daily_sales'],['All released daily sales','Fully stocked evaluation days']):
            a=f[f.target_kind.eq(target)].set_index('method');vs=[a.loc[m,'wape'] for m in methods];ax.bar(range(len(methods)),vs,color=PALETTE[:len(methods)]);ax.set_xticks(range(len(methods)),[label(m).replace('LightGBM, ','LGBM\n') for m in methods],rotation=25,ha='right');ax.set_title(f'{title}\n(n={int(a.n.iloc[0])})');ax.yaxis.set_major_formatter(PercentFormatter(1));ax.set_ylabel('Daily WAPE')
        self.savefig('03_forecast_accuracy','Forecast quality depends on which sales are observable',
            'All four frozen forecasts cover the same 30 series and seven-day holdout. All-day sales can be censored; fully stocked days are conditionally observable. Future masks, weather, promotion and actual sales are never forecast features.',fig)
        fig,ax=plt.subplots(figsize=(9,4));a=self.forecasts
        truth=a.drop_duplicates(['series_id','day_index']).groupby('dt').observed_sales.sum();ax.plot(pd.to_datetime(truth.index),truth.values,'o-',label='Observed sales (censored)',color='#253A41',lw=2.5)
        for m in sorted(a.method.unique()):
            g=a[a.method.eq(m)].groupby('dt').prediction.sum();ax.plot(pd.to_datetime(g.index),g.values,'o--',label=label(m),alpha=.8)
        ax.set_ylabel('Sum of globally normalized sales amounts');ax.tick_params(axis='x',rotation=25);ax.legend(fontsize=8,ncol=2)
        self.savefig('04_daily_forecasts','The holdout forecast is made before the seven-day window',
            'Panel aggregate predictions are made at the final training origin. Later forecast lags use predictions. Globally normalized sales are neither physical quantities nor a currency.',fig)
        fig,axes=plt.subplots(2,3,figsize=(11,7))
        for j,sc in enumerate(SCENARIOS[:3]):
            g=self.numeric[self.numeric.scenario.eq(sc)];policies=['B1','B2','B3','B4'];means=g.groupby('policy');cost=[means.get_group(p).cost.mean() for p in policies];std=[means.get_group(p).cost.std() for p in policies]
            axes[0,j].bar(policies,cost,yerr=std,capsize=3,color=PALETTE[:4]);axes[0,j].set_title(SCENARIO_LABEL[sc]);axes[0,j].set_ylabel('Seven-day simulated cost index');axes[1,j].bar(policies,[means.get_group(p).fill_rate.mean() for p in policies],color=PALETTE[:4]);axes[1,j].yaxis.set_major_formatter(PercentFormatter(1));axes[1,j].set_ylim(0,1.05);axes[1,j].set_ylabel('Mean fulfilled-demand share')
        self.savefig('05_numerical_cost_service','Conventional policies establish a demanding baseline',
            'Thirty seeds per policy and scenario. Error bars show one standard deviation across conditional simulation seeds. B1: seasonal base-stock; B2: LightGBM (s,S); B3: native GRU-NB with MILP; B4: B3 with an evidence gate. Physical stock expires by FIFO, while the inherited MILP remains age-unaware.',fig)
        fig,axes=plt.subplots(1,2,figsize=(10,4));g=self.numeric.groupby(['scenario','policy'])[['hard_violations','held_decisions']].sum().unstack('policy')
        for ax,col,title in zip(axes,['hard_violations','held_decisions'],['Actual hard-constraint violations','Held decision-days']):
            g[col].reindex(SCENARIOS[:3]).rename(index=SCENARIO_LABEL).plot.bar(ax=ax,color=PALETTE[:4],width=.8);ax.set_title(title);ax.set_xlabel('');ax.tick_params(axis='x',rotation=20);ax.set_ylabel('Count across 30 runs');ax.legend(fontsize=8)
        self.savefig('06_numerical_safety','A safe hold has a service cost that should be measured',
            'Counts are summed across 30 runs per cell. Hard violations use actual active constraints and differ from clean-reference deviations. Holds count decision days. A low violation count should be assessed alongside service.',fig)
        fig,axes=plt.subplots(2,2,figsize=(9.5,7.5))
        for ax,sc in zip(axes.flat,SCENARIOS):
            g=self.agent_means[self.agent_means.scenario.eq(sc)].set_index('arm');x=np.arange(4);ax.bar(x,[g.loc[a,'cost'] for a in ARMS],color=PALETTE[:4]);ax.set_xticks(x,[ARM_LABEL[a].replace(' + ','\n+ ') for a in ARMS],fontsize=12);ax.set_title(SCENARIO_LABEL[sc]);ax.set_ylabel('Seven-day cost index')
            for i,a in enumerate(ARMS):ax.text(i,g.loc[a,'cost'],f"Fill {percent(g.loc[a,'fill_rate'])}",ha='center',va='bottom',fontsize=11)
            ax.set_ylim(0,float(g.cost.max())*1.24)
        self.savefig('07_agent_factorial','The LLM shares the same forecasts, tools and safeguards',
            'Four matched controller/repair configurations, four scenarios and two seeds, each using seven days. Repair is verified inventory-record reconciliation. It is not statistical recovery of censored customer demand. All configurations share the numerical tools and source evidence.',fig)
        fig,axes=plt.subplots(1,2,figsize=(10,4));g=self.comparison[self.comparison.scenario.eq('derived_field_collapse')]
        short=[('M5\n14 days\n2 seeds' if x=='M5' else 'Fresh, fixed\n7 days\n2 seeds' if 'fixed' in x else 'Fresh, calibrated\n7 days\n30 seeds') for x in g.dataset];axes[0].bar(short,g.repair_cost_change_pct,color=PALETTE[:len(g)]);axes[0].axhline(0,color='#888',lw=.8);axes[0].set_ylabel('Repair-on minus repair-off cost (%)');axes[1].bar(short,g.repair_fill_change_pp,color=PALETTE[:len(g)]);axes[1].axhline(0,color='#888',lw=.8);axes[1].set_ylabel('Fill-rate change (percentage points)')
        for ax in axes:ax.tick_params(axis='x',labelsize=10)
        self.savefig('08_m5_fresh_comparison','Two information environments; two separately estimated effects',
            'Separate within-dataset repair effects under an inventory-field fault. M5 uses 14 days and FreshRetailNet seven; their quantity units, expiry assumptions and economics differ. The calibrated result is an exploratory reused-holdout follow-up with 30 simulator seeds; the other bars have two. These bars do not estimate a commercial return or the causal effect of changing datasets.',fig)
        enriched=self.forecasts.merge(self.panel[['series_id','censor_stratum']],on='series_id',validate='many_to_one')
        strata=[]
        for target in ['observed_daily_proxy','fully_in_stock_daily_sales']:
            a=enriched if target=='observed_daily_proxy' else enriched[enriched.fully_stocked]
            for (method,stratum),g in a.groupby(['method','censor_stratum']):
                denominator=float(g.observed_sales.sum())
                strata.append({'target_kind':target,'method':method,'training_censor_stratum':int(stratum),'n':len(g),'sales_denominator':denominator,'wape':float(np.abs(g.prediction-g.observed_sales).sum()/denominator) if denominator else np.nan,'signed_bias':float((g.prediction-g.observed_sales).sum()/denominator) if denominator else np.nan})
        self.forecast_strata=pd.DataFrame(strata);self.forecast_strata.to_csv(self.out/'FreshRetailNet_forecast_by_training_censor_stratum.csv',index=False)
        fig,axes=plt.subplots(1,2,figsize=(10,4))
        for ax,target,title in zip(axes,['observed_daily_proxy','fully_in_stock_daily_sales'],['All recorded sales','Fully stocked days']):
            for method,g in self.forecast_strata[self.forecast_strata.target_kind.eq(target)].groupby('method'):
                g=g.sort_values('training_censor_stratum');ax.plot(g.training_censor_stratum,g.wape,'o-',label=label(method))
            ax.set_xticks([0,1,2],['Lower','Middle','Higher']);ax.set_xlabel('Training-only censoring stratum');ax.set_title(title);ax.set_ylabel('Final daily WAPE');ax.yaxis.set_major_formatter(PercentFormatter(1));ax.legend(fontsize=7)
        self.savefig('15_forecast_censor_strata','Historical censoring changes the forecasting problem','Ten series per stratum, defined before evaluation using the first 62 training days. Errors are weighted by the sum of actual scored sales within each group. Evaluation availability selects the fully stocked scoring subset only; it does not change membership or forecasts. These small strata are descriptive.',fig)
        self.sensitivity_plots()
        if not self.model_calls.empty:
            fig,ax=plt.subplots(figsize=(8,4));groups=[];values=[]
            for col in ['arm','run_arm']:
                if col in self.model_calls:
                    groups=self.model_calls.groupby(col).size();break
            if len(groups):
                ax.bar([ARM_LABEL.get(a,a) for a in groups.index],groups.values,color=PALETTE[2:4]);ax.set_ylabel('Captured fresh physical model attempts')
            else:
                ax.bar(['Physical attempts'],[len(self.model_calls)],color=PALETTE[2]);ax.set_ylabel('Captured attempt records')
            self.savefig('11_local_model_calls','Inference is an additional operating burden',
                'Fresh physical attempt records from the independent audit. Cache imports and typed extraction outputs are not additional calls. Local API charges are zero; CPU hosting and oversight costs are unpriced.',fig)

    def sensitivity_plots(self):
        d=self.sensitivity;columns=set(d.columns)
        # The sensitivity owner may use explicit names; fail rather than infer an unsupported metric.
        view=next((x for x in ['forecast_view','view','forecast_history','history'] if x in columns),None)
        calculator=next((x for x in ['calculator','policy','controller'] if x in columns),None)
        shelf=next((x for x in ['shelf_life_days','shelf_life'] if x in columns),None)
        multiplier=next((x for x in ['demand_multiplier','assumed_demand_multiplier'] if x in columns),None)
        waste=next((x for x in ['waste_fraction','expiry_fraction','waste_rate','spoilage_rate','expired_fraction','waste_per_opening_plus_receipts'] if x in columns),None)
        fill=next((x for x in ['fill_rate','service'] if x in columns),None)
        cost=next((x for x in ['cost','total_cost','cost_index','total_cost_index'] if x in columns),None)
        if None in [view,calculator,shelf,multiplier,fill,cost]:raise ValueError(f'Unsupported policy schema: {sorted(columns)}')
        self.sens_columns={'view':view,'calculator':calculator,'shelf':shelf,'multiplier':multiplier,'waste':waste,'fill':fill,'cost':cost}
        core=d[d.experiment.eq('F3')].copy() if 'experiment' in d else d[d[calculator].astype(str).isin(['order_up_to','age_aware'])].copy()
        if core.empty:core=d[d[shelf].notna()&d[multiplier].notna()].copy()
        self.sens_core=core
        if waste:
            fig,axes=plt.subplots(3,1,figsize=(8,10))
            for ax,L in zip(axes,[1,3,7]):
                a=core[core[shelf].eq(L)&core[multiplier].eq(1)].groupby([view,calculator])[[fill,waste]].mean().reset_index()
                for i,r in a.iterrows():
                    raw='raw' in str(r[view]);name=('Raw' if raw else 'Recovered')+', '+str(r[calculator]).replace('_','-')
                    ax.scatter(r[waste],r[fill],s=100,color=PALETTE[0] if raw else PALETTE[2],marker='s' if r[calculator]=='age_aware' else 'o',label=name)
                ax.set_title(f'Assumed shelf life: {L} day'+('s' if L>1 else ''));ax.set_xlabel('Expired quantity / received + opening stock');ax.set_ylabel('Fulfilled-demand share');ax.yaxis.set_major_formatter(PercentFormatter(1));ax.xaxis.set_major_formatter(PercentFormatter(1));ax.margins(.25)
            handles,labels=axes[0].get_legend_handles_labels();fig.legend(handles,labels,loc='lower center',bbox_to_anchor=(.5,.008),ncol=2,fontsize=11);fig._report_layout_rect=(0,.075,1,1)
            self.savefig('09_service_waste_frontier','Fresh retail exposes a service–waste trade-off',
                'Independent calculator experiment, with 30 shared seeds per cell and demand multiplier 1. Shelf lives are hypothetical. Quantity and economic conventions differ from the core MILP experiment, so their absolute results are not pooled.',fig)
        fig,axes=plt.subplots(1,2,figsize=(10,4))
        for ax,calc in zip(axes,['order_up_to','age_aware']):
            a=core[core[calculator].eq(calc)];means=a.groupby([shelf,multiplier,view])[cost].mean().unstack(view)
            raw=next((x for x in means.columns if 'raw' in str(x)),None);rec=next((x for x in means.columns if x!=raw),None)
            if raw is None or rec is None:raise ValueError('Sensitivity needs raw and recovered forecast views')
            contrasts=100*(means[rec]/means[raw]-1);mat=contrasts.unstack(multiplier).reindex([1,3,7]);v=max(1,float(np.nanmax(np.abs(mat.to_numpy()))));im=ax.imshow(mat,cmap='RdBu_r',vmin=-v,vmax=v,aspect='auto');ax.set_title(calc.replace('_',' '));ax.set_xticks(range(len(mat.columns)),mat.columns);ax.set_yticks(range(3),['1 day','3 days','7 days']);ax.set_xlabel('Assumed demand multiplier');ax.set_ylabel('Hypothetical shelf life')
            for i in range(3):
                for j in range(len(mat.columns)):ax.text(j,i,f'{mat.iloc[i,j]:+.1f}%',ha='center',va='center',fontsize=11)
            fig.colorbar(im,ax=ax,label='Recovered minus raw mean indexed cost (%)')
        self.savefig('10_sensitivity_costs','Demand recovery must earn its keep under different assumptions',
            'Relative change from raw to recovered forecast history within each common-demand cell. Thirty seeds, shelf lives 1/3/7 days and demand multipliers 1/1.25/1.5. Negative values mean lower assumed indexed cost. They are not measured profit or purchasing savings.',fig)
        self.sensitivity_means=core.groupby([calculator,view,shelf,multiplier]).agg(runs=(cost,'size'),cost=(cost,'mean'),fill_rate=(fill,'mean'),**({'waste_fraction':(waste,'mean')} if waste else {})).reset_index()
        self.sensitivity_means.to_csv(self.out/'FreshRetailNet_sensitivity_means.csv',index=False)

    def add(self,title,paragraphs=(),tables=(),figures=()):
        self.sections.append({'title':title,'paragraphs':list(paragraphs),'tables':list(tables),'figures':list(figures)})

    def make_narrative(self):
        p=self.population;hours=self.hourly[(self.hourly.series_id=='ALL')&(self.hourly.mask_duration.astype(str)=='all')]
        f=self.forecast[(self.forecast.partition=='eval')&(self.forecast.fold.astype(str)=='pooled')&(self.forecast.series_id=='ALL')]
        winner=self.forecast_selection['forecast']['method'];rec=self.forecast_selection['recovery']['method'];valid=self.forecast_selection['forecast']['score']
        identical=int(self.pairs.cost_fill_identical.sum());pair_count=len(self.pairs)
        means=self.agent_means.set_index(['arm','scenario']);before=means.loc[('llm_no_recovery','derived_field_collapse')];after=means.loc[('llm_recovery','derived_field_collapse')]
        delta=100*(after.cost/before.cost-1);fill_delta=100*(after.fill_rate-before.fill_rate)
        self.findings={'parser_llm_identical_pairs':identical,'parser_llm_total_pairs':pair_count,'repair_cost_change_pct':float(delta),'repair_fill_change_pp':float(fill_delta),'forecast_selected':winner,'recovery_selected':rec}
        self.add('Executive findings',[])
        self.add('1. Study scope, timing and evidence boundaries',[],tables=[('Completed scope', ['Experiment','Series','Seeds','Days','Runs / decisions'],[
            ['Core numerical policies','30','30','7',f'{len(self.numeric)} / {int(self.numeric.trace_count.sum())}'],
            ['Matched parser/LLM factorial','30','2','7',f'{len(self.agent)} / {int(self.agent.trace_count.sum())}'],
            ['Age-aware sensitivity grid','30','30','7','1,080 runs'],
            ['Stockout-gate mechanism ablation','30','30','7','60 runs'],
            ['Forecast final holdout','30','Not simulation seeds','7','210 targets × 4 methods']
        ])])
        self.add('2. Source validation and computational panel',[],tables=[('Selected panel by first 62-day censoring stratum',['Stratum','Series','Mean hourly stockout share','Mean normalized sales'],[[str(k),str(len(g)),percent(g.censor_rate.mean(),2),number(g.mean_sales.mean(),4)] for k,g in self.panel.groupby('censor_stratum')])],figures=['01_source_availability'])
        self.add('3. F1: demand recovery where a target is identifiable',[],tables=[('Pooled artificial-mask operating-hour recovery',['Method','N predictions','MAE','WAPE','Signed bias'],[[label(r.method),str(int(r.n)),number(r.mae,5),percent(r.wape,2),percent(r.relative_bias,2)] for r in hours.itertuples()])],figures=['02_artificial_mask_recovery'])
        self.add('4. F2: forecasting censored and fully stocked outcomes',[],tables=[('Final seven-day forecast scores',['Method','Target','N','WAPE','Bias'],[[label(r.method),'All sales' if r.target_kind=='observed_daily_proxy' else 'Fully stocked',str(int(r.n)),percent(r.wape,2),percent(r.relative_bias,2)] for r in f.itertuples()])],figures=['03_forecast_accuracy','04_daily_forecasts'])
        self.add('5. Core numerical replenishment: cost, service and safety',[],tables=[('Core numerical means; counts summed across 30 seeds',['Scenario','Policy','Cost index','Fill','Stockout','Violations','Held days'],[[SCENARIO_LABEL[r.scenario],r.policy,number(r.cost),percent(r.fill_rate),percent(r.stockout_rate),str(int(r.hard_violations)),str(int(r.held_decisions))] for r in self.numeric_means.itertuples()])],figures=['05_numerical_cost_service','06_numerical_safety'])
        self.add('6. Live LLM factorial: contribution and safeguards',[],tables=[('Agent factorial means; two seeds per row',['Arm','Scenario','Cost index','Fill','Held days','Violations','Freshcalls'],[[ARM_LABEL[r.arm],SCENARIO_LABEL[r.scenario],number(r.cost),percent(r.fill_rate),str(int(r.held_decisions)),str(int(r.hard_violations)),str(int(r.llm_calls))] for r in self.agent_means.itertuples()])],figures=['07_agent_factorial']+(['11_local_model_calls'] if any(x['key']=='11_local_model_calls' for x in self.figures) else []))
        self.add('7. Independent age-aware policy and stockout-gate sensitivity',[],tables=[('Independent sensitivity means; full grid in CSV',['Calculator','History','Life','Demand×','Cost index','Fill','Expiry share'],[[str(getattr(r,self.sens_columns['calculator'])),str(getattr(r,self.sens_columns['view'])),str(getattr(r,self.sens_columns['shelf'])),str(getattr(r,self.sens_columns['multiplier'])),number(r.cost,4),percent(r.fill_rate),percent(getattr(r,'waste_fraction',np.nan))] for r in self.sensitivity_means.itertuples()])],figures=[x['key'] for x in self.figures if x['key'] in ['09_service_waste_frontier','10_sensitivity_costs']])
        self.add('8. M5 versus FreshRetailNet: a short descriptive comparison',[],tables=[('Within-dataset inventory-repair contrasts',['Dataset','Scenario','Days','Cost change','Fill change(pp)','Identical parser/LLM pairs'],[[r.dataset,SCENARIO_LABEL[r.scenario],str(r.days),number(r.repair_cost_change_pct)+'%',number(r.repair_fill_change_pp),f'{r.cost_fill_identical_pairs}/{r.parser_llm_pairs}'] for r in self.comparison.itertuples()])],figures=['08_m5_fresh_comparison'])
        self.add('9. What the results mean for a business',[])
        self.add('10. Audit, limitations and reproducibility',[],tables=[('Independent audit receipt (selected scalar fields)',['Audit item','Recorded value'],scalar_frame(self.audit)[:100])])
        self.add('Appendix A. Artifact guide and exact run-level results',[],tables=[('Selected computational panel',['Series','Stratum','Train mean sales','Train stockout share'],[[r.series_id,str(r.censor_stratum),number(r.mean_sales,5),percent(r.censor_rate,2)] for r in self.panel.itertuples()])])
        # Save every core run in the standalone report, with selected nonredundant metrics.
        for sc in SCENARIOS[:3]:
            g=self.numeric[self.numeric.scenario.eq(sc)].sort_values(['policy','seed'])
            self.add('Appendix B. Core individual runs — '+SCENARIO_LABEL[sc],tables=[('All 30 seeds per policy',['Policy','Seed','Cost','Fill','Stockout','Held','Violations'],[[r.policy,str(r.seed),number(r.cost,3),percent(r.fill_rate,2),percent(r.stockout_rate,2),str(int(r.held_decisions)),str(int(r.hard_violations))] for r in g.itertuples()])])
        self.add('Appendix C. Live-agent individual runs',tables=[('All 32 live factorial runs',['Arm','Scenario','Seed','Cost','Fill','Held','Violations'],[[ARM_LABEL[r.arm],SCENARIO_LABEL[r.scenario],str(r.seed),number(r.cost,3),percent(r.fill_rate,2),str(int(r.held_decisions)),str(int(r.hard_violations))] for r in self.agent.sort_values(['arm','scenario','seed']).itertuples()])])

    def markdown(self):
        parts=['# FreshRetailNet bounded replenishment study\n',f'Built from saved evidence on {datetime.now(timezone.utc).isoformat()}.\n']
        figs={f['key']:f for f in self.figures}
        for section in all_report_sections(self):
            parts+=['\n## '+section['title']+'\n']
            parts += ['\n'+prose_text(p)+'\n' for p in section['paragraphs']]
            for title,cols,rows in section['tables']:
                parts+=['\n'+prose_text(title)+'\n','| '+' | '.join(cols)+' |','| '+' | '.join(['---']*len(cols))+' |']
                parts+=['| '+' | '.join(str(x).replace('|','/') for x in row)+' |' for row in rows]
                parts+=['']
            for key in section['figures']:
                f=figs[key];parts+=['\n!['+f['title']+'](figures/'+Path(f['path']).name+')\n',prose_text(f['caption'])+'\n']
        p=self.out/'FreshRetailNet_study_report.md';p.write_text('\n'.join(parts));return p

    def pdf(self):
        for face,file in [('ReportSans','DejaVuSans.ttf'),('ReportSansBold','DejaVuSans-Bold.ttf'),('ReportSansItalic','DejaVuSans-Oblique.ttf'),('ReportSansBoldItalic','DejaVuSans-BoldOblique.ttf')]:
            pdfmetrics.registerFont(TTFont(face,'/usr/share/fonts/truetype/dejavu/'+file))
        pdfmetrics.registerFontFamily('ReportSans',normal='ReportSans',bold='ReportSansBold',italic='ReportSansItalic',boldItalic='ReportSansBoldItalic')
        styles=getSampleStyleSheet()
        for style in styles.byName.values():style.fontName='ReportSansBold' if 'Heading' in style.name or style.name=='Title' else 'ReportSans'
        styles.add(ParagraphStyle(name='ReportBody',parent=styles['BodyText'],fontName='ReportSans',fontSize=10,leading=14,spaceAfter=8));styles.add(ParagraphStyle(name='ReportSmall',parent=styles['BodyText'],fontSize=8,leading=10));styles.add(ParagraphStyle(name='ReportCaption',parent=styles['BodyText'],fontSize=8.5,leading=11,spaceAfter=12,textColor=colors.HexColor('#44565D')))
        styles['Heading1'].textColor=colors.HexColor('#173D43');styles['Heading1'].fontSize=16;styles['Heading1'].leading=20
        styles['Heading2'].textColor=colors.HexColor('#173D43');styles['Heading2'].fontSize=12;styles['Heading2'].leading=15
        story=[Paragraph('FreshRetailNet-50K',styles['Title']),Paragraph('A bounded study of evidence-gated replenishment',styles['Heading1']),Spacer(1,12),Paragraph('Detailed results, scientific figures, audit evidence and a descriptive comparison with M5',styles['ReportBody']),Spacer(1,12),Paragraph('Business-school research report · 8 October 2026',styles['ReportBody']),Paragraph('30 selected series · 360 numerical runs · 32 live-agent runs · 1,140 policy runs',styles['ReportBody']),Paragraph('Source population validation: 50,000 series. Counterfactual cost and waste remain simulated; naturally lost demand has no released ground truth.',styles['ReportCaption']),PageBreak()]
        if self.followup:
            story.insert(-1,Paragraph(f"Separate exploratory transfer: {len(self.followup['runs']):,} additional runs; training-only cap calibration after the primary failure",styles['ReportBody']))
        figs={f['key']:f for f in self.figures}
        width=6.85*inch
        for idx,s in enumerate(all_report_sections(self)):
            if idx:story.append(PageBreak())
            story.append(Paragraph(s['title'],styles['Heading1']))
            for p in s['paragraphs']:story.append(Paragraph(escape(prose_text(p)),styles['ReportBody']))
            for title,cols,rows in s['tables']:
                story.append(Paragraph(escape(prose_text(title)),styles['Heading2']))
                data=[[Paragraph(escape(str(x)),styles['ReportSmall']) for x in cols]]+[[Paragraph(escape(str(x)),styles['ReportSmall']) for x in r] for r in rows]
                widths=[width/len(cols)]*len(cols)
                if len(cols)==2:widths=[width*.55,width*.45]
                t=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCEAE8')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#637B80')),('LINEBELOW',(0,1),(-1,-1),.2,colors.HexColor('#D1DADC')),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]));story+=[t,Spacer(1,10)]
            for key in s['figures']:
                f=figs[key];im=Image(f['path']);ratio=im.imageHeight/im.imageWidth;im.drawWidth=width;im.drawHeight=width*ratio
                if im.drawHeight>6.8*inch:im.drawHeight=6.8*inch;im.drawWidth=im.drawHeight/ratio
                story.append(KeepTogether([im,Paragraph(escape(prose_text(f['caption'])),styles['ReportCaption'])]))
        path=self.out/'FreshRetailNet_study_report.pdf'
        def footer(c,doc):
            c.saveState();c.setFont('ReportSans',8);c.setFillColor(colors.HexColor('#63757C'));c.drawString(40,26,'Evidence-gated replenishment | FreshRetailNet and M5 | Saved-evidence report');c.drawRightString(570,26,str(doc.page));c.restoreState()
        doc=SimpleDocTemplate(str(path),pagesize=(612,792),rightMargin=40,leftMargin=40,topMargin=40,bottomMargin=40,title='FreshRetailNet bounded replenishment study',author='MasterDissertation experimental artifact');doc.build(story,onFirstPage=footer,onLaterPages=footer);return path

    def dissertation(self):
        doc=Document(self.source);orig=list(doc.paragraphs);styles={s.name:s for s in doc.styles};styles.update({s.style_id:s for s in doc.styles})
        for name in ['Normal','Body Text']:
            if name in styles:styles[name].font.name='Times New Roman';styles[name].font.size=Pt(11);styles[name].paragraph_format.line_spacing=1.25;styles[name].paragraph_format.space_after=Pt(6)
        for name in ['Heading 1','Heading 2','Heading 3']:
            if name in styles:
                styles[name].font.name='Times New Roman';styles[name].font.color.rgb=RGBColor.from_string('173D43');styles[name].paragraph_format.keep_with_next=True
                if name=='Heading 1':styles[name].paragraph_format.page_break_before=True
        for section in doc.sections:
            section.page_width=Cm(21);section.page_height=Cm(29.7);section.left_margin=section.right_margin=Cm(2.3);section.top_margin=section.bottom_margin=Cm(2.2)
        def replace(index,text):orig[index].text=prose_text(text)
        def para(text,style='Body Text'):
            p=doc.add_paragraph(prose_text(text),style=styles.get(style,styles['Normal']));return p._p
        def heading(text,level=2):return para(text,'Heading '+str(level))
        def table(title,cols,rows):
            caption=para(title);caption.get_or_add_pPr().append(OxmlElement('w:keepNext'));blocks=[caption];t=doc.add_table(rows=1,cols=len(cols));
            for c,x in zip(t.rows[0].cells,cols):c.text=str(x)
            for row in rows:
                for c,x in zip(t.add_row().cells,row):c.text=str(x)
            for i,row in enumerate(t.rows):
                no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
                        for run in p.runs:run.font.size=Pt(8.5);run.font.name='Times New Roman';run.bold=i==0
                    if i==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E4ECEB');cell._tc.get_or_add_tcPr().append(sh)
            rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep);blocks.append(t._tbl);return blocks
        figs={f['key']:f for f in self.figures};figure_index=0;embedded=[]
        def figure(key):
            nonlocal figure_index
            f=figs[key];figure_index+=1;p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(f['path'],width=Inches(6.05));embedded.append(f)
            return [p._p,para(f'Figure 6.{figure_index}. '+f['title']+'. '+f['caption'])]
        def insert_before(index,blocks):
            anchor=orig[index]._p
            for block in blocks:anchor.addprevious(block)
        n=self.findings;before=self.agent_means.set_index(['arm','scenario']).loc[('llm_no_recovery','derived_field_collapse')];after=self.agent_means.set_index(['arm','scenario']).loc[('llm_recovery','derived_field_collapse')]
        for index,text in clean_doc_replacements(self).items():replace(index,text)
        if self.followup:
            extension=self.findings['exploratory_followup']
            explanation=f" A separate {extension['runs']:,}-run follow-up was motivated by the inherited spending cap's failure to pass clean orders. A registered training-only calibration selected a cap of {extension['training_selected_cap']:,.0f}; the evaluation reuses the holdout and is explicitly exploratory. Its {extension['seeds']} simulation seeds per cell do not turn that post-primary question into an untouched confirmation."
            orig[26].add_run(explanation)
            orig[381].add_run(explanation)
            original=self.agent[self.agent.arm.eq('llm_no_recovery')&self.agent.scenario.eq('normal')]
            transferred=self.followup['runs']
            transferred=transferred[transferred.arm.eq('llm_no_recovery')&transferred.scenario.eq('normal')&transferred.seed.isin(original.seed)]
            conclusion=f"The primary transfer result also changes the managerial answer. The inherited 1,500-point autonomous spending cap held every normal-operation decision, even when source evidence was clean. On the same original seeds, historical calibration to 1,850 points changed mean fill from {percent(original.fill_rate.mean())} to {percent(transferred.fill_rate.mean())} and held days from {int(original.held_decisions.sum())} to {int(transferred.held_decisions.sum())}. The exploratory follow-up found identical cost and fill in {extension['identical_cost_fill_pairs']}/{extension['parser_llm_pairs']} parser–LLM pairs. Useful autonomy therefore requires operating limits that fit the business scale, as well as trustworthy evidence. The gain is shared by the verified architecture; it is not an isolated language-model return."
            comparison=extension['calibrated_normal_vs_B3']
            conclusion+=f" The calibrated controller’s full 30-seed normal fill was {percent(comparison['calibrated_fill'])}, below the conventional constrained B3 baseline’s {percent(comparison['b3_fill'])}. Its lower indexed cost accompanies lower service, so calibration does not remove the remaining opportunity cost of the gate."
            insert_before(380,[para(conclusion)])

        implementation=doc.tables[13]
        implementation.add_column(Cm(4.2))
        statuses=['Reported status','Complete source profiles and pinned hashes','Complete; simulated fields explicitly labelled','Three point-estimate methods; natural truth unavailable','Four forecasts on 30 series; no full probabilistic suite','FIFO and separate calculators; core MILP is age-unaware','B1/B2/B3/B4 and matched B10; other architecture families untested','Descriptive paired effects; conditional seed variation','Not conducted; no participant claim']
        for row,text in zip(implementation.rows,statuses):row.cells[3].text=text
        implementation.autofit=False
        for column,width in zip(implementation.columns,[2.4,4.3,5.2,4.1]):
            column.width=Cm(width)
            for cell in column.cells:
                cell.width=Cm(width)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:run.font.size=Pt(8.5)
        for run in orig[456].runs:run.font.name='Liberation Mono';run.font.size=Pt(9)
        orig[456].paragraph_format.line_spacing=1.05
        if self.followup:orig[456].add_run('\nexploratory_gate_transfer: '+str(len(self.followup['runs']))+' runs / '+str(len(self.followup['protocol']['evaluation']['seeds']))+' seeds\ntransfer_status: post-primary hypothesis / reused holdout\ntraining_selected_spend_cap: '+str(self.followup['calibration']['selected_cap']))
        for run in orig[456].runs:run.font.name='Liberation Mono';run.font.size=Pt(9)
        # Preserve chapter order, inserting results underneath existing section anchors.
        def blocks_for(indices,include_paragraphs=False):
            blocks=[]
            for i in indices:
                s=self.sections[i]
                if include_paragraphs:blocks += [para(x) for x in s['paragraphs']]
                for title,cols,rows in s['tables']:blocks+=table('Results table. '+title,cols,rows)
                for key in s['figures']:blocks+=figure(key)
            return blocks
        insert_before(20,[para('Appendix D FreshRetailNet Results and Artifact Register'),para('Appendix E Complete M5 Figure Register'),para('Appendix F Complete Historical M5 Results')])
        insert_before(27,[heading('Executive evidence summary',2)]+[para(x) for x in self.sections[0]['paragraphs']])
        m5v1=csv(ROOT/'study_results/v1/M5_six_hour_results.csv')
        m5numeric=m5v1[m5v1.comparison_group.eq('numerical_7_days')]
        v1means=m5numeric.groupby(['policy','scenario']).agg(cost=('cost','mean'),fill=('fill_rate','mean'),violations=('hard_violations','sum'),held=('held_decisions','sum')).reset_index()
        m5blocks=[para('An earlier M5 stage also completed 372 runs: 360 numerical controls over seven days and 30 seeds, eight two-day prose-feasibility runs, and four two-day structured controls. The stages are not pooled. The first 96 numerical runs used seeds 0–7; a later 264-run extension used seeds 8–29 after that batch completed. The pooled 30-seed numerical means are therefore exploratory historical descriptions. The initial live prose experiment held every decision and therefore did not demonstrate an operational LLM saving. Its evidence motivated the bounded semantic-head and verified-repair improvements subsequently evaluated in V2.'),para('In the initial feed-gap numerical controls, B3 committed 63 actual hard-constraint violations while B4 committed none and held 90 of 210 decision days. This supports a safety role for the evidence gate, with an explicit delay trade-off. These are original M5 results; the FreshRetailNet portability failure described next shows why a threshold should not be transferred without checking the new quantity scale.')]
        m5blocks+=table('Historical M5 numerical means: seven days, 30 seeds per cell',['Policy','Scenario','Simulated USD','Fill','Violations','Held days'],[[r.policy,SCENARIO_LABEL[r.scenario],number(r.cost),percent(r.fill),str(int(r.violations)),str(int(r.held))] for r in v1means.itertuples()])
        insert_before(333,m5blocks)
        completion_rows=[
            ['F1 recovery','Artificial masking; zero/profile/LightGBM point estimates','Natural lost-demand truth, TimesNet/Tobit and calibrated intervals'],
            ['F2 forecasting','Four candidates, four validation folds, seven-day holdout','Full-population modeling and probabilistic TFT/DeepAR comparisons'],
            ['F3 perishable decisions','Independent continuous calculator grid; explicit conservation','An age-aware version of the inherited core MILP'],
            ['F4 valid availability','60 targeted gate-ablation runs','Human interpretation and production exception costs'],
            ['F5 LLM attribution','Shared-tool parser/LLM record-repair factorial','All semantic difficulty levels and raw/recovered-forecast factorial variants'],
            ['F6 sensitivity','Shelf lives 1/3/7, common demand multipliers 1/1.25/1.5','A complete economics, cadence and disruption sensitivity study'],
            ['Human audit','No participants; machine audit of records','Measured review speed, trust and staff effort']]
        insert_before(302,table('Executed coverage of the FreshRetailNet questions',['Question','Completed evidence','Remaining gap'],completion_rows))
        insert_before(336,blocks_for([1,2],False))
        insert_before(336,[para(self.sections[6]['paragraphs'][1])])
        insert_before(339,blocks_for([3,4,5,7],False))
        insert_before(342,blocks_for([6],False))
        insert_before(345,[para(x) for x in self.sections[6]['paragraphs'] if x.startswith('A concrete simulator case')])
        if self.followup_sections:
            extension=[]
            for section in self.followup_sections:
                extension.append(heading('6.5.1 Exploratory transfer of the spending threshold',3))
                extension += [para(x) for x in section['paragraphs']]
                for title,cols,rows in section['tables']:extension+=table(title,cols,rows)
                for key in section['figures']:extension+=figure(key)
            insert_before(345,extension)
        insert_before(347,[para(x) for x in self.sections[6]['paragraphs'][5:] if not x.startswith('A concrete simulator case')]+blocks_for([10],False))
        insert_before(353,blocks_for([8],False)+[para(x) for x in self.sections[9]['paragraphs']]+blocks_for([9],False))
        # Complete individual core rows and all sensitivity rows remain portable CSVs.
        doc.add_page_break();doc.add_paragraph('Appendix D FreshRetailNet Results and Artifact Register',style=styles.get('Heading 1',styles['Normal']))
        for i in range(11,len(self.sections)):
            s=self.sections[i];doc.add_paragraph(s['title'].replace('Appendix A.','D.1').replace('Appendix B.','D.2').replace('Appendix C.','D.3'),style=styles.get('Heading 2',styles['Normal']))
            for x in s['paragraphs']:para(x)
            for title,cols,rows in s['tables']:table(title,cols,rows)
        doc.add_paragraph('Appendix E Complete M5 Figure Register',style=styles.get('Heading 1',styles['Normal']))
        para('The previous M5 studies remain part of the comparative record. These original scientific figures are preserved from their final reports and describe earlier M5 windows. They are not FreshRetailNet results. The repository retains both original M5 PDFs, all 404 individual runs and the complete audit archives.')
        m5_manifest=json_read(ROOT/'study_results/dissertation/dissertation_build_manifest.json')
        m5_records=m5_manifest['figures']
        assert len(m5_records)==25
        m5_figures=[ROOT/r['path'] for r in m5_records]
        assert all(sha(p)==r['sha256'] for p,r in zip(m5_figures,m5_records))
        for i,p in enumerate(m5_figures,1):
            paragraph=doc.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER;paragraph.paragraph_format.keep_with_next=True;paragraph.add_run().add_picture(str(p),width=Inches(6.05));para(f'Figure E.{i}. '+m5_records[i-1]['caption']);embedded.append({'path':str(p),'sha256':sha(p),'key':'m5_'+p.stem})
        landscape=doc.add_section(WD_SECTION_START.NEW_PAGE)
        landscape.orientation=WD_ORIENT.LANDSCAPE;landscape.page_width=Cm(29.7);landscape.page_height=Cm(21);landscape.left_margin=landscape.right_margin=Cm(1.8)
        doc.add_paragraph('Appendix F Complete Historical M5 Results',style=styles.get('Heading 1',styles['Normal']))
        para('These tables retain every completed M5 run as historical evidence: 360 seven-day numerical controls, eight two-day prose feasibility runs, four two-day structured controls and 32 fourteen-day V2 runs. The stages are reported separately because their horizons and interventions differ. Simulated USD is the original M5 convention; it is not a currency conversion for FreshRetailNet.')
        historical=Document(ROOT/'study_results/dissertation/Evidence_Gated_Autonomy_Dissertation_Final.docx')
        historical_rows=0
        for t in historical.tables:
            preceding=t._tbl.getprevious()
            text=''.join(preceding.itertext()) if preceding is not None else ''
            if text.startswith('Table B.'):
                cap=next((p.text for p in historical.paragraphs if p._p is preceding),None)
                if cap is None:cap=''.join(preceding.itertext())
                para(cap.replace('Table B.','Table F.'))
                doc.element.body.insert(len(doc.element.body)-1,deepcopy(t._tbl))
                if not text.startswith('Table B.19.'):historical_rows+=len(t.rows)-1
        assert historical_rows==404, f'M5 individual run rows missing: {historical_rows}'
        self.historical_rows=historical_rows
        for section in doc.sections:
            footer=section.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('Page ')
            for instr in ['PAGE','NUMPAGES']:
                field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),instr);footer._p.append(field)
                if instr=='PAGE':footer.add_run(' of ')
        doc.core_properties.title='Evidence-Gated Autonomy for Retail Replenishment: M5 and FreshRetailNet'
        doc.core_properties.subject='Completed bounded experimental evidence, business implications and full artifact register';doc.core_properties.modified=datetime.now(timezone.utc).replace(tzinfo=None)
        output=self.out/'Dual_Benchmark_Dissertation_Final.docx';doc.save(output)
        manifest={'source':str(self.source),'source_sha256':sha(self.source),'output_sha256':sha(output),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'inline_figures':len(doc.inline_shapes),'fresh_figures':len(self.figures),'m5_figures':len(m5_figures),'eight_original_chapters_preserved':all(any('CHAPTER'+str(k) in re.sub(r'\s+','',p.text) for p in doc.paragraphs) for k in range(1,9)),'embedded_figures':embedded,'claims':self.findings}
        (self.out/'dissertation_build_manifest.json').write_text(json.dumps(manifest,indent=2));return output

    def build(self):
        self.plots();self.make_narrative();polish(self);followup_content(self);md=self.markdown();pdf=self.pdf();docx=self.dissertation()
        outputs=[]
        for p in sorted(self.out.rglob('*')):
            if p.is_file() and p.name!='report_manifest.json' and p.suffix!='.log':outputs.append({'path':str(p.relative_to(self.out)),'bytes':p.stat().st_size,'sha256':sha(p)})
        manifest={'status':'BUILT_PENDING_RENDER_QA','built_utc':datetime.now(timezone.utc).isoformat(),'code_sha256':sha(__file__),'inputs':self.inputs,'outputs':outputs,'findings':self.findings,'figure_count':len(self.figures),'individual_runs':{'numerical':len(self.numeric),'live_agents':len(self.agent),'independent_policy':len(self.sensitivity),'exploratory_gate_transfer':len(self.followup['runs']) if self.followup else 0},'source_document_preserved':sha(self.source)==self.inputs[str(self.source)]['sha256'] if str(self.source) in self.inputs else True}
        (self.out/'report_manifest.json').write_text(json.dumps(manifest,indent=2));print(json.dumps({'pdf':str(pdf),'docx':str(docx),'markdown':str(md),'figures':len(self.figures),'status':manifest['status']},indent=2))


def load_followup(self,allow_partial=False):
    self.followup=None;self.followup_sections=[]
    root=self.root/'gate_transfer'
    if not (root/'protocol.json').exists():return
    protocol=json_read(root/'protocol.json')
    completion=json_read(root/'execution_complete.json',{})
    if completion.get('status')!='COMPLETE':
        if allow_partial:return
        raise RuntimeError('Exploratory gate transfer is registered but incomplete; retain its partial evidence before reporting final completion')
    frame=csv(require(root/'summary.csv'));expected=protocol['evaluation']['expected_runs']
    assert len(frame)==expected
    audit=json_read(require(root/'audit.json'));assert str(audit.get('status','')).startswith('VERIFIED')
    calibration=json_read(require(root/'calibration/selection_receipt.json'))
    independent=json_read(require(root/'audit_independent/audit.json'));assert independent.get('status')=='VERIFIED_COMPLETE', 'Follow-up independent audit is incomplete'
    assert independent['totals']['runs']==len(frame)
    assert independent['totals']['decisions']==int(frame.trace_count.sum())
    assert independent['totals']['parser_llm_matched_pairs']==len(self.pair_controls(frame))
    numeric_cases=json_read(require(root/'audit_independent/numeric_extraction_cases.json'))
    self.followup={'root':root,'protocol':protocol,'completion':completion,'runs':frame,'audit':audit,'calibration':calibration,'numeric_cases':numeric_cases,
        'independent_audit':independent,'pairs':self.pair_controls(frame),'means':frame.groupby(['arm','scenario']).agg(runs=('cost','size'),cost=('cost','mean'),fill_rate=('fill_rate','mean'),held_decisions=('held_decisions','sum'),hard_violations=('hard_violations','sum'),llm_calls=('llm_calls','sum'),tokens=('tokens','sum')).reset_index()}
    frame.to_csv(self.out/'FreshRetailNet_exploratory_gate_transfer_all_runs.csv',index=False)
    self.followup['pairs'].to_csv(self.out/'FreshRetailNet_exploratory_gate_transfer_parser_llm_pairs.csv',index=False)
    self.followup['means'].to_csv(self.out/'FreshRetailNet_exploratory_gate_transfer_means.csv',index=False)
    for path in [root/'protocol.json',root/'summary.csv',root/'audit.json',root/'calibration/selection_receipt.json',root/'audit_independent/audit.json',root/'audit_independent/numeric_extraction_cases.json']:
        self.inputs[str(path.relative_to(ROOT))]={'sha256':sha(path),'bytes':path.stat().st_size}
    additional=[]
    means=frame.groupby(['arm','scenario'])[['cost','fill_rate','held_decisions','hard_violations']].mean()
    for sc in SCENARIOS:
        before=means.loc[('llm_no_recovery',sc)];after=means.loc[('llm_recovery',sc)];matched=self.followup['pairs'][self.followup['pairs'].scenario.eq(sc)]
        additional.append({'dataset':'FreshRetailNet (calibrated gate)','scenario':sc,'series':30,'seeds':len(protocol['evaluation']['seeds']),'days':7,
            'repair_cost_change_pct':100*(after.cost/before.cost-1) if before.cost else np.nan,'repair_fill_change_pp':100*(after.fill_rate-before.fill_rate),
            'repair_holds_change_per_run':after.held_decisions-before.held_decisions,'parser_llm_pairs':len(matched),'cost_fill_identical_pairs':int(matched.cost_fill_identical.sum()),
            'mean_llm_minus_parser_cost_pct':float(np.mean(100*(matched.llm_cost/matched.parser_cost-1))),
            'comparison_type':'Exploratory follow-up after primary portability failure; reused holdout; no confirmatory dataset effect'})
    self.comparison.loc[self.comparison.dataset.eq('FreshRetailNet'),'dataset']='FreshRetailNet (fixed gate)'
    self.comparison=pd.concat([self.comparison,pd.DataFrame(additional)],ignore_index=True)
    self.comparison.to_csv(self.out/'M5_FreshRetailNet_descriptive_comparison.csv',index=False)


def followup_content(self):
    primary=self.agent[self.agent.arm.eq('parser_no_recovery')&self.agent.scenario.eq('normal')]
    normal_held=int(primary.held_decisions.sum());normal_days=int(primary.trace_count.sum())
    portability=f'The original gate did not transfer successfully. Its inherited autonomous spending cap was 1,500 cost-index points, while a clean proposed FreshRetailNet order could require about 1,656 points. In the primary normal parser condition, {normal_held} of {normal_days} decision days were held. This failure is reported before any repair: an architecture can be computationally feasible and still be commercially inactive if its threshold is carried into a new quantity scale without checking.'
    self.sections[0]['paragraphs'].insert(1,portability)
    self.sections[5]['paragraphs'].append(portability)
    self.sections[6]['paragraphs'].insert(1,portability)
    if self.followup is None:return
    f=self.followup;protocol=f['protocol'];cal=f['calibration'];frame=f['runs'];means=f['means'];pairs=f['pairs'];seeds=len(protocol['evaluation']['seeds']);cap=float(cal['selected_cap'])
    fig,axes=plt.subplots(1,2,figsize=(10,4))
    spends=cal['proposed_spends'];axes[0].bar(range(7),spends,color=PALETTE[0]);axes[0].axhline(1500,color=PALETTE[3],label='Inherited cap 1,500');axes[0].axhline(cap,color=PALETTE[1],label=f'Training-selected cap {cap:,.0f}');axes[0].set_xticks(range(7),range(1,8));axes[0].set_xlabel('Historical proposal number');axes[0].set_ylabel('Historical proposed spending index');axes[0].legend(fontsize=10)
    before=self.agent[self.agent.arm.eq('parser_no_recovery')&self.agent.scenario.eq('normal')]
    after=frame[frame.arm.eq('parser_no_recovery')&frame.scenario.eq('normal')&frame.seed.isin(before.seed)]
    assert set(before.seed)==set(after.seed) and int(before.trace_count.sum())==int(after.trace_count.sum())
    axes[1].bar(['Primary fixed cap','Exploratory calibrated cap'],[before.held_decisions.sum()/before.trace_count.sum(),after.held_decisions.sum()/after.trace_count.sum()],color=[PALETTE[3],PALETTE[1]]);axes[1].yaxis.set_major_formatter(PercentFormatter(1));axes[1].set_ylabel('Share of held normal decision days');axes[1].set_xlabel('Same parser condition, paired seeds 13 and 29');axes[1].set_ylim(0,1.1)
    self.savefig('12_gate_portability','A spending threshold needs calibration to its new scale',f'The right panel compares the same original seeds, 13 and 29, with {int(before.trace_count.sum())} normal parser decision days under each cap. Full 30-seed means appear in the following figure. The fixed-gate failure motivated this exploratory follow-up after primary evaluation began. Seven historical calibration proposals use training-only dates and a separately fitted model. The cap formula was registered before calibration. Reusing the official evaluation window means the follow-up is not confirmatory.',fig)
    fig,axes=plt.subplots(2,2,figsize=(9.5,7.5));indexed=means.set_index(['arm','scenario'])
    for ax,sc in zip(axes.flat,SCENARIOS):
        ax.bar(range(4),[indexed.loc[(a,sc),'cost'] for a in ARMS],color=PALETTE[:4]);ax.set_xticks(range(4),[ARM_LABEL[a].replace(' + ','\n+ ') for a in ARMS],fontsize=12);ax.set_title(SCENARIO_LABEL[sc]);ax.set_ylabel('Seven-day simulated cost index')
        for i,a in enumerate(ARMS):ax.text(i,indexed.loc[(a,sc),'cost'],f"Fill {percent(indexed.loc[(a,sc),'fill_rate'])}",ha='center',va='bottom',fontsize=11)
        ax.set_ylim(0,float(means[means.scenario.eq(sc)].cost.max())*1.25)
    self.savefig('13_calibrated_agent_factorial','Exploratory calibrated transfer: useful execution and fair attribution',f'{seeds} paired simulator seeds per configuration and scenario, all 30 series and seven days. The original saved forecast, prompts, model and numerical tools are retained. Only the spending threshold and its version change semantically. This follow-up was chosen after inspecting primary portability and reuses the holdout.',fig)
    repair=[]
    for sc in SCENARIOS:
        a=frame[frame.arm.eq('llm_no_recovery')&frame.scenario.eq(sc)].set_index('seed')
        b=frame[frame.arm.eq('llm_recovery')&frame.scenario.eq(sc)].set_index('seed')
        dif=b.cost-a.cost
        for seed,value in dif.items():repair.append({'scenario':sc,'seed':seed,'repair_minus_no_repair_cost':value,'repair_minus_no_repair_fill_pp':100*(b.loc[seed,'fill_rate']-a.loc[seed,'fill_rate'])})
    repair=pd.DataFrame(repair);repair.to_csv(self.out/'FreshRetailNet_exploratory_repair_paired_seed_effects.csv',index=False)
    fig,ax=plt.subplots(figsize=(9,4))
    for j,sc in enumerate(SCENARIOS):
        values=repair[repair.scenario.eq(sc)].repair_minus_no_repair_cost.to_numpy();offset=np.linspace(-.13,.13,len(values));ax.scatter(j+offset,values,color=PALETTE[j],alpha=.7,s=18);ax.plot([j-.18,j+.18],[values.mean()]*2,color='#253A41',lw=2)
    ax.set_xticks(range(4),[SCENARIO_LABEL[sc] for sc in SCENARIOS],rotation=15);ax.axhline(0,color='#888',lw=.8);ax.set_ylabel('Repair-on minus repair-off cost index');ax.set_xlabel('Paired conditional simulator seed effects')
    self.savefig('14_calibrated_paired_effects','Record-repair benefit should survive individual paired seeds','Each point is a matched simulator seed in the exploratory follow-up, with a short line for its mean. Seeds vary a shared assumed environment, not independent retailers or historical forecast origins. The graph makes adverse and null effects visible without a confirmatory population claim.',fig)
    n_pairs=len(pairs);n_equal=int(pairs.cost_fill_identical.sum());collapse=indexed.loc[('llm_no_recovery','derived_field_collapse')];repaired=indexed.loc[('llm_recovery','derived_field_collapse')]
    cost_delta=100*(repaired.cost/collapse.cost-1);fill_delta=100*(repaired.fill_rate-collapse.fill_rate)
    self.findings['exploratory_followup']={'runs':len(frame),'decisions':int(frame.trace_count.sum()),'seeds':seeds,'training_selected_cap':cap,'identical_cost_fill_pairs':n_equal,'parser_llm_pairs':n_pairs,'repair_collapse_cost_change_pct':float(cost_delta),'repair_collapse_fill_change_pp':float(fill_delta)}
    paragraphs=[
        portability,
        'The failure led to a separately registered follow-up. Its hypothesis was chosen after the primary study had started; the official holdout had already been used. The study therefore labels all subsequent transfer results exploratory. The original fixed-gate outputs are unchanged and remain in the report and archive.',
        f'Cap calibration uses seven earlier historical dates, training a separate GRU through day 75 and simulating June 12–18 (indices 76–82). The rule was fixed before those proposals were generated: min(0.8 × budget, ceil(1.10 × the seven-proposal 95th percentile / 50) × 50). The historical 95th percentile was {cal["p95"]:,.3f}; the selected cap was {cap:,.0f}, below the 2,400 maximum. The independent two-person threshold remains 2,500. All seven proposals were included, including {cal.get("held_decisions",0)} days held by the unchanged extra-spend risk rule. Proposal values, dates, model identity and selection timestamps are saved.',
        f'The transfer experiment contains {len(frame):,} runs and {int(frame.trace_count.sum()):,} decision days: four configurations, four scenarios and {seeds} seeds. It uses the primary saved GRU tensors and the same simulator, supplier sources, LLM, prompts, constraints and checking tools. Only the spending cap and its version change semantically; output locations, cache and spend ledger are separately named for delivery and audit.',
        f'Cost and fill were identical in {n_equal}/{n_pairs} parser–LLM pairs. Under the inventory-field fault, verified repair changed mean LLM cost by {cost_delta:+.2f}% and fill by {fill_delta:+.2f} percentage points. These results distinguish fixing an unusable operating threshold from the value of the language model. They are not an untouched confirmation, and any shared repair benefit remains attributable to the architecture.',
        'For a business, this is an implementation lesson as much as a modeling result. A limit calibrated for unit sales in one retailer can produce systematic holds after quantity normalization changes. Conservative caps remain valuable, but the business must verify that legitimate orders can pass them. That review should use historical proposals and explicit risk budgets, rather than repeatedly adjusting limits until a favourable holdout result appears.'
    ]
    collapsed_pairs=repair[repair.scenario.eq('derived_field_collapse')]
    improved_cost=int(collapsed_pairs.repair_minus_no_repair_cost.lt(0).sum());worse_cost=int(collapsed_pairs.repair_minus_no_repair_cost.gt(0).sum());improved_fill=int(collapsed_pairs.repair_minus_no_repair_fill_pp.gt(0).sum())
    paragraphs.append(f"The paired seed results show the benefit and its limits. Under the inventory-field fault, indexed cost falls in {improved_cost}/{len(collapsed_pairs)} seeds and rises in {worse_cost}; fill improves in {improved_fill}/{len(collapsed_pairs)}. Verified repair restores the normal-case economic and service results for every seed. The stale-feed condition is unchanged because fresh source evidence is unavailable. These are conditional simulator effects; the same recovery benefit occurs with the capable parser.")
    tables=[('Exploratory transfer means; primary results remain separate',['Arm','Scenario','Runs','Cost index','Fill','Held days','Violations'],[[ARM_LABEL[r.arm],SCENARIO_LABEL[r.scenario],str(r.runs),number(r.cost),percent(r.fill_rate),str(int(r.held_decisions)),str(int(r.hard_violations))] for r in means.itertuples()]),
        ('Training-only proposed spending; all seven days retained',['Historical decision','Spending index'],[[str(i+76),number(x,3)] for i,x in enumerate(spends)])]
    b3=self.numeric[self.numeric.policy.eq('B3')&self.numeric.scenario.eq('normal')].set_index('seed')
    gated=frame[frame.arm.eq('llm_no_recovery')&frame.scenario.eq('normal')].set_index('seed')
    assert set(b3.index)==set(gated.index)
    remaining=pd.DataFrame({'seed':sorted(b3.index)})
    for name,values in [('b3_cost',b3.cost),('calibrated_cost',gated.cost),('b3_fill',b3.fill_rate),('calibrated_fill',gated.fill_rate),('b3_held_days',b3.held_decisions),('calibrated_held_days',gated.held_decisions)]:remaining[name]=remaining.seed.map(values)
    remaining['calibrated_minus_b3_cost']=remaining.calibrated_cost-remaining.b3_cost
    remaining['calibrated_minus_b3_fill_pp']=100*(remaining.calibrated_fill-remaining.b3_fill)
    remaining.to_csv(self.out/'FreshRetailNet_calibrated_vs_B3_normal_seed_pairs.csv',index=False)
    remaining_cost_change=100*(gated.cost.mean()/b3.cost.mean()-1);remaining_fill_change=100*(gated.fill_rate.mean()-b3.fill_rate.mean())
    self.findings['exploratory_followup']['calibrated_normal_vs_B3']={'seeds':len(b3),'b3_cost':float(b3.cost.mean()),'calibrated_cost':float(gated.cost.mean()),'cost_change_pct':float(remaining_cost_change),'b3_fill':float(b3.fill_rate.mean()),'calibrated_fill':float(gated.fill_rate.mean()),'fill_change_pp':float(remaining_fill_change),'b3_held_days':int(b3.held_decisions.sum()),'calibrated_held_days':int(gated.held_decisions.sum())}
    tables.append(('Remaining clean-operation gate trade-off; 30 paired seeds, exploratory',['Controller','Runs','Cost index','Fill','Held days','True violations'],[['Conventional B3 MILP',str(len(b3)),number(b3.cost.mean()),percent(b3.fill_rate.mean(),2),str(int(b3.held_decisions.sum())),str(int(b3.hard_violations.sum()))],['Calibrated LLM with the shared gate',str(len(gated)),number(gated.cost.mean()),percent(gated.fill_rate.mean(),2),str(int(gated.held_decisions.sum())),str(int(gated.hard_violations.sum()))]]))
    paragraphs.append(f"Calibration restores useful execution, but the remaining gate still carries an opportunity cost. Over the same 30 seeds, normal-operation fill is {percent(gated.fill_rate.mean(),2)} for the calibrated controller and {percent(b3.fill_rate.mean(),2)} for conventional B3, a {remaining_fill_change:+.2f}-point change. Mean indexed cost is {gated.cost.mean():,.2f} against {b3.cost.mean():,.2f} ({remaining_cost_change:+.2f}%), with {int(gated.held_decisions.sum())} held days against {int(b3.held_decisions.sum())}. Both use the same saved forecaster, simulator and numerical tools and both have zero actual constraint violations. Lower indexed spending accompanies lower service; these values do not establish overall LLM superiority. The parser gives the same calibrated result, and this comparison remains exploratory.")
    self.sections[0]['paragraphs'].append(f"The calibrated controller still fills {percent(gated.fill_rate.mean(),2)} of normal simulated demand, compared with {percent(b3.fill_rate.mean(),2)} for conventional B3 over the same 30 seeds. Its lower indexed cost comes with lower service and remaining holds. Calibration improves an unusable gate; it does not demonstrate a dominant LLM policy.")
    independent=f['independent_audit'];totals=independent['totals']
    flags=frame.groupby('arm')[['hard_violations','reference_deviations','executed_actions','held_decisions']].sum()
    tables.append(('Exploratory actual violations versus clean-reference distance',['Arm','True rule checks violated','Reference-only flagged days','Nonzero action days','Held days'],[[ARM_LABEL[arm],str(int(row.hard_violations)),str(int(row.reference_deviations)),str(int(row.executed_actions)),str(int(row.held_decisions))] for arm,row in flags.iterrows()]))
    paragraphs.append('Reference-only deviations are retained separately from actual active-rule violations in this follow-up as well. The legacy harmful_execution composite includes either endpoint. Its label alone cannot identify realized injury, waste or economic harm; outcomes and the underlying reason for each flag must be inspected.')
    audit_fields=['runs','decisions','physical_API_attempts','tokens','client_or_schema_error_events','committed_true_constraint_violations','held_decisions','selector_instruction_compliant','selector_semantic_error_responses','parser_llm_matched_pairs','parser_llm_pairs_primary_metrics_equal','parser_llm_pairs_all_action_days_equal','physical_mass_balance_cases_passed','forbidden_model_context_findings']
    tables.append(('Independent audit of the exploratory transfer',['Measure','Recorded count'],[(AUDIT_LABEL.get(key,key),totals[key]) for key in audit_fields if key in totals]))
    paragraphs.append(f"The independent follow-up audit verifies {totals['runs']:,} runs and {totals['decisions']:,} physical decision records. It captures {totals.get('physical_API_attempts',0):,} fresh attempts and {totals.get('tokens',0):,} tokens, with {totals.get('client_or_schema_error_events',0)} client/schema errors. There are {totals.get('committed_true_constraint_violations',0)} actual committed requested-plan violations. Selector compliance remains a separate table entry, so a blocked semantic error is not counted as a successful choice.")
    exact_terms=sum(case['exact_numeric_terms'] for case in f['numeric_cases']);submitted_terms=sum(case['submitted_terms'] for case in f['numeric_cases'])
    paragraphs.append(f"The bounded numerical head returns {exact_terms}/{submitted_terms} exact submitted terms in {len(f['numeric_cases'])} fresh extraction responses. Separately, {totals['selector_instruction_compliant']}/{totals['physical_selector_responses']} fresh recovery selections follow the instruction. The {totals['selector_semantic_error_responses']} stale-feed mis-selections are refused by independent checks. A zero client/schema-error count therefore coexists with substantive reasoning errors; safe final execution depends on the wider control process.")
    grounding=pd.DataFrame(independent['grounding_aggregates'])
    grounding.to_csv(self.out/'FreshRetailNet_exploratory_source_grounding_audit.csv',index=False)
    columns=['arm','grounding_attempted_decisions','complete_source_set_exact_decisions','state_gated_unattempted_decisions','false_active_rules','omitted_active_rules']
    tables.append(('Exploratory source grounding; conditional coverage is explicit',['Configuration','Grounding attempted','Exact full source sets','State-gated days','False clauses','Omitted clauses'],[[ARM_LABEL.get(row[0],row[0]),*row[1:]] for row in grounding[columns].values.tolist()]))
    attempted=int(grounding.grounding_attempted_decisions.sum());exact=int(grounding.complete_source_set_exact_decisions.sum());unattempted=int(grounding.state_gated_unattempted_decisions.sum())
    paragraphs.append(f"Grounding was attempted on {attempted:,} of {totals['decisions']:,} follow-up decision days; {exact:,} of those attempted source sets were exact. Each available set has 163 clauses, yielding {int(grounding.exact_active_rules.sum()):,} exact active-rule records. The other {unattempted:,} days were state-gated before grounding and remain in the denominator of all decision days. These records combine a bounded model extraction head with deterministic parsing and verification; full-set agreement is not the model’s accuracy on every clause.")
    original_seeds=set(self.agent.seed.unique())
    matched=frame[frame.seed.isin(original_seeds)]
    transition=[]
    for (arm,scenario),g in self.agent.groupby(['arm','scenario']):
        a=matched[(matched.arm==arm)&(matched.scenario==scenario)]
        assert len(a)==len(g)
        transition.append({'arm':arm,'scenario':scenario,'matched_seeds':len(a),'primary_cost':g.cost.mean(),'calibrated_cost':a.cost.mean(),'cost_change_pct':100*(a.cost.mean()/g.cost.mean()-1),'primary_fill':g.fill_rate.mean(),'calibrated_fill':a.fill_rate.mean(),'fill_change_pp':100*(a.fill_rate.mean()-g.fill_rate.mean()),'primary_held_days':int(g.held_decisions.sum()),'calibrated_held_days':int(a.held_decisions.sum())})
    transition=pd.DataFrame(transition);transition.to_csv(self.out/'FreshRetailNet_gate_transition_original_seed_pairs.csv',index=False)
    tables.append(('Gate transfer on the same two original seeds; exploratory holdout reuse',['Arm','Scenario','Cost change','Fill change (pp)','Held days before/after'],[[ARM_LABEL[r.arm],SCENARIO_LABEL[r.scenario],f'{r.cost_change_pct:+.2f}%',f'{r.fill_change_pp:+.2f}',f'{r.primary_held_days}/{r.calibrated_held_days}'] for r in transition.itertuples()]))
    normal=transition[(transition.arm=='llm_no_recovery')&(transition.scenario=='normal')].iloc[0]
    paragraphs.append(f"On the original matched seeds 13 and 29 in normal operation, cap transfer changed mean LLM cost from {normal.primary_cost:,.2f} to {normal.calibrated_cost:,.2f} ({normal.cost_change_pct:+.2f}%) and fill from {percent(normal.primary_fill)} to {percent(normal.calibrated_fill)} ({normal.fill_change_pp:+.2f} points). Held decision days changed from {normal.primary_held_days} to {normal.calibrated_held_days}. These paired values isolate the cap change more clearly than comparing the primary two-seed mean with all 30 follow-up seeds, but the question remains exploratory.")
    self.sections[0]['paragraphs'].append(f"In that exploratory follow-up, the original normal-operation seed pair changed from {percent(normal.primary_fill)} to {percent(normal.calibrated_fill)} fill, with held days {normal.primary_held_days} to {normal.calibrated_held_days}. This is a verified cap-transfer effect, not an incremental LLM advantage: parser–LLM cost and fill are equal in {n_equal}/{n_pairs} follow-up pairs.")
    self.followup_sections=[{'title':'6a. Exploratory spending-cap portability follow-up','paragraphs':paragraphs,'tables':tables,'figures':['12_gate_portability','13_calibrated_agent_factorial','14_calibrated_paired_effects']}]
    self.sections[0]['paragraphs'].append(f'A separately registered follow-up used a training-only spending-cap formula and {seeds} simulation seeds per cell, producing {len(frame):,} additional runs. It is exploratory because the fixed-gate failure motivated it after primary observations, and it reuses the holdout. Its complete outputs appear separately; none replaces the primary failure.')


def all_report_sections(self):
    return self.sections[:7]+self.followup_sections+self.sections[7:]


def polish(self):
    """Use human-readable prose while keeping identifiers and all source units intact."""
    n=self.findings;p=self.population
    means=self.agent_means.set_index(['arm','scenario'])
    before=means.loc[('llm_no_recovery','derived_field_collapse')]
    after=means.loc[('llm_recovery','derived_field_collapse')]
    prose={
    0:[
        'The FreshRetailNet study is complete on a frozen panel of 30 store–product series. It adds 360 conventional replenishment runs, 32 matched parser and LLM runs, and 1,140 independent policy experiments to the existing M5 evidence. Each main simulation run uses all 30 series. The source contains 50,000 series, so this panel covers 0.06% of the release. It supports a bounded engineering conclusion rather than a population estimate.',
        f'The central attribution check found identical cost and fill in {n["parser_llm_identical_pairs"]} of {n["parser_llm_total_pairs"]} matched parser–LLM scenario and seed pairs. The language model was able to participate in a verified replenishment workflow. The tested supplier grammar did not, however, demonstrate an additional economic benefit from choosing an LLM over a capable parser.',
        f'Under an injected inventory-field failure, verified record repair changed mean LLM-arm cost from {number(before.cost)} to {number(after.cost)} ({n["repair_cost_change_pct"]:+.2f}%) and fill from {percent(before.fill_rate)} to {percent(after.fill_rate)} ({n["repair_fill_change_pp"]:+.2f} percentage points). The parser received the same repair tools. Any shared improvement belongs to the recovery architecture, rather than to the language model alone.',
        'Fresh retail adds a second problem: a trustworthy stockout record can still leave demand partly unobserved. This study evaluates recovery on artificially hidden sales, forecasts a protected seven-day holdout, and examines perishability inside disclosed simulators. It does not measure true lost demand during natural stockouts, real retailer profit, actual waste savings, or changes in human trust.'
    ],
    1:[
        f'The protocol was registered after the request at {self.protocol.get("requested_at_utc")}. The experiment cutoff was {self.protocol.get("experiment_cutoff_utc")}, leaving reporting and publication time before {self.protocol.get("report_and_push_deadline_utc")}. The experiment was frozen at {self.protocol.get("frozen_at_utc", "the timestamp recorded in protocol.json")}. Panel and model selection used training data only. This report consumes the saved outputs and does not restart an experiment.',
        'The evidence has four distinct meanings. Observed fields are released sales, identifiers and hourly stockout annotations. Derived fields include historical censoring summaries and source validation. Inferred fields are recovered demand estimates and forecasts. Simulated fields include stock balances, purchasing, suppliers, lead times, expiry and alternative-policy outcomes. A correctly recorded stockout is therefore different from a deliberately corrupted inventory record.',
        'For compatibility with inherited integer-order tools, the core simulator represents quantities as 100 times the published globally normalized sales amount. These are scaled model quantities, not kilograms, cartons or original physical units. Its legacy schemas contain USD tags, but all FreshRetailNet results are reported as a simulated cost index. The independent age-aware experiment uses continuous normalized quantities and its own dimensionless economics.'
    ],
    2:[
        f'The pinned dataset revision is {self.protocol["source_revision"]}. Source checks covered {p["rows"]:,} training rows, {p["series"]:,} store–product series, {p["products"]} products, {p["stores"]} stores and {p["cities"]} cities. The downloaded files contain 865 distinct product IDs, resolving the card and paper counting discrepancy for this revision. Daily sales agree with the sum of hourly sales within {p["maximum_daily_hourly_difference"]:.2e}. Duplicate keys, hierarchy conflicts, invalid stockout encodings and inconsistent stockout counts were absent.',
        f'The publisher’s implementation and the released arrays confirm that 1 means out of stock. Its 16 operating-hour slots run from 06:00 through 21:00. The training hourly stockout share is {percent(p["operating_hour_stockout_fraction"],3)}, while {percent(p["daily_any_stockout_fraction"],3)} of daily rows contain at least one stockout hour. These are different denominators. There are 16,073 training rows with discount values outside [0,1]; the sales records were retained and discount was excluded from forecasting features.',
        'Panel selection used only the first 62 training days. Eligible series needed 62 complete daily rows and mean normalized daily sales between 0.02 and 1.0. Of 37,097 eligible series, three equal-rank historical censoring strata were formed. Ten series were then chosen by deterministic SHA-256 ranking from each stratum. This balances historical censoring exposure for a computational experiment; it does not produce a representative sample of stores or products.'
    ],
    3:[
        'Availability flags tell us when demand could have been hidden, but the public data do not tell us how much was lost. The identifiable test therefore hides contiguous blocks of 2, 4 or 8 operating hours on originally fully stocked validation days. The hidden values provide a measurable target. The original daily total is also withheld from reconstruction, so it cannot reveal the sum of the hidden observations.',
        'The test compares no recovery, an in-stock hourly profile, and a historical LightGBM regression. Models are fitted before each validation origin. The regression won the registered hidden-hour MAE criterion: 0.04230 against 0.04302 for zero imputation, a modest 1.67% improvement. Its hidden-hour WAPE was nevertheless 98.33% and signed relative bias was −83.10%. Sparse hourly outcomes make small MAE compatible with substantial under-recovery; this is not evidence that naturally lost demand has been reconstructed accurately.',
        'Aggregation tells a different story. The hourly profile’s daily reconstruction WAPE was 19.62%, compared with 29.26% for no recovery, a 32.96% reduction. Its daily signed bias was +3.07%. The learned method’s daily WAPE was 25.80% and bias −24.32%. The hourly profile had worse hidden-hour WAPE, at 135.18%, but better daily reconstruction. A planner should therefore inspect the outcome needed for the decision rather than assume that one recovery leaderboard answers every question.',
        'There were 495 originally fully stocked validation days, each masked at three durations: 1,485 dependent mask cases. Each method produced 6,930 hidden-hour predictions. These repeated masks are not 6,930 independent businesses. Estimated additions on natural stockout hours remain in separate inferred fields; they are not observed customer demand. The methods are point estimators, so no interval calibration claim is made.'
    ],
    4:[
        'Four expanding validation origins at training day indices 62, 69, 76 and 83 each forecast seven days. Recovery models, profiles and forecasting models are refitted from each allowable prefix. The final models use all 90 training days, after which the official seven-day evaluation split is opened once. Forecasts use historical sales lags, calendar date, day of week and series identity. Future actual sales, availability, weather and promotions are excluded.',
        'Validation selected raw-history LightGBM on fully stocked daily WAPE, at 40.52%. On the untouched final holdout, its corresponding WAPE was 38.84%, against 42.84% for seasonal naive: a 9.33% relative reduction. Learned-recovery LightGBM scored 38.95%, and profile-recovery LightGBM scored 40.07%. These results provide no substantial forecasting advantage from the selected recovery step on the fully stocked holdout.',
        'The fully stocked holdout contains 114 of 210 daily observations. Against all recorded sales, WAPE was 40.41% for raw LightGBM, 40.26% for learned-recovery LightGBM, 40.61% for profile recovery and 45.63% for seasonal naive. The small ordering change across scoring subsets matters: all recorded sales can be censored, while fully stocked outcomes are observable but conditionally selected. Neither target reveals natural lost demand.',
        f'All {len(self.leakage["checks"])} leakage checks left forecasts bitwise unchanged after future sales, hourly sales, availability and contextual fields were altered. Final selection was saved before the evaluation read. This verifies the implemented feature boundary; it does not by itself validate the demand and commercial assumptions used later in simulation. No LLM is involved in these recovery or forecast comparisons.'
    ],
    5:[
        'The core policies are B1 seasonal base-stock, B2 LightGBM (s,S), B3 native GRU negative-binomial forecasting with MILP, and B4 the same system with an evidence gate. Each policy and scenario uses 30 simulation seeds, one decision origin, seven days and all 30 series. The opportunity random streams are shared across competing policies rather than being keyed to their names.',
        'The FreshRetailNet physical environment ages stock and serves it by FIFO under a hypothetical three-day shelf life. The inherited MILP is age-unaware: its objective does not optimize age buckets. These are useful controls, but they cannot be described as a full comparison of perishable MILP policies. The independent calculator experiment later examines expiry-aware decisions explicitly.',
        'The cost index combines assumed purchasing, ordering and transfers, holding, shortage penalties and expiry. Opening stock comes from the disclosed warm-up convention and is not charged again. There is no terminal salvage credit. A seven-day result is sensitive to these conventions. Fill, stockouts and expiry are outcomes in an assumed counterfactual world, rather than observed results of a purchase change at Dingdong.'
    ],
    6:[
        'The live factorial crosses a capable deterministic parser with local Qwen2.5-1.5B, and verified inventory-record repair off with repair on. It contains 32 runs and 224 decision days across four scenarios and seeds 13 and 29. The model uses Q4_K_M quantization, temperature 0, an 8,192-token context, model seed 42 and three CPU threads. The parser and LLM receive the same source evidence, supplier grammar, forecasts, optimizer and independent checking tools.',
        f'Cost and fill were exactly equal in {n["parser_llm_identical_pairs"]}/{n["parser_llm_total_pairs"]} matched parser–LLM comparisons. This shows that the model can fit into a bounded replenishment process on the tested grammar. It does not establish an additional economic return from the language layer. Checked source interpretation, bounded execution and refusal to bypass a failed certificate should be evaluated separately from profit or service superiority.',
        f'For the inventory-field fault, enabling repair changed the LLM-arm mean cost by {n["repair_cost_change_pct"]:+.2f}% and fill by {n["repair_fill_change_pp"]:+.2f} percentage points. The capable parser has access to the same verified repair. This is record reconciliation: it checks a derived stock field against fresh source and movement evidence. Statistical demand recovery, measured earlier, estimates sales that were not observed; the two operations should not share a single “AI recovery” benefit label.',
        'A semantic mistake can be blocked by an independent verifier before it becomes an unsafe committed order. Accordingly, the audit reports wrong selector choices, attempts, cache imports, holds and actual committed-plan violations separately. Zero committed violations must not be translated into a claim of zero model errors. Reused cached responses are not fresh independent trials.'
    ],
    7:[
        'The independent sensitivity grid crosses raw and validation-selected recovered forecast histories, order-up-to and age-aware calculators, shelf lives of 1, 3 and 7 days, demand multipliers of 1, 1.25 and 1.5, and 30 seeds. This gives 1,080 runs. Another 60 runs compare a censoring-aware gate with an intentionally defective rule that holds each series after a valid prior stockout. At the origin, that flag comes from the final training day; later it comes from the policy’s own preceding simulated shortage. The rule therefore tests a feedback mechanism, using no future released mask. It is not the primary strong-parser comparator.',
        'Within each cell, both policies receive the same observed-sales demand proxy multiplied by the same registered scenario factor. A recovered estimate from one competing arm is never treated as ground truth. Orders change counterfactual availability and stockouts. Published stockout flags remain evidence about the original retailer’s observations, rather than outcomes that are forced onto every new stocking policy.',
        'This simulator uses continuous normalized quantities, explicit FIFO expiry and inventory conservation. Opening stock is two days of raw historical mean demand, spread across age buckets and supplied as the initial condition. Expected receipts are 98% of ordered quantity; keyed lead-time probabilities are 90%, 8% and 2% for one, two and three days. Unit purchase, shortage and expiry indices are 1, 5 and 1, with holding 0.015 and fixed order cost 0.02 per line. Receipt day counts as a usable life day. Terminal inventory and paid pipeline are recorded without a salvage credit. Where target calculations extend beyond the seven reported days, the frozen forecast repeats by weekday; no later observed outcome is added.',
        'The calculators also differ in target timing. The simple order-up-to rule covers today and tomorrow against current stock and pipeline. The age-aware rule projects stock surviving today’s service and expiry, then covers up to two future days subject to shelf life. Their contrast is therefore not a pure isolated effect of adding age information. Their cost coefficients and quantity conventions also differ from the core MILP, so absolute costs are not pooled across experiments.',
        'At the default three-day shelf life and 1.25 demand multiplier, age-aware raw forecasts cost 303.3008 with 92.0804% fill and 7.5433% expired quantity per demand. Learned-recovery forecasts cost 300.5691 with 92.6071% fill and 8.0615% expiry per demand: cost fell about 0.90% and fill rose 0.53 points, while expiry rose 0.52 points. Some other cells increased cost. Recovery is therefore a trade-off, rather than an unconditional improvement.',
        'The same default censoring-aware gate cost 300.5691 with 92.6071% fill and no held series-days. The deliberately naive stockout-hold gate cost 365.9235 with 82.4591% fill and a mean of 56.9333 held series-days. This benefit belongs to interpreting valid availability evidence correctly. It uses no live language inference and cannot be reported as an LLM gain.'
    ],
    8:[
        'Both panels contain 30 series, but their markets, products, dates, horizons, quantity units and economics differ. M5 supplies a long daily-unit sales hierarchy. FreshRetailNet supplies 90 training days of globally normalized sales with real hourly availability annotations. M5’s matched agent runs last 14 days; the new runs last seven. Absolute cost levels cannot therefore form a meaningful commercial ranking.',
        'The compact comparison reports separately estimated within-dataset contrasts. In M5, parser and LLM cost and fill were identical in all 16 matched pairs. Record repair under the injected field fault reduced mean cost from 894.30 to 796.51, or 10.93%, and raised fill from 89% to 96%, or seven points. The corresponding FreshRetailNet contrasts appear below. These are descriptive engineering effects, not an estimate of what changing datasets does to a retailer’s profit.',
        'FreshRetailNet adds an information question that M5 cannot identify directly: a correct record can show constrained availability without showing how much demand was lost. That supports demand-visibility and perishability controls. It does not turn the language model into a better numerical forecaster. Additional semantic value would need fair, independently held-out tests of difficult supplier and exception language.'
    ],
    9:[
        'The business case supported by this evidence is for a controlled replenishment process. A retailer needs dependable stock evidence, a clearly labelled estimate when availability hides demand, typed supplier rules and checks before spending. An LLM can participate inside that process. The experiments have not established that it should replace the conventional numerical tools or a capable rules engine.',
        'Useful language tasks remain plausible: translating supplier clauses into checkable fields, requesting a permitted exception tool, and producing a decision packet connected to its sources. The study demonstrates bounded feasibility for a controlled synthetic grammar. It does not measure planner time savings, human trust, unfamiliar natural-contract performance or production ROI. Those are distinct investment tests.',
        'A sensible deployment starts in shadow mode with actual inventories and supplier documents, then moves to reviewed orders and spending caps. Compare exception-resolution quality and delay against the existing system. Expand autonomy only where the measured additional benefit covers error, oversight, hosting and integration costs. Where a capable parser gives the same business outcome more cheaply, it remains a reasonable default.'
    ],
    10:[
        'The independent audit reads run folders, decision objects, grounding outputs and model request/response records. It verifies source hashes and event chains, matched comparisons, inventory-repair lineage and exclusion of future or evaluator-only information. Exact requests are written before HTTP calls. Development failures, setup receipts and any partial directories remain in the evidence archive.',
        'Thirty seeds describe simulator variability conditional on one panel, one origin and assumed economics. The live factorial has two seeds and reused cached evidence. The short forecasting holdout and its conditionally stocked subset restrict generalization. No confirmatory p-value, noninferiority claim, participant result or retailer return is asserted.',
        'The wider dissertation proposes additional B1–B11 architectures, a semantic difficulty ladder, free-form multi-agent alternatives, prompt-injection testing and a human study. Those remain unevaluated unless their own saved executions exist. This bounded study completes the registered controls and specified F1–F6 subtests; it does not silently promote the whole proposed research program into a completed result.',
        'Data attribution: Dingdong-Inc, FreshRetailNet-50K, CC BY 4.0; Wang et al. (2025), pinned dataset revision 08c1fab7f9257bc73679d415d65d644165d351d4. M5 retains its competition provenance. The supplied literature review and references are preserved in the updated Word dissertation.'
    ],
    11:[
        'The accompanying files preserve every completed individual run. Numerical, live-agent and independent-policy CSVs have 360, 32 and 1,140 rows respectively. Forecast files contain series/day predictions, observable targets and flags used only for scoring. Artificial-mask files retain each hidden target. Audit files contain model calls, grounding, selector and decision records.',
        'The report directory contains FreshRetailNet_numerical_all_runs.csv, FreshRetailNet_agent_all_runs.csv and FreshRetailNet_policy_sensitivity_all_runs.csv, as well as their descriptive means and matched contrasts. M5_FreshRetailNet_descriptive_comparison.csv records the compatible relative comparisons. The full evidence archive adds raw objects, saved models, frozen code, source receipts and failed attempts.'
    ]}
    question_rows=[
        ['RQ1: performance','Conventional policies and matched parser/LLM outcomes are reported for each core scenario.','All single-agent and free-form architectures remain unexecuted.'],
        ['RQ2: evidence gate','M5 detects known faults; FreshRetailNet exposes spending-cap portability and delay costs.','Effects depend on thresholds and simulator assumptions.'],
        ['RQ3: censored demand','Recovery is measurable on masking; final forecast benefit is mixed; default service/expiry trade-off is small.','Natural unserved demand has no identified target.'],
        ['RQ4: availability vs error','The explicit stockout-gate ablation shows a valid observation need not justify an integrity hold.','The naive comparator is a mechanism test, not a capable production parser.'],
        ['RQ5: semantic value','Cost and fill match the capable parser on the disclosed synthetic grammar.','Unfamiliar, conflicting natural documents are not tested.'],
        ['RQ6: agent organization','Typed outputs and independent checks are operationally audited.','No head-to-head test isolates additional agent roles.'],
        ['RQ7: trace quality','Source-linked records, physical balances and attempts are machine-audited.','Human diagnosis and trust were not measured.'],
        ['RQ8: generalizability','Within-dataset effects and scale-transfer failure are compared descriptively.','One small panel and origin cannot establish retailer-wide external validity.']]
    hypothesis_rows=[
        ['H1','Descriptively consistent','Matched grammar outcomes show no additional LLM cost/fill advantage; this is not a noninferiority test.'],
        ['H2','Mechanism verified; incremental benefit bounded','Detected evidence faults trigger holds or verified repair. Conventional MILP already has zero actual rule violations; cap portability restricts useful execution.'],
        ['H3','Mixed and conditional','Recovery selection, forecast accuracy and service–expiry trade-offs differ; natural demand recovery is not proved.'],
        ['H4','Supported within the targeted ablation','Holding on valid stockout observations worsens simulated service against the censoring-aware rule.'],
        ['H5','Supported on templates; remaining clause untested','The capable parser matches the LLM. Gains on unfamiliar documents were not evaluated.'],
        ['H6','Untested as a comparative hypothesis','No free-form multi-agent baseline isolates typed-state or specialization effects.'],
        ['H7','Untested','No participant study measures fault diagnosis, acceptance or trust.'],
        ['H8','Conditional engineering evidence','Known-fault controls are useful, but the transferred cap fails until a separately calibrated exploratory follow-up.']]
    self.sections[9]['tables'] += [('Answers to the retained research questions',['Question','Completed answer','Scope limit'],question_rows),('Hypotheses: descriptive evidence status, without confirmatory tests',['Hypothesis','Status','Reason'],hypothesis_rows)]
    daily=self.daily_recovery[(self.daily_recovery.series_id=='ALL')&(self.daily_recovery.mask_duration.astype(str)=='all')]
    self.sections[3]['tables'].append(('Pooled daily reconstruction from artificial masks',['Method','Masked cases','WAPE','Signed bias'],[[label(r.method),str(int(r.n)),percent(r.wape,2),percent(r.relative_bias,2)] for r in daily.itertuples()]))
    self.sections[4]['figures'].append('15_forecast_censor_strata')
    numerical_flags=self.numeric.groupby('policy')[['hard_violations','reference_deviations','executed_actions','held_decisions']].sum()
    self.sections[5]['tables'].append(('Actual violations and clean-reference flags are different endpoints',['Policy','Rule checks violated','Reference-only flagged days','Nonzero action days','Held days'],[[policy,str(int(row.hard_violations)),str(int(row.reference_deviations)),str(int(row.executed_actions)),str(int(row.held_decisions))] for policy,row in numerical_flags.iterrows()]))
    agent_flags=self.agent.groupby('arm')[['hard_violations','reference_deviations','executed_actions','held_decisions']].sum()
    self.sections[6]['tables'].append(('Primary execution and reference-distance flags',['Arm','True rule checks violated','Reference-only flagged days','Nonzero action days','Held days'],[[ARM_LABEL[arm],str(int(row.hard_violations)),str(int(row.reference_deviations)),str(int(row.executed_actions)),str(int(row.held_decisions))] for arm,row in agent_flags.iterrows()]))
    for i,paragraphs in prose.items():self.sections[i]['paragraphs']=paragraphs
    if self.eval_population:
        e=self.eval_population
        self.sections[2]['paragraphs'].append(f"After the forecast selection receipt was frozen, evaluation integrity checks examined all {e['rows']:,} released rows. They retain all {e['series']:,} series for seven days, with no duplicate keys or stock-flag alignment failures. Thus the complete source check covers {p['rows']+e['rows']:,} rows. Holdout profiling does not change the training-selected modeling panel.")

    # Append independently audited model performance with all distinct denominators.
    totals=self.audit.get('totals',{})
    if totals:
        naive=self.numeric[self.numeric.policy.isin(['B1','B2'])]
        optimized=self.numeric[self.numeric.policy.eq('B3')]
        language=self.agent[self.agent.arm.str.startswith('llm_')]
        self.sections[5]['paragraphs'].append(f"There were {int(naive.hard_violations.sum()):,} active-rule violations in committed requested plans from the naive B1/B2 policies. The conventional B3 MILP already had {int(optimized.hard_violations.sum())} violations, as did the gated agent configurations. This safety difference cannot be credited uniquely to either the LLM or the gate. The physical environment clamps supplier fulfillment; a requested-plan violation is not evidence of negative physical stock.")
        violation_days=int((self.decisions.recomputed_hard_violations.gt(0)&self.decisions.committed).sum())
        self.sections[5]['paragraphs'].append(f"The violation total counts failed constraint checks across {violation_days:,} committed decision days, rather than that many distinct purchase decisions. A plan can violate more than one rule. The audited family breakdown below shows which checks failed.")
        families=json_read(self.root/'audit/violation_families.json')['violation_families']
        self.sections[5]['tables'].append(('Failed active constraint checks in naive requested plans',['Policy','Supplier capacity','Maximum order','Budget','Total checks'],[[policy,str(values.get('supplier_capacity',0)),str(values.get('max_order',0)),str(values.get('budget',0)),str(sum(values.values()))] for policy,values in families.items()]))
        self.sections[10]['paragraphs'].append(f"The overall audit's {totals.get('committed_true_constraint_violations',0):,} requested-plan violations are concentrated in naive numerical B1/B2. The live LLM arms have {int(language.hard_violations.sum())}; the conventional MILP B3 also has zero. An operational architecture should report feasibility and service together rather than claim a language-specific safety benefit from this difference.")
        self.sections[6]['paragraphs'].append('In the primary live study, 24 physical extraction calls returned 38 exact numerical terms out of 38. That is the bounded numerical head, not the whole supplier corpus. Twelve fresh recovery selections were instruction-compliant in six cases; six wrong selections were refused by independent checks. Across 224 agent decisions, 188 attempted source grounding and all 188 source sets were exact. Each set has 163 clauses, yielding 30,644 checked active clauses; 36 state-gated decisions did not attempt grounding. Exact conditional source coverage must not be relabelled unconditional model accuracy.')
        self.sections[6]['paragraphs'].append('The primary agent configurations made 40 nonzero action-days and held 184 decision-days in total. The nonzero actions occur under capacity cuts, whose smaller orders can pass the inherited spending cap. Normal, field-fault and stale-feed scenarios are fully held. These successful bounded actions demonstrate feasibility in a narrow condition, while the normal-operation failure motivates the explicitly separate transfer follow-up.')
        self.sections[5]['paragraphs'].append('The inherited harmful_execution field combines two tests: violating an actual active constraint, or departing sufficiently from a clean-evidence reference plan. The default quantity-distance thresholds are 12 model units and 50% relative difference. A reference-only flag does not establish a realized harmful order, but it remains a diagnostic rather than being discarded. B3 has 380 such flagged action-days out of 630, while its actual constraint violations are zero.')
        self.sections[6]['paragraphs'].append('All 40 primary nonzero agent action-days carry a reference-distance flag, with zero actual constraint violations. The legacy composite therefore labels them harmful_execution even though the flag arises from plan distance alone. The report retains those counts and makes no claim of zero modeled risk or zero disagreement with the reference. Feasibility, clean-reference agreement and achieved service are distinct endpoints.')

        cases=json_read(self.root/'audit/recovery_cases.json',[])
        case=next((c for c in cases if c.get('arm')=='llm_recovery' and c.get('applied') and c.get('day')==92),None)
        if case:
            row=max((r for r in case['rows'] if r['changed']),key=lambda r:r['quantity_after'])
            decision=self.decisions[(self.decisions.arm=='llm_recovery')&(self.decisions.day==case['day'])&(self.decisions.run==case['run'])].iloc[0] if 'run' in case else self.decisions[(self.decisions.arm=='llm_recovery')&(self.decisions.day==case['day'])&self.decisions.run.str.contains('derived_field_collapse')].iloc[0]
            self.sections[6]['paragraphs'].append(f"A concrete simulator case shows both the value and the limit of this design. On decision day {case['day']}, the damaged derived inventory for {row['series_id']} was {row['quantity_before']:g}, while fresh source and movement balance both showed {row['source_quantity']:g}. The model selected the permitted reconciliation tool. Independent checks restored {row['quantity_after']:g} scaled model units, changed {case['changed_quantities']} permitted quantities and removed the integrity hard failure. The resulting order was still held: proposed spending of {decision.proposed_spend:,.2f} exceeded the 1,500 autonomous cap. Correct exception handling can repair evidence without making the operating policy usable.")
            (self.out/'FreshRetailNet_worked_record_repair_case.json').write_text(json.dumps({'audit_case':case,'selected_row':row,'decision':decision.to_dict()},indent=2))

        selected=[(AUDIT_LABEL.get(k,k),totals[k]) for k in ['runs','decisions','numerical_runs','agent_runs','physical_API_attempts','physical_API_response_records','tokens','client_or_schema_error_events','physical_extraction_responses','physical_selector_responses','committed_true_constraint_violations','held_decisions','selector_instruction_compliant','selector_semantic_error_responses','parser_llm_matched_pairs','parser_llm_pairs_primary_metrics_equal','parser_llm_pairs_all_action_days_equal','physical_mass_balance_cases_passed','forbidden_model_context_findings'] if k in totals]
        self.sections[10]['tables']=[('Primary registered study audit: distinct operational counts',['Measure','Recorded count'],selected)]
        if self.grounding.shape[0]:
            columns=[c for c in ['arm','grounding_attempted_decisions','complete_source_set_exact_decisions','state_gated_unattempted_decisions','false_active_rules','omitted_active_rules'] if c in self.grounding]
            self.grounding.to_csv(self.out/'FreshRetailNet_source_grounding_audit.csv',index=False)
            if len(columns)>=3:
                display={'arm':'Configuration','grounding_attempted_decisions':'Grounding attempted','complete_source_set_exact_decisions':'Exact full source sets','state_gated_unattempted_decisions':'State-gated days','false_active_rules':'False clauses','omitted_active_rules':'Omitted clauses'}
                rows=self.grounding[columns].values.tolist()
                for row in rows:row[0]=ARM_LABEL.get(row[0],row[0])
                self.sections[10]['tables'].append(('Source grounding coverage; full columns in the audit CSV',[display[c] for c in columns],rows))
        self.sections[10]['paragraphs'].insert(1,f'The primary registered study audit captured {totals.get("physical_API_attempts","the recorded")} physical API attempts and {totals.get("tokens","the recorded")} tokens. There were {totals.get("client_or_schema_error_events","the recorded")} client or schema error events and {totals.get("committed_true_constraint_violations","the recorded")} committed true-constraint violations. Instruction-compliant selector responses and semantic selector errors are reported separately below; successful final safety is not a substitute for reliable tool choice. The exploratory follow-up has its own audit and operating counts in its preceding section.')
    captions={
    '01_source_availability':'Training population rates use 72 million operating-hour flags and 4.5 million daily rows. The selected-panel rate uses its first 62 training days. Daily and hourly percentages have different denominators. None measures the quantity of naturally lost demand.',
    '02_artificial_mask_recovery':'Originally observed sales are hidden in 2-, 4- or 8-hour validation blocks. Models fit only earlier data. The target is deliberately withheld sales, rather than unobserved natural stockout demand. Mask cases overlap and are not independent retailers.',
    '03_forecast_accuracy':'All four frozen forecasts cover the same 30 series and seven-day holdout. All-day sales can be censored; fully stocked days are conditionally observable. Future masks, weather, promotion and actual sales are never forecast features.',
    '04_daily_forecasts':'Panel aggregate predictions are made at the final training origin. Later forecast lags use predictions. Globally normalized sales are neither physical quantities nor a currency.',
    '05_numerical_cost_service':'Thirty seeds per policy and scenario. Error bars show one standard deviation across conditional simulation seeds. B1: seasonal base-stock; B2: LightGBM (s,S); B3: native GRU-NB with MILP; B4: B3 with an evidence gate. Physical stock expires by FIFO, while the inherited MILP remains age-unaware.',
    '06_numerical_safety':'Counts are summed across 30 runs per cell. Hard violations use actual active constraints and differ from clean-reference deviations. Holds count decision days. A low violation count should be assessed alongside service.',
    '07_agent_factorial':'Four matched controller/repair configurations, four scenarios and two seeds, each using seven days. Repair is verified inventory-record reconciliation. It is not statistical recovery of censored customer demand. All configurations share the numerical tools and source evidence.',
    '08_m5_fresh_comparison':'Separate within-dataset repair effects under an inventory-field fault. M5 uses 14 days and FreshRetailNet seven; their quantity units, expiry assumptions and economics differ. The calibrated result is an exploratory reused-holdout follow-up with 30 simulator seeds; the other bars have two. These bars do not estimate a commercial return or the causal effect of changing datasets.',
    '09_service_waste_frontier':'Independent calculator experiment, with 30 shared seeds per cell and demand multiplier 1. Shelf lives are hypothetical. Quantity and economic conventions differ from the core MILP experiment, so their absolute results are not pooled.',
    '10_sensitivity_costs':'Relative change from raw to recovered forecast history within each common-demand cell. Thirty seeds, shelf lives 1/3/7 days and demand multipliers 1/1.25/1.5. Negative values mean lower assumed indexed cost. They are not measured profit or purchasing savings.',
    '11_local_model_calls':'Primary registered study physical attempts from the independent audit. The exploratory follow-up has separate counts. Cache imports and typed extraction outputs are not additional calls. Local API charges are zero; CPU hosting and oversight costs are unpriced.'}
    for f in self.figures:f['caption']=captions.get(f['key'],f['caption'])
    for s in self.sections:
        clean=[]
        for title,columns,rows in s['tables']:
            title=title.replace('30seeds','30 seeds').replace('twoseeds','two seeds').replace('fullgrid','full grid').replace('32livefactorialruns','32 live factorial runs')
            clean.append((title,columns,rows))
        s['tables']=clean


def clean_doc_replacements(self):
    n=self.findings;p=self.population
    means=self.agent_means.set_index(['arm','scenario']);before=means.loc[('llm_no_recovery','derived_field_collapse')];after=means.loc[('llm_recovery','derived_field_collapse')]
    return {
    3:'Master’s dissertation • Completed bounded two-benchmark study • October 2026',
    13:'Chapter 6 Experimental Findings and Managerial Implications',
    19:'Appendix C Completed Studies and Remaining Evidence Gaps',
    299:'Table 5.5. FreshRetailNet research questions; executed coverage and remaining gaps are documented in Chapter 6 and Appendix C.',
    333:'6.2 FreshRetailNet: observed evidence and simulated outcomes',
    336:'6.3 Demand recovery, forecast performance and perishable decisions',
    339:'6.4 The incremental contribution of the LLM',
    342:'6.5 Evidence gating, record repair and useful holds',
    345:'6.6 Reliability and the limits of multi-agent attribution',
    347:'6.7 Decision traces and human review: evidence and future validation',
    380:'8.2 Completed evidence and its limits',
    468:'Appendix C Completed Studies and Remaining Evidence Gaps',

    4:'This manuscript integrates the completed M5 studies with a bounded FreshRetailNet experiment. All 50,000 source series were profiled for integrity; forecasting and replenishment results concern the disclosed 30-series computational panel. Observed, inferred and simulated quantities remain distinct.',
    26:'The evaluation connects M5 engineering evidence to a completed FreshRetailNet study on 30 training-selected store–product series. The new study contains 360 conventional replenishment runs, 32 matched parser and LLM runs, and 1,140 independent policy experiments. Recovery was measured against artificially hidden sales; four forecasts were scored on a protected seven-day holdout. The language model used the same source checks, numerical tools and safeguards as a capable parser. Cost and fill were identical in '+str(n['parser_llm_identical_pairs'])+' of '+str(n['parser_llm_total_pairs'])+' matched FreshRetailNet pairs; all matched M5 pairs were also identical. Inventory-record repair changed FreshRetailNet fault-scenario cost by '+number(n['repair_cost_change_pct'])+'% and fill by '+number(n['repair_fill_change_pp'])+' percentage points. Shared benefits belong to the recovery architecture, rather than an isolated language-model advantage. The findings support auditable automation and availability-aware analysis. They do not measure real retailer profit, naturally lost demand or human trust. Additional language value requires fair testing on harder, independently held-out business documents and field outcomes.',
    256:'FreshRetailNet-50K was released by Dingdong in 2025 as a benchmark for recovering demand hidden by stockouts (Wang et al., 2025). The downloaded revision contains 50,000 store–product series, 898 stores, 18 cities and 865 distinct product IDs. The paper identifies 863 SKUs and the card 865; this study records both source claims and uses the 865 IDs present in revision 08c1fab7f9257bc73679d415d65d644165d351d4.',
    259:'The publisher’s implementation and released arrays confirm that 1 means out of stock. The 16 operating-hour slots run from 06:00 through 21:00. In training, 19.8771% of operating hours are flagged and 44.2668% of daily rows have at least one stockout hour. Daily stockout counts match the sums of hourly flags. These percentages describe different denominators, and neither reveals the unobserved quantity of lost demand.',
    266:'The executed protocol retains the official 90-day training split and sealed seven-day evaluation. Four expanding validation folds begin at day indices 62, 69, 76 and 83. Panel membership uses only the first 62 days. Forecast selection uses pooled fully stocked validation WAPE; recovery selection uses pooled artificial-mask hourly MAE. Both selections are saved before final evaluation is opened.',
    273:'The executed recovery comparators are no imputation, a train-only in-stock hourly profile and historical LightGBM regression. The regression is named accurately; Tobit, TimesNet and the other ambitious candidate models were not executed within this budget. Each model fits before its validation origin. Published sales remain unchanged, with recovery estimates stored separately.',
    274:'The identifiable recovery test hides deterministic contiguous blocks of 2, 4 and 8 operating hours on originally fully stocked validation days. The hidden values and original full daily amount are withheld from reconstruction inputs. Saved outputs measure hourly and daily MAE, WAPE and signed bias with explicit denominators. The masks use reproducible hash ranking; they do not claim to reproduce the full empirical distribution of natural stockouts. These are point estimates, without an asserted interval-calibration result.',
    277:'The four executed forecasts are seasonal naive, LightGBM on raw sales, LightGBM on hourly-profile recovered history and LightGBM on learned recovered history. LightGBM variants use the same 150-estimator budget and historical lag, date and series features. Recursive forecasts use predictions for later lags. Future realized weather, promotions and stockout masks are excluded. TFT, DeepAR and prediction-interval calibration were not executed for this new holdout.',
    278:'Final results separately score all recorded daily sales and fully stocked daily outcomes, using WAPE, MAE and signed bias with their denominators. Every frozen candidate remains visible. True lost demand is unavailable during natural stockouts, so a higher prediction cannot be labelled proven recovery of real lost sales. Published-paper results remain literature evidence.',
    283:'The bounded experiment gives competing policies a common observed-sales-proxy demand stream. It does not use one recovered arm’s estimates as evaluation truth. The core adapter scales published normalized sales by 100 for inherited order-tool compatibility. The independent continuous calculator experiment tests common demand multipliers of 1, 1.25 and 1.5. Each policy’s orders create new simulated availability; historical stockout masks are not replayed as outcomes under the changed stocking policy.',
    284:'The core physical environment uses FIFO expiry and a hypothetical three-day life, while its inherited MILP remains age-unaware. A separate calculator experiment varies life across 1, 3 and 7 days, raw and recovered forecast views, demand multipliers and 30 seeds. Its continuous quantities and dimensionless cost coefficients differ from the core, so absolute outcomes are not pooled. These assumed lives are not measured Dingdong product attributes.',
    297:'The six FreshRetailNet questions guide different forms of evidence. F1 and F2 receive masking and holdout forecasting tests. F3 and F6 receive independent age-aware simulations and sensitivity. F4 receives a targeted stockout-gate mechanism ablation. F5 receives a parser-versus-LLM factorial using shared forecasts and tools. The full architecture suite, independent natural supplier documents and human study remain outside the completed scope.',
    315:'Thirty matched seeds describe conventional-policy variability conditional on one panel, origin and assumed world. The live factorial has seeds 13 and 29. Days, SKU rows and cached responses are not independent replications. All seed rows and descriptive contrasts are available; no confirmatory significance, noninferiority or population-precision claim is made.',
    316:'The completed grid contains 360 conventional core runs, 32 live factorial runs, 1,080 independent calculator sensitivity runs and 60 gate ablations. Each run uses 30 series over seven days. The numerical grid’s 30 seeds do not create 30-seed evidence for the two-seed LLM comparison.',
    324:'The full 4.85 million daily-row release is checked for integrity. Modeling and closed-loop conclusions use a frozen 30-series panel selected from first-62-day sales and censoring evidence. The sample is balanced across three historical censoring strata by deterministic ranking and covers 0.06% of series. It is computationally purposive rather than representative.',
    325:'CHAPTER 6  Experimental Findings and Managerial Implications',
    335:'The completed checks show why a valid stockout flag calls for a different response from a broken feed. Correct availability evidence can limit what sales reveal. Demand recovery therefore supplies a separate inference. An inventory-field failure is instead held or reconciled through permitted source and movement evidence. The following results retain that distinction.',
    337:'The saved recovery and forecast results are leakage-checked and reported below. Artificial masking supplies known targets; natural stockouts do not. Validation chose raw-history LightGBM for forecasting and the hourly LightGBM model for recovery before the final holdout was read. Every alternative remains in the tables. The selected recovery method’s modest hourly MAE advantage does not establish reliable reconstruction of naturally lost demand.',
    338:'A planner needs more than an accuracy ranking. Under the default independent age-aware setting, recovered forecasts modestly reduced indexed cost and improved fill, while increasing expiry per demand. Some other cells increased cost. The service–waste frontier and sensitivity table therefore show a conditional trade-off, rather than an unconditional return from recovery.',
    340:'The live factorial found identical cost and fill in '+str(n['parser_llm_identical_pairs'])+'/'+str(n['parser_llm_total_pairs'])+' matched parser–LLM pairs. Supplier texts remain synthetic and controlled. The model’s participation demonstrates bounded execution with numerical tools and independent verification, but no additional economic advantage over a capable parser on these texts.',
    341:'A retailer still needs evidence that language assistance resolves exceptions that its rules cannot handle economically. Checked interpretation and exception coordination are plausible benefits. Inference delay, monitoring, failures and oversight are incremental costs. The audit retains erroneous choices even when independent tools prevent unsafe execution. Semantic accuracy and final execution safety are separate endpoints.',
    343:'Under the injected inventory-field fault, enabling repair changed mean LLM-arm cost from '+number(before.cost)+' to '+number(after.cost)+' ('+number(n['repair_cost_change_pct'])+'%) and fill from '+percent(before.fill_rate)+' to '+percent(after.fill_rate)+' ('+number(n['repair_fill_change_pp'])+' percentage points). The parser receives the same repair. Valid historical stockouts remain demand-visibility evidence; the deliberately naive stockout-hold rule is reported only as a targeted mechanism ablation.',
    346:'The system has typed responsibilities and independent verification, but this study does not isolate the number of agents or execute every single-agent and free-form-chat comparison proposed earlier. Measured reliability comes from source grounding, completed traces, model calls and committed outcomes. Cache imports are not new trials. Conclusions concern the verified workflow rather than an untested benefit of additional agents.',
    349:'No formal human study was conducted. Structured traces make decisions reviewable and their references auditable. They do not by themselves prove faster review, lower staff effort or better calibrated trust. Those remain participant-study outcomes for later validation.',
    351:'The two datasets support a useful but bounded synthesis. Verified record repair can restore continuity under a known fault; the LLM does not add measured cost or fill value over the capable parser on the tested grammar. FreshRetailNet contributes observed stockout timing, artificial-mask tests and perishable trade-offs. These are information and architecture findings, rather than proof that the language model forecasts better. Cross-dataset contrasts remain descriptive because horizons and economics differ.',
    360:'Evidence gating shifts some risk from acting on bad information to delaying an action. The primary cap-transfer failure makes that trade-off visible: a controller can satisfy its checks while systematically withholding otherwise feasible orders. Holds should therefore be reported beside service, cost and the time needed to resolve an exception. A held decision means the permitted action lacks sufficient support under the current policy; it does not establish that ordering nothing is the best business decision.',
    373:'Fourth, outcomes from simulated alternative replenishment policies are conditional on assumptions about demand, delivery, prices, shelf lives and objectives. The independent calculator sensitivity does not share the core MILP’s economic scale or target timing. Fifth, supplier documents use a controlled synthetic grammar and a single small local LLM. No human study was conducted. The primary agent experiment has two seeds; the calibrated follow-up has 30 but reuses the holdout after the portability failure was observed. These studies do not establish population significance or production returns.',
    374:'The valid claims concern disclosed engineering comparisons: recovery and forecasting on observable targets, conditional policy outcomes in simulation, and the behaviour of structured checks under known test faults. The negative findings also matter. A conventional constrained optimizer already avoided the tested hard-rule violations, and the capable parser matched the LLM’s cost and fill. External business impact requires independent inventory, procurement and outcome data from a partner retailer.',
    381:'Both benchmarks now contribute completed engineering evidence. The M5 archive retains the original numerical and feasibility study and the 32-run V2 pilot. FreshRetailNet adds 360 numerical runs, 32 matched live runs, 1,080 independent calculator sensitivity runs, 60 gate ablations, and protected forecast and recovery tests on 30 series. The evidence is conditional on the disclosed panel and assumptions. It does not validate all 50,000 series, natural lost demand, commercial waste savings or human trust. Incremental language-model superiority over a capable parser was not demonstrated by matched cost and fill.',
    382:'Neither public benchmark reveals every business quantity needed to evaluate actual purchase orders. True lost demand is unobserved during natural stockouts, and the cost of an alternative inventory policy can only be evaluated inside an assumed counterfactual world. The completed checks support the implemented feature and evidence boundaries. They do not remove the need to test those assumptions against independent operational outcomes.',
    384:'The next priority is broader independent product, store and time panels, real inventory and procurement ground truth, and difficult natural supplier documents with fair parser coverage. Recovery needs external checks where demand was not observed. Human review and partner shadow-mode studies could measure whether checked language assistance improves exception resolution enough to justify its operating burden.',
    449:'The implementation program below records what the bounded study completed and what remains outside its scope. Source validation, the 30-series adapters, point-recovery and forecast tests, physical expiry, controlled agent comparisons and descriptive reporting are complete. Full architecture ablations and human participants remain unevaluated.',
    454:'This pseudocode explains the timing and validation rules; the executable scripts and saved manifests are the authoritative implementation. The downloaded release passed key, chronology, hourly sales and flag-alignment checks. Discount values outside [0,1] were retained and reported, with discount excluded from forecasting features. Every exception and exclusion remains in the source-validation receipts.',
    456:'research_version: freshretailnet-bounded-v1\nstatus: completed_bounded_panel\nsource_revision: 08c1fab7f9257bc73679d415d65d644165d351d4\nsource_license: CC-BY-4.0\nsource_validation: 4,850,000 daily rows / 50,000 series\nmodeling_panel: 30 training-selected series / 0.06% of release\nfinal_horizon: 7 days after 90 training days\nprimary_numerical: 360 runs / 30 simulation seeds\nprimary_agent_factorial: 32 runs / 2 simulation seeds\nindependent_policy_sensitivity: 1,080 runs\nstockout_gate_ablation: 60 runs\nquantity_scale_core: 100 * globally_normalized_sales\neconomic_unit: simulated cost index\nfuture_observed_features: prohibited\nnatural_lost_demand_ground_truth: unavailable\nhuman_participants: not conducted\nauthoritative_manifest: results/freshretailnet/protocol.json',
    458:'The completed figures and tables cover source availability and panel composition, identifiable artificial-mask recovery, daily reconstruction, protected forecasts and training censoring strata, core cost/service/safety, matched parser–LLM outcomes, independent service–expiry and demand sensitivity, operating attempts and within-dataset M5 comparisons. Appendices retain all individual core runs and original M5 figures. Full sensitivity and exploratory transfer rows are accompanying CSVs. Human explanation outcomes, a full semantic-difficulty suite and production financial returns are not included because they were not measured.',
    461:'The 19 published columns were verified in the pinned Parquet files: city_id, store_id, management_group_id, first_category_id, second_category_id, third_category_id, product_id, dt, sale_amount, hours_sale, stock_hour6_22_cnt, hours_stock_status, discount, holiday_flag, activity_flag, precpt, avg_temperature, avg_humidity and avg_wind_level. Sales/hierarchy and historical availability are used by the bounded adapter. Future realized contextual fields are not planning features (Dingdong-Inc, 2025).',
    472:'Unevaluated commitments include full-population FreshRetailNet modeling, ground truth for naturally unserved demand, all B1–B11 architecture ablations, independent heterogeneous supplier-language tests, a formal human audit and actual ERP purchasing. Completed 30-series forecasts, recovery tests and simulator runs are reported in the following appendices. Unexecuted tests have no invented outcome.'}


def validate_report(root,source=DEFAULT_SOURCE):
    """Check rendered delivery, the untouched source and every embedded figure."""
    import zipfile
    import fitz
    root=Path(root);out=root/'report';manifest=json_read(require(out/'report_manifest.json'))
    docx=out/'Dual_Benchmark_Dissertation_Final.docx';document=Document(require(docx))
    source=Path(source);original=Document(require(source))
    paragraphs=[p.text for p in document.paragraphs]
    checks={}
    chapters=[]
    for paragraph in document.paragraphs:
        match=re.match(r'^CHAPTER\s+(\d+)\b',paragraph.text)
        if match and paragraph.style.name=='Heading 1':chapters.append(int(match.group(1)))
    checks['all_eight_chapters']=chapters==list(range(1,9))
    checks['literature_preserved']=all(p.text in paragraphs for p in original.paragraphs[72:112] if p.text)
    checks['references_preserved']=all(p.text in paragraphs for p in original.paragraphs[387:447] if p.text)
    checks['primary_all_392_individual_rows']=sum(len(t.rows)-1 for t in document.tables if t._tbl.getprevious() is not None and ''.join(t._tbl.getprevious().itertext()).startswith(('All 30 seeds per policy','All 32 live factorial runs')))==392
    build=json_read(out/'dissertation_build_manifest.json')
    with zipfile.ZipFile(docx) as z:media={hashlib.sha256(z.read(p)).hexdigest() for p in z.namelist() if p.startswith('word/media/')}
    checks['all_figures_embedded_exactly']=all(x['sha256'] in media for x in build['embedded_figures'])
    checks['figure_count']=len(document.inline_shapes)==build['fresh_figures']+build['m5_figures']
    checks['m5_404_individual_rows']=sum(len(t.rows)-1 for t in document.tables if t._tbl.getprevious() is not None and ''.join(t._tbl.getprevious().itertext()).startswith('Table F.') and not ''.join(t._tbl.getprevious().itertext()).startswith('Table F.19.'))==404
    checks['source_doc_unchanged']=sha(source)==build['source_sha256']
    checks['no_compact_drafting_tokens']=not any(x in '\n'.join(paragraphs) for x in ['first62trainingdays','37,097eligible','2,4or8contiguous','day62,69,76and83each','ThepreviousM5studies'])
    def input_path(path):
        if path==build['source']:return source
        return Path(path) if Path(path).is_absolute() else ROOT/path
    checks['input_hashes_unchanged']=all(sha(require(input_path(path)))==item['sha256'] for path,item in manifest['inputs'].items())
    transfer_protocol=json_read(root/'gate_transfer/protocol.json')
    if transfer_protocol:
        checks['registered_transfer_rows_complete']=manifest['individual_runs']['exploratory_gate_transfer']==transfer_protocol['evaluation']['expected_runs']
        checks['registered_transfer_owner_audit_verified']=str(json_read(require(root/'gate_transfer/audit.json')).get('status','')).startswith('VERIFIED')
        checks['registered_transfer_independent_audit_verified']=json_read(require(root/'gate_transfer/audit_independent/audit.json')).get('status')=='VERIFIED_COMPLETE'
    rendered={}
    paths=[require(out/'FreshRetailNet_study_report.pdf'),require(out/'Dual_Benchmark_Dissertation_Final.pdf')]
    for path in paths:
        pdf=fitz.open(path);text_pages=[page.get_text() for page in pdf];problems=[]
        for i,page in enumerate(pdf):
            body=page.get_text(clip=fitz.Rect(0,30,page.rect.width,page.rect.height-40))
            if len(body.strip())<20:problems.append({'page':i+1,'problem':'unexpected blank page'})
            for x in page.get_text('dict')['blocks']:
                box=fitz.Rect(x['bbox'])
                if box.x0<-.5 or box.y0<-.5 or box.x1>page.rect.width+.5 or box.y1>page.rect.height+.5:problems.append({'page':i+1,'problem':('image' if x['type']==1 else 'text')+' beyond page','bbox':list(box)})
        rendered[path.name]={'pages':len(pdf),'sha256':sha(path),'bytes':path.stat().st_size,'layout_findings':problems}
        checks[path.stem+'_nonempty_pages_and_content_bounds']=not problems
        (out/(path.stem+'_rendered_text.txt')).write_text('\n\n'.join(text_pages))
    checks['word_source_epoch_removed']='No FreshRetailNet experiments have yet been run' not in '\n'.join(paragraphs)
    checks['all_checks_passed']=all(checks.values())
    receipt={'status':'VERIFIED' if checks['all_checks_passed'] else 'FAILED','checked_utc':datetime.now(timezone.utc).isoformat(),'checks':checks,'rendered':rendered,'docx_sha256':sha(docx),'paragraphs':len(document.paragraphs),'tables':len(document.tables),'inline_figures':len(document.inline_shapes),'reference_paragraphs_preserved':sum(bool(p.text) for p in original.paragraphs[387:447])}
    (out/'report_validation.json').write_text(json.dumps(receipt,indent=2));print(json.dumps(receipt,indent=2))
    if not checks['all_checks_passed']:raise RuntimeError('Rendered dissertation validation failed; inspect report_validation.json')
    manifest['status']='VERIFIED';manifest['validated_utc']=receipt['checked_utc'];manifest['outputs']=[{'path':str(p.relative_to(out)),'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(out.rglob('*')) if p.is_file() and p.name!='report_manifest.json' and p.suffix!='.log'];(out/'report_manifest.json').write_text(json.dumps(manifest,indent=2))
    return receipt


def escape(s):return str(s).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')

def prose_text(s):
    return str(s)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--root',type=Path,default=ROOT/'results/freshretailnet');p.add_argument('--source',type=Path,default=DEFAULT_SOURCE);p.add_argument('--allow-partial',action='store_true');p.add_argument('--validate',action='store_true');a=p.parse_args();validate_report(a.root,a.source) if a.validate else Study(a.root,a.source,a.allow_partial).build()
