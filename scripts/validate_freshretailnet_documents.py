"""Check exported dissertation/report integrity and layout against the source."""
import argparse,hashlib,json,re,zipfile
from pathlib import Path
from collections import Counter
import fitz
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'study_results/research_context/Evidence_Gated_Autonomy_M5_FreshRetailNet_Full_Dissertation.docx'

def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def inspect_pdf(path):
    d=fitz.open(path);issues=[];image_pages=[];text=[]
    for number,page in enumerate(d,1):
        content=page.get_text();text.append(content);displayed_images=page.get_image_info()
        body_words=[w for w in page.get_text('words') if w[1]>=55 and w[3]<=page.rect.height-55]
        body_images=[item for item in displayed_images if item['bbox'][3]>55 and item['bbox'][1]<page.rect.height-55]
        if not body_words and not body_images:issues.append({'page':number,'issue':'Empty page body'})
        if len(content.strip())<15 and not displayed_images:issues.append({'page':number,'issue':'Blank or nearly blank page'})
        for word in page.get_text('words'):
            if word[0]<-1 or word[1]<-1 or word[2]>page.rect.width+1 or word[3]>page.rect.height+1:issues.append({'page':number,'issue':'Text outside page','word':word[4]})
        if displayed_images:image_pages.append(number)
    combined='\n'.join(text)
    return {'file':str(path),'sha256':sha(path),'pages':len(d),'image_pages':image_pages,'issues':issues,'text':combined}

def main():
    parser=argparse.ArgumentParser();parser.add_argument('docx',type=Path);parser.add_argument('--pdf',type=Path);parser.add_argument('--report-pdf',type=Path);parser.add_argument('--out',type=Path,required=True);args=parser.parse_args()
    args.out.parent.mkdir(parents=True,exist_ok=True)
    source=Document(SOURCE);doc=Document(args.docx);paragraphs=[p.text for p in doc.paragraphs];text='\n'.join(paragraphs)
    original_refs=[p.text for p in source.paragraphs if p.style.name=='Reference' and p.text.strip()]
    missing_refs=[p for p in original_refs if p not in paragraphs]
    chapters=[p.text for p in doc.paragraphs if p.style.name=='Heading 1' and re.match(r'^CHAPTER\s+[1-8]\b',p.text,re.I)]
    missing_chapters=[i for i in range(1,9) if not any(re.match(fr'^CHAPTER\s+{i}\b',p,re.I) for p in chapters)]
    with zipfile.ZipFile(args.docx) as archive:
        names=set(archive.namelist());media=[p for p in names if p.startswith('word/media/')]
        image_hashes=[hashlib.sha256(archive.read(p)).hexdigest() for p in media]
        # Every internal relationship must resolve within the Word package.
        from lxml import etree
        broken=[]
        for rel in [p for p in names if p.endswith('.rels')]:
            tree=etree.fromstring(archive.read(rel));base=Path(rel).parent.parent
            for element in tree:
                if element.get('TargetMode')=='External':continue
                target=element.get('Target','');resolved=str(base/target)
                import posixpath
                resolved=posixpath.normpath(resolved).lstrip('/')
                if resolved not in names:broken.append({'relationship':rel,'target':target})
    scientific_figures=list((args.docx.parent/'figures').glob('*.png'))
    missing_figures=[str(p) for p in scientific_figures if sha(p) not in image_hashes]
    old_manifest=json.loads((ROOT/'study_results/dissertation/dissertation_build_manifest.json').read_text())
    historical_missing=[]
    for entry in old_manifest['figures']:
        checksum=entry.get('sha256') if isinstance(entry,dict) else None
        if checksum and checksum not in image_hashes:historical_missing.append(entry)
    issues=[]
    if missing_chapters:issues.append({'missing_chapters':missing_chapters})
    if len(chapters)!=8:issues.append({'chapter_count':len(chapters),'expected':8})
    if missing_refs:issues.append({'missing_original_references':missing_refs})
    if broken:issues.append({'broken_word_relationships':broken})
    if missing_figures:issues.append({'missing_new_figures':missing_figures})
    if historical_missing:issues.append({'missing_historical_M5_figures':historical_missing})
    for phrase in ['No FreshRetailNet experiments have yet been run','FreshRetailNet experimentation is proposed, not presented as completed','status: design_only_for_freshretailnet']:
        if phrase in text:issues.append({'stale_status_phrase':phrase})
    result={'status':'VERIFIED' if not issues else 'FAILED','docx':str(args.docx),'docx_sha256':sha(args.docx),'source_docx_sha256':sha(SOURCE),'chapters':chapters,'paragraphs':len(paragraphs),'tables':len(doc.tables),'inline_images':len(doc.inline_shapes),'media_files':len(media),'original_references':len(original_refs),'issues':issues,'pdf_checks':[]}
    for path in [args.pdf,args.report_pdf]:
        if path:
            check=inspect_pdf(path);export_text=args.out.parent/(path.stem+'_rendered_text.txt');export_text.write_text(check.pop('text'));result['pdf_checks'].append(check)
            if check['issues']:issues.extend(check['issues'])
    result['status']='VERIFIED' if not issues else 'FAILED';args.out.parent.mkdir(parents=True,exist_ok=True);args.out.write_text(json.dumps(result,indent=2)+'\n');print(json.dumps({k:v for k,v in result.items() if k!='chapters'},indent=2))
    if issues:raise SystemExit(1)

if __name__=='__main__':main()
