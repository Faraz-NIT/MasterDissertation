"""Validate the new manuscript, source values, Word format and PDF layout."""
from __future__ import annotations
import argparse,hashlib,json,re,zipfile
from datetime import datetime,timezone
from pathlib import Path
import fitz
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'study_results/business_dissertation'
DOCX=OUT/'LLM_Replenishment_Business_Dissertation.docx';PDF=DOCX.with_suffix('.pdf')
def js(p):return json.loads(Path(p).read_text())
def sha(p):
 with Path(p).open('rb') as s:return hashlib.file_digest(s,'sha256').hexdigest()
def words(s):return len(re.findall(r"\b[\w’'-]+\b",s))
def clean(s):return re.sub(r'\s+',' ',s).strip()
def contents(pdf):
 headings=js(OUT/'heading_register.json');found={}
 for n,page in enumerate(pdf,1):
  large=[s['text'] for b in page.get_text('dict')['blocks'] for l in b.get('lines',[]) for s in l['spans'] if s['size']>=12.8]
  big=clean(' '.join(large))
  for h in headings:
   if clean(h['title']) in big:found.setdefault(h['title'],n)
 if len(found)!=len(headings):raise ValueError(['Unmapped headings']+[h['title'] for h in headings if h['title'] not in found])
 return found

def validate(update=False):
 pdf=fitz.open(PDF);actual_pages=contents(pdf)
 if update:
  (OUT/'contents_pages.json').write_text(json.dumps(actual_pages,indent=2)+'\n');print(json.dumps({'status':'MAPPED','headings':len(actual_pages),'pdf_pages':len(pdf)}));return
 issues=[];checks={}
 def check(name,condition):
  checks[name]=bool(condition)
  if not condition:issues.append(name)
 doc=Document(DOCX);manifest=js(OUT/'build_manifest.json');register=js(OUT/'figure_register.json');brief=js(OUT/'rewrite_brief.json')
 text=(OUT/'manuscript.md').read_text();source_prose=[x for x in text.splitlines() if x.strip() and not x.startswith('#') and not x.startswith('[[')]
 source_word_count=sum(words(x) for x in source_prose)
 check('main_prose_minimum_15000',source_word_count>=15000)
 check('word_count_matches_build_manifest',source_word_count==manifest['main_prose_words_excluding_front_matter_references_appendices'])
 actual_prose=[];active=False
 for p in doc.paragraphs:
  if p.style.name=='Heading 1':
   if re.match(r'^1\. ',p.text):active=True
   if p.text=='References':active=False
  if active and p.style.name=='Normal' and p.text.strip():actual_prose.append(p.text)
 check('whole_new_main_manuscript_in_word',actual_prose==source_prose)
 check('main_style_font_size_spacing_alignment',doc.styles['Normal'].font.name=='Times New Roman' and doc.styles['Normal'].font.size.pt==12 and doc.styles['Normal'].paragraph_format.line_spacing==1.5 and doc.styles['Normal'].paragraph_format.alignment==WD_ALIGN_PARAGRAPH.JUSTIFY)
 check('all_margins_one_inch',all(getattr(s,x).twips==1440 for s in doc.sections for x in ['top_margin','bottom_margin','left_margin','right_margin']))
 check('no_theme_font_override',not any('theme' in a.lower() for f in doc.styles.element.iter(qn('w:rFonts')) for a in f.attrib))
 check('six_main_chapters',len([p for p in doc.paragraphs if p.style.name=='Heading 1' and re.match(r'^[1-6]\. ',p.text)])==6)
 check('brief_results',next(v for k,v in manifest['chapter_words'].items() if k.startswith('4.'))<2000)
 check('brief_conclusion',next(v for k,v in manifest['chapter_words'].items() if k.startswith('6.'))<900)
 check('model_and_training_distinction_present','not trained or fine-tuned on M5 or FreshRetailNet' in text and '32 GiB' in text and 'No GPU was available' in text)
 check('reader_equality_and_exploratory_status_explicit','240 reader comparisons' in text and 'exploratory' in text and 'capable rule-based reader achieved the same' in text)
 check('contents_pages_match_export',js(OUT/'contents_pages.json')==actual_pages)
 with zipfile.ZipFile(DOCX) as z:
  embedded=[hashlib.sha256(z.read(p)).hexdigest() for p in z.namelist() if p.startswith('word/media/')]
  xml=etree.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
  bookmarks=set(xml.xpath('//w:bookmarkStart/@w:name',namespaces=ns));links=xml.xpath('//w:hyperlink/@w:anchor',namespaces=ns)
 check('all_contents_and_figure_links_resolve',set(links).issubset(bookmarks))
 check('twenty_two_exact_embedded_figures',len(doc.inline_shapes)==22 and len(embedded)==22 and all(f['sha256'] in embedded and sha(OUT/f['path'])==f['sha256'] for f in register['figures']))
 check('every_figure_has_dataset_label',all(f['dataset'] in {'M5','RETAILNET','M5 + RETAILNET'} for f in register['figures']))
 check('eight_explanatory_tables',len(doc.tables)==8)
 check('input_evidence_hashes_unchanged',all(sha(ROOT/p)==r['sha256'] for p,r in register['inputs'].items()))
 check('previous_published_documents_and_graphs_unchanged',all(sha(ROOT/'study_results/freshretailnet'/r['path'])==r['sha256'] for r in brief['protected_previous_report_outputs']))
 # Recalculate the headline evidence independently of figure-building code.
 m=pd.read_csv(ROOT/'study_results/v2/M5_v2_results.csv');f=pd.read_csv(ROOT/'study_results/freshretailnet/FreshRetailNet_exploratory_gate_transfer_all_runs.csv');primary=pd.read_csv(ROOT/'study_results/freshretailnet/FreshRetailNet_agent_all_runs.csv')
 effects={};all_equal=True
 for name,frame,prefix in [('M5',m,'llm_v2_'),('RetailNet',f,'llm_')]:
  mean=frame.groupby(['arm','scenario'])[['cost','fill_rate']].mean();before=mean.loc[(prefix+'no_recovery','derived_field_collapse')];after=mean.loc[(prefix+'recovery','derived_field_collapse')]
  effects[name]={'cost_change_pct':100*(after.cost/before.cost-1),'fill_change_pp':100*(after.fill_rate-before.fill_rate)}
  for status in ['no_recovery','recovery']:
   llm=frame[frame.arm.eq(prefix+status)];parser=frame[frame.arm.eq('parser_'+status)];p=llm.merge(parser,on=['scenario','seed'],suffixes=('_llm','_reader'),validate='one_to_one')
   for metric in ['cost','fill_rate','held_decisions','hard_violations']:all_equal=all_equal and np.array_equal(p[metric+'_llm'],p[metric+'_reader'])
 check('headline_m5_effect_recomputed',round(effects['M5']['cost_change_pct'],2)==-10.93 and round(effects['M5']['fill_change_pp'],2)==7.00)
 check('headline_retailnet_effect_recomputed',round(effects['RetailNet']['cost_change_pct'],2)==-11.71 and round(effects['RetailNet']['fill_change_pp'],2)==12.25)
 check('matched_readers_equal_in_run_tables',all_equal)
 check('actual_matched_agent_violations_zero',m.hard_violations.sum()==0 and f.hard_violations.sum()==0 and primary.hard_violations.sum()==0)
 plotted=register['plotted_values']
 check('figure_effects_equal_recomputed_means',np.isclose(plotted['m5_repair']['cost_change_pct'],effects['M5']['cost_change_pct']) and np.isclose(plotted['retailnet_repair']['fill_change_pp'],effects['RetailNet']['fill_change_pp']))
 paired=pd.read_csv(ROOT/'study_results/freshretailnet/FreshRetailNet_exploratory_repair_paired_seed_effects.csv');q=paired[paired.scenario.eq('derived_field_collapse')]
 check('adverse_seeds_preserved',len(q)==30 and int((q.repair_minus_no_repair_cost<0).sum())==26 and int((q.repair_minus_no_repair_fill_pp>0).sum())==30)
 forecast=js(ROOT/'results/freshretailnet/forecast/summary.json');metrics={r['method']:r['wape']*100 for r in forecast['eval_metrics'] if r['target_kind']=='fully_in_stock_daily_sales'}
 check('forecast_values_recomputed',round(metrics['lightgbm_raw'],2)==38.84 and round(metrics['seasonal_naive_raw'],2)==42.84 and metrics['lightgbm_raw']<metrics['lightgbm_gbm_recovered'])
 # Every main-body author/year citation has a corresponding reference entry.
 refs=[p.text for p in doc.paragraphs if p.style.name=='Reference'];unmatched=[]
 for found in re.finditer(r'\(([^()]*\d{4}[^()]*)\)',text):
  for citation in found[1].split(';'):
   year=re.search(r'\b(19\d{2}|20\d{2})\b',citation)
   if year:
    author=citation.strip().split(',')[0].split(' & ')[0].replace(' et al.','').strip()
    if author==year[1]:
     context=text[max(0,found.start()-120):found.start()]
     matches=any(r.split(',')[0] in context and year[1] in r for r in refs)
    else:matches=any(r.startswith(author) and year[1] in r for r in refs)
    if not matches:unmatched.append(citation)
 check('main_author_year_citations_have_references',not unmatched)
 layout=[];sparse=[];images=[];pdftext=[]
 for n,page in enumerate(pdf,1):
  content=page.get_text();pdftext.append(content);shown=page.get_image_info()
  bodywords=[w for w in page.get_text('words') if w[1]>=55 and w[3]<=page.rect.height-55]
  if not bodywords and not shown:layout.append({'page':n,'issue':'Empty page body'})
  for w in page.get_text('words'):
   if w[0]<-1 or w[1]<-1 or w[2]>page.rect.width+1 or w[3]>page.rect.height+1:layout.append({'page':n,'issue':'Text outside page','text':w[4]})
  if len(bodywords)<45 and not shown and n>2:sparse.append(n)
  if shown:images.append(n)
 check('all_pdf_pages_in_bounds_and_nonblank',not layout)
 check('all_22_figure_captions_rendered',all(f'Figure {i}.' in '\n'.join(pdftext) for i in range(1,23)))
 main_start=actual_pages['1. Retail replenishment: why the decision matters'];main_end=actual_pages['References']-1
 receipt={'status':'VERIFIED' if not issues else 'FAILED','checked_at_utc':datetime.now(timezone.utc).isoformat(),'checks':checks,'issues':issues,'unmatched_citations':unmatched,'docx_sha256':sha(DOCX),'pdf_sha256':sha(PDF),'main_prose_words':source_word_count,'main_prose_whitespace_words':sum(len(p.split()) for p in actual_prose),'chapter_words':manifest['chapter_words'],'total_pdf_pages':len(pdf),'main_chapter_pages':main_end-main_start+1,'main_page_range':[main_start,main_end],'embedded_figures':len(embedded),'tables':len(doc.tables),'references':len(refs),'headline_effects':effects,'layout_issues':layout,'sparse_pages_for_visual_review':sparse,'image_pages':images}
 (OUT/'validation_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n');(OUT/'rendered_text.txt').write_text('\n'.join(pdftext))
 if not issues:
  manifest['status']='VERIFIED';manifest['validated_at_utc']=receipt['checked_at_utc'];manifest['docx_sha256']=receipt['docx_sha256'];manifest['pdf_sha256']=receipt['pdf_sha256'];manifest['total_pdf_pages']=len(pdf);manifest['main_chapter_pages']=receipt['main_chapter_pages'];(OUT/'build_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
 print(json.dumps({k:receipt[k] for k in ['status','issues','main_prose_words','main_prose_whitespace_words','total_pdf_pages','main_chapter_pages','embedded_figures','sparse_pages_for_visual_review']},indent=2))
 if issues:raise SystemExit(1)

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--update-contents',action='store_true');validate(p.parse_args().update_contents)
