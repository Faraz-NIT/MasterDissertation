"""Build the complete new dissertation from newly drafted chapters and clean-study evidence."""
from __future__ import annotations
import hashlib,json,re
from datetime import datetime,timezone
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor
import pandas as pd

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'study_results/ollama_dissertation_20261012'
TARGET=OUT/'LLM_Agents_Replenishment_Dissertation.docx'
INK='20343D';BLUE='24548A';GREEN='27775D'

EXECUTIVE=json.loads((OUT/'executive_summary.json').read_text())

def sha(p):
 with Path(p).open('rb') as s:return hashlib.file_digest(s,'sha256').hexdigest()
def read(p):return json.loads(Path(p).read_text())
def main_words(text):
 prose='\n'.join(x for x in text.splitlines() if x.strip() and not x.startswith('#') and not x.startswith('[['))
 return len(re.findall(r"\b[\w’'-]+\b",prose))
def field(p,instruction):
 e=OxmlElement('w:fldSimple');e.set(qn('w:instr'),instruction);p._p.append(e)
def link(p,text,target,internal=False):
 e=OxmlElement('w:hyperlink')
 if internal:e.set(qn('w:anchor'),target)
 else:
  rel=p.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);e.set(qn('r:id'),rel)
 r=OxmlElement('w:r');props=OxmlElement('w:rPr')
 color=OxmlElement('w:color');color.set(qn('w:val'),BLUE);props.append(color)
 fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');props.append(fonts)
 r.append(props);t=OxmlElement('w:t');t.text=text;r.append(t);e.append(r);p._p.append(e)
def add_text(p,text):
 pos=0
 for m in re.finditer(r'https?://\S+',text):
  p.add_run(text[pos:m.start()]);url=m.group().rstrip('.;,');link(p,url,url);p.add_run(m.group()[len(url):]);pos=m.end()
 p.add_run(text[pos:])

from evidence import get_refs,tables

def build():
 body='\n\n'.join(p.read_text().strip() for p in sorted((OUT/'chapters').glob('*.md')))+'\n'
 (OUT/'manuscript.md').write_text(body)
 appendix=(OUT/'appendices.md').read_text();assert main_words(body)>=15000
 register=read(OUT/'figure_register.json');figs={x['key']:x for x in register['figures']};assert len(figs)==21
 doc=Document();section=doc.sections[0]
 section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
 section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
 section.header_distance=section.footer_distance=Inches(.45);section.different_first_page_header_footer=True
 for s in doc.styles:
  if s.type in [1,2]:
   s.font.name='Times New Roman'
   if s.element.rPr is not None:
    fonts=s.element.rPr.find(qn('w:rFonts'))
    if fonts is not None:
     fonts.set(qn('w:eastAsia'),'Times New Roman');fonts.set(qn('w:cs'),'Times New Roman')
 normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 normal.paragraph_format.widow_control=True
 for name,size in [('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
  s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor.from_string(INK);s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.15
  if name=='Heading 1':s.paragraph_format.page_break_before=True
 for name,size in [('Caption',10.5),('Compact',11),('Reference',12),('Contents Entry',11)]:
  if name not in doc.styles:doc.styles.add_style(name,1)
  s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.paragraph_format.line_spacing=1.1;s.paragraph_format.space_after=Pt(5)
  s.font.bold=False;s.font.italic=False;s.font.color.rgb=RGBColor.from_string(INK)
  s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
 # Remove theme font aliases so office renderers honour the explicit family.
 for fonts in doc.styles.element.iter(qn('w:rFonts')):
  for attr in list(fonts.attrib):
   if 'theme' in attr.lower():del fonts.attrib[attr]
 doc.styles['Reference'].paragraph_format.left_indent=Inches(.22);doc.styles['Reference'].paragraph_format.first_line_indent=Inches(-.22)
 header=section.header.paragraphs[0];header.text='LLM AGENTS FOR RETAIL REPLENISHMENT | M5 & RETAILNET';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 for r in header.runs:r.font.size=Pt(9)
 footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('Page ');field(footer,'PAGE');footer.add_run(' of ');field(footer,'NUMPAGES')
 for r in footer.runs:r.font.size=Pt(10)
 title=doc.add_paragraph();title.paragraph_format.space_before=Pt(100);title.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=title.add_run('LLM Agents for\nRetail Replenishment');r.bold=True;r.font.size=Pt(26);r.font.color.rgb=RGBColor.from_string(INK)
 p=doc.add_paragraph('Evidence from M5 and FreshRetailNet-50K');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(18);p.runs[0].font.size=Pt(17)
 p=doc.add_paragraph('A local Ollama study of service, supplier information and planning');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(12)
 p=doc.add_paragraph('Master’s dissertation\nemlyon business school\nOctober 2026');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(65)
 doc.add_heading('Executive summary',level=1)
 for text in EXECUTIVE:doc.add_paragraph(text)
 # Contents are clickable and filled with checked PDF page numbers after first render.
 records=[]
 for text in [body,appendix]:
  for line in text.splitlines():
   if re.match(r'^#{1,2} ',line):records.append({'title':line.split(' ',1)[1],'level':len(line.split(' ',1)[0]),'bookmark':f'h{len(records):03d}'})
 body_n=sum(1 for line in body.splitlines() if re.match(r'^#{1,2} ',line))
 records.insert(body_n,{'title':'References','level':1,'bookmark':'references'})
 record_by_title={r['title']:r for r in records}
 pages=read(OUT/'contents_pages.json') if (OUT/'contents_pages.json').exists() else {}
 doc.add_heading('Contents',level=1)
 for r in records:
  p=doc.add_paragraph(style='Contents Entry');p.paragraph_format.tab_stops.add_tab_stop(Inches(6.22),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
  if r['level']==2:p.paragraph_format.left_indent=Inches(.18)
  link(p,r['title'],r['bookmark'],internal=True);p.add_run('\t'+str(pages.get(r['title'],'—')))
 doc.add_heading('List of figures',level=1)
 ordered_keys=re.findall(r'\[\[FIGURE:(\w+)\]\]',body+appendix)
 for i,key in enumerate(ordered_keys,1):
  f=figs[key];p=doc.add_paragraph(style='Compact');link(p,f'Figure {i}. {f["dataset"]}: {f["title"]}',f'figure{i}',internal=True)
 figure_number=0;table_number=0;heading_id=1;inserted=[]
 table_defs=tables()
 def bookmark(p,name):
  nonlocal heading_id
  start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(heading_id));start.set(qn('w:name'),name)
  end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(heading_id));p._p.insert(0,start);p._p.append(end);heading_id+=1
 def figure(key):
  nonlocal figure_number
  f=figs[key];assert sha(OUT/f['path'])==f['sha256'];figure_number+=1
  p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True;run=p.add_run(f'Figure {figure_number}. {f["dataset"]}: {f["title"]}');run.bold=True;run.font.color.rgb=RGBColor.from_string(GREEN if f['dataset']=='RetailNet' else BLUE if f['dataset']=='M5' else INK);bookmark(p,f'figure{figure_number}')
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(str(OUT/f['path']),width=Inches(6.22))
  p=doc.add_paragraph('Source and scope: '+f['caption'],style='Caption');p.paragraph_format.keep_together=True
  inserted.append({'number':figure_number,**f})
 def table(key):
  nonlocal table_number
  title,headers,rows=table_defs[key];table_number+=1
  p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True;p.add_run(f'Table {table_number}. {title}').bold=True
  t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=True
  for i,v in enumerate(headers):t.rows[0].cells[i].text=v
  prop=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(prop)
  for row in rows:
   cells=t.add_row().cells
   for cell,val in zip(cells,row):cell.text=str(val)
  for j,row in enumerate(t.rows):
   prop=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(prop)
   for cell in row.cells:
    for p in cell.paragraphs:
     p.style='Compact';p.paragraph_format.space_after=Pt(3)
     for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(11);run.bold=(j==0)
    if j==0:
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF2');cell._tc.get_or_add_tcPr().append(shade)
  doc.add_paragraph().paragraph_format.space_after=Pt(2)
 def manuscript(text):
  for line in text.splitlines():
   if not line.strip():continue
   h=re.match(r'^(#{1,3}) (.*)',line)
   if h:
    p=doc.add_heading(h[2],level=len(h[1]));r=record_by_title.get(h[2])
    if r:bookmark(p,r['bookmark'])
   elif line.startswith('[[FIGURE:'):figure(line[9:-2])
   elif line.startswith('[[TABLE:'):table(line[8:-2])
   else:add_text(doc.add_paragraph(),line)
 manuscript(body)
 p=doc.add_heading('References',level=1);bookmark(p,'references')
 refs=get_refs()
 for text in refs:add_text(doc.add_paragraph(style='Reference'),text)
 (OUT/'references.md').write_text('\n\n'.join(refs)+'\n')
 manuscript(appendix)
 doc.core_properties.title='LLM Agents for Retail Replenishment: Evidence from M5 and FreshRetailNet-50K'
 doc.core_properties.subject='New business-school dissertation based on the completed clean Ollama study; concise findings and dataset-labelled graphs'
 doc.core_properties.author='';doc.core_properties.keywords='M5; FreshRetailNet; replenishment; Qwen; Ollama; business'
 doc.core_properties.modified=datetime.now(timezone.utc).replace(tzinfo=None)
 doc.save(TARGET)
 (OUT/'heading_register.json').write_text(json.dumps(records,indent=2)+'\n')
 counts={}
 for section in re.split(r'(?m)^# ',body)[1:]:counts[section.splitlines()[0]]=main_words('\n'.join(section.splitlines()[1:]))
 manifest={'status':'BUILT_PENDING_RENDER','main_prose_words_excluding_front_matter_references_appendices':main_words(body),'chapter_words':counts,'figures':inserted,'tables':table_number,'references':len(refs),'evaluated_study_commit':'4721351c3745fd51db9fa10b86b7e9d9839a821c','format':{'font':'Times New Roman','body_points':12,'line_spacing':1.5,'margins_inches':1,'alignment':'justified'},'docx_sha256':sha(TARGET),'inputs':register['inputs'],'manuscript_sha256':sha(OUT/'manuscript.md'),'appendices_sha256':sha(OUT/'appendices.md'),'figure_register_sha256':sha(OUT/'figure_register.json'),'builder_sha256':sha(__file__),'contents_has_verified_pages':bool(pages),'built_at_utc':datetime.now(timezone.utc).isoformat()}
 (OUT/'build_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
 print(json.dumps({k:manifest[k] for k in ['status','main_prose_words_excluding_front_matter_references_appendices','chapter_words','tables','references']},indent=2))

if __name__=='__main__':build()
