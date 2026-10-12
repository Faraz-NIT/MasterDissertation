"""Validate the complete rewrite against Word formatting, PDF layout and saved evidence."""
from __future__ import annotations
import argparse
import json
import re
import zipfile
from datetime import datetime,timezone
import fitz
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
from config import ROOT,OUT,STUDY,DOCX,PDF,sha,read,prose,main_words


def clean(text):return re.sub(r'\s+',' ',text).strip()


def contents(pdf):
    headings=read(OUT/'heading_register.json');found={}
    for n,page in enumerate(pdf,1):
        spans=[span['text'] for block in page.get_text('dict')['blocks'] for line in block.get('lines',[]) for span in line['spans'] if span['size']>=12.8]
        text=clean(' '.join(spans))
        for heading in headings:
            if clean(heading['title']) in text:found.setdefault(heading['title'],n)
    missing=[h['title'] for h in headings if h['title'] not in found]
    if missing:raise RuntimeError('Unmapped headings: '+str(missing))
    return found


def estimate(values):
    values=np.asarray(values,dtype=float);random=np.random.default_rng(20261009)
    draws=values[random.integers(0,len(values),size=(10000,len(values)))].mean(axis=1)
    return {'mean':float(values.mean()),'ci_low':float(np.quantile(draws,.025)),'ci_high':float(np.quantile(draws,.975))}


def validate(update=False):
    pdf=fitz.open(PDF);pages=contents(pdf)
    if update:
        (OUT/'contents_pages.json').write_text(json.dumps(pages,indent=2)+'\n')
        print(json.dumps({'status':'MAPPED','headings':len(pages),'pdf_pages':len(pdf)}));return
    issues=[];checks={}
    def check(name,value):
        checks[name]=bool(value)
        if not value:issues.append(name)
    doc=Document(DOCX);manifest=read(OUT/'build_manifest.json');register=read(OUT/'figure_register.json');brief=read(OUT/'rewrite_brief.json')
    body=(OUT/'manuscript.md').read_text();expected_prose=prose(body);actual_prose=[];active=False
    for paragraph in doc.paragraphs:
        if paragraph.style.name=='Heading 1':
            if paragraph.text.startswith('1. '):active=True
            if paragraph.text=='References':active=False
        if active and paragraph.style.name=='Normal' and paragraph.text.strip():actual_prose.append(paragraph.text)
    whitespace=sum(len(p.split()) for p in actual_prose)
    check('main_prose_minimum_15000',main_words(body)>=15000 and whitespace>=15000)
    check('entire_new_manuscript_present_in_word',actual_prose==expected_prose)
    check('source_word_count_matches_manifest',main_words(body)==manifest['main_prose_words_excluding_front_matter_references_appendices'])
    style=doc.styles['Normal']
    check('times_new_roman_12pt_1_5_justified',style.font.name=='Times New Roman' and style.font.size.pt==12 and style.paragraph_format.line_spacing==1.5 and style.paragraph_format.alignment==WD_ALIGN_PARAGRAPH.JUSTIFY)
    check('one_inch_margins',all(getattr(s,name).twips==1440 for s in doc.sections for name in ['top_margin','bottom_margin','left_margin','right_margin']))
    check('no_theme_font_override',not any('theme' in attr.lower() for fonts in doc.styles.element.iter(qn('w:rFonts')) for attr in fonts.attrib))
    chapters=[p for p in doc.paragraphs if p.style.name=='Heading 1' and re.match(r'^[1-7]\. ',p.text)]
    check('seven_new_main_chapters',len(chapters)==7)
    counts=manifest['chapter_words']
    check('concise_results',next(v for k,v in counts.items() if k.startswith('5.'))<1500)
    check('concise_conclusion',next(v for k,v in counts.items() if k.startswith('7.'))<850)
    check('trained_model_and_inference_distinguished',"not trained or fine-tuned on M5 or FreshRetailNet" in body and '32 GiB' in body and 'No GPU was used' in body and 'one numerical fitting thread' in body)
    check('new_study_only',all(term not in body for term in ['10.93%','11.71%','23,331.29','ega-qwen2.5:1.5b-v2']))
    check('structured_comparator_and_limits_explicit','same numerical planner' in body and 'slightly better in RetailNet' in body and 'previously unseen research holdouts' in body and 'No participant study was conducted' in body)
    check('contents_pages_match_current_pdf',read(OUT/'contents_pages.json')==pages)
    with zipfile.ZipFile(DOCX) as archive:
        embedded=[__import__('hashlib').sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith('word/media/')]
        xml=etree.fromstring(archive.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        bookmarks=set(xml.xpath('//w:bookmarkStart/@w:name',namespaces=ns));anchors=xml.xpath('//w:hyperlink/@w:anchor',namespaces=ns)
    check('contents_and_figure_links_resolve',set(anchors).issubset(bookmarks))
    check('all_21_figures_exactly_embedded',len(doc.inline_shapes)==21 and len(embedded)==21 and all(f['sha256'] in embedded and sha(OUT/f['path'])==f['sha256'] for f in register['figures']))
    check('dataset_label_on_every_figure',all(f['dataset'] in ['M5','RetailNet','M5 + RetailNet'] for f in register['figures']))
    check('ten_explanatory_tables',len(doc.tables)==10)
    check('evidence_inputs_unchanged',all(sha(ROOT/path)==entry['sha256'] for path,entry in register['inputs'].items()))
    check('completed_study_outputs_unchanged',all(sha(ROOT/path)==digest for path,digest in brief['protected_study_files'].items()))
    check('exact_gilbert_structural_reference',brief['structural_reference']['sha256']=='66540dc875e1861d2adb8c716f2e717afaff26723972a6c20aed08e84d7eb584')
    # Independent numerical derivation, rather than trusting the document's table builder.
    seeds=pd.read_csv(STUDY/'analysis/seed_results.csv');overall=seeds.groupby(['dataset','arm','seed'])[['fill_rate','cost_per_demand','expiry_share']].mean().reset_index()
    means=overall.groupby(['dataset','arm'])[['fill_rate','cost_per_demand','expiry_share']].mean();recomputed=[]
    source_effects=read(STUDY/'analysis/overall_effects.json')
    for dataset in ['M5','RetailNet']:
        group=overall[overall.dataset==dataset];agent=group[group.arm=='llm_agents'].set_index('seed')
        for comparison in ['conventional','structured_planner']:
            baseline=group[group.arm==comparison].set_index('seed');ids=agent.index.intersection(baseline.index)
            service=estimate(100*(agent.loc[ids,'fill_rate']-baseline.loc[ids,'fill_rate']))
            cost=estimate(100*(agent.loc[ids,'cost_per_demand']/baseline.loc[ids,'cost_per_demand']-1))
            original=next(r for r in source_effects if r['dataset']==dataset and r['comparison']==comparison)
            check('paired_effect_recalculated_'+dataset+'_'+comparison,all(np.isclose(service[k],original['fill_effect_pp'][k],atol=1e-10) and np.isclose(cost[k],original['cost_effect_pct'][k],atol=1e-10) for k in service))
            recomputed.append({'dataset':dataset,'comparison':comparison,'service':service,'cost':cost})
    daily=pd.read_csv(STUDY/'analysis/daily_results.csv');agents=daily[daily.arm=='llm_agents'];priority=[]
    for dataset in ['M5','RetailNet']:
        g=agents[agents.dataset==dataset];active=g[(g.scenario=='promotion')&(g.day>=2)]
        priority.append({'dataset':dataset,'all_days':len(g),'all_correct':int(g.priority_correct.sum()),'active_days':len(active),'active_correct':int(active.priority_correct.sum())})
    figure=next(f for f in register['figures'] if f['key']=='priority_accuracy')
    check('priority_figure_recalculated_from_decisions',priority==figure['plotted_values'])
    check('complete_extraction_measured_on_all_agent_days',len(agents)==2520 and agents.source_exact.eq(1).all() and agents.source_correct_slots.eq(13).all())
    check('zero_actual_holds_and_committed_violations',daily.held.sum()==0 and daily.actual_violations.sum()==0)
    audit=read(STUDY/'analysis/audit_receipt.json');execution=read(STUDY/'execution_receipt.json')
    check('completed_registered_calls_and_grid',execution['status']=='COMPLETE' and execution['runs']==1080 and execution['fresh_primary_model_calls']==7560 and audit['status']=='VERIFIED' and audit['unique_model_response_timestamps']==7560)
    refs=[p.text for p in doc.paragraphs if p.style.name=='Reference'];unmatched=[]
    for found in re.finditer(r'\(([^()]*\d{4}[^()]*)\)',body):
        for citation in found[1].split(';'):
            year=re.search(r'\b(19\d{2}|20\d{2})\b',citation)
            if not year:continue
            author=citation.strip().split(',')[0].split(' & ')[0].replace(' et al.','').strip()
            if author==year[1]:
                context=body[max(0,found.start()-150):found.start()].casefold()
                matched=any(r.split(',')[0].casefold() in context and year[1] in r for r in refs)
            else:matched=any(r.casefold().startswith(author.casefold()) and year[1] in r for r in refs)
            if not matched:unmatched.append(citation)
    check('main_author_year_citations_have_references',not unmatched)
    layout=[];sparse=[];image_pages=[];texts=[]
    for n,page in enumerate(pdf,1):
        texts.append(page.get_text());images=page.get_image_info();bodywords=[w for w in page.get_text('words') if w[1]>=55 and w[3]<=page.rect.height-55]
        if not bodywords and not images:layout.append({'page':n,'issue':'Empty page body'})
        for word in page.get_text('words'):
            if word[0]<-1 or word[1]<-1 or word[2]>page.rect.width+1 or word[3]>page.rect.height+1:layout.append({'page':n,'issue':'Out-of-page text','text':word[4]})
        if len(bodywords)<45 and not images and n>2:sparse.append(n)
        if images:image_pages.append(n)
    check('all_pdf_pages_nonblank_and_within_bounds',not layout)
    pdftext='\n'.join(texts)
    check('all_figure_captions_present_in_pdf',all(f'Figure {n}.' in pdftext for n in range(1,22)))
    check('headline_values_present_in_pdf',all(value in pdftext for value in ['65.05%','90.33%','83.53%','84.31%','70.34%','4.52%','98.97%','87.46%']))
    first=pages['1. Retail replenishment and the business case for change'];last=pages['References']-1
    receipt={'status':'VERIFIED' if not issues else 'FAILED','checked_at_utc':datetime.now(timezone.utc).isoformat(),'checks':checks,'issues':issues,
        'main_prose_words':main_words(body),'main_prose_whitespace_words':whitespace,'chapter_words':counts,'pdf_pages':len(pdf),'main_chapter_pages':last-first+1,'main_page_range':[first,last],
        'embedded_figures':len(embedded),'tables':len(doc.tables),'references':len(refs),'docx_sha256':sha(DOCX),'pdf_sha256':sha(PDF),
        'recomputed_effects':recomputed,'priority_counts':priority,'unmatched_citations':unmatched,'layout_issues':layout,'sparse_pages_for_visual_review':sparse,'image_pages':image_pages}
    (OUT/'validation_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n');(OUT/'rendered_text.txt').write_text(pdftext)
    if not issues:
        manifest.update(status='VERIFIED',docx_sha256=receipt['docx_sha256'],pdf_sha256=receipt['pdf_sha256'],pdf_pages=len(pdf),main_chapter_pages=receipt['main_chapter_pages'])
        (OUT/'build_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(json.dumps({k:receipt[k] for k in ['status','issues','main_prose_words','main_prose_whitespace_words','pdf_pages','main_chapter_pages','embedded_figures','sparse_pages_for_visual_review']},indent=2))
    if issues:raise SystemExit(1)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--update-contents',action='store_true');validate(parser.parse_args().update_contents)
