"""Assemble appendices only, using archived evidence without rerunning experiments."""
from __future__ import annotations
import hashlib
import importlib.util
import json
import re
import sys
from datetime import datetime,timezone
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'study_results/full_appendices_20261017'
BOOK=ROOT/'study_results/ollama_dissertation_20261012'
MAIN=ROOT/'study_results/ollama_clean_study_20261009'
ARCH=ROOT/'study_results/architecture_study_20261013'
CHAIN=ROOT/'study_results/langchain_extension_20261014'
DOCX=OUT/'Full_Dissertation_Appendices.docx'
USED=set()


def sha(path):
    with Path(path).open('rb') as stream:return hashlib.file_digest(stream,'sha256').hexdigest()


def read(path):
    USED.add(Path(path));return json.loads(Path(path).read_text())


def csv(path):
    USED.add(Path(path));return pd.read_csv(path)


def interval(item,precision=3):
    return f"{item['mean']:+.{precision}f} [{item['ci_low']:+.{precision}f}, {item['ci_high']:+.{precision}f}]"


def assemble():
    sys.path.insert(0,str(ROOT/'scripts/ollama_dissertation'))
    from evidence import tables as original_tables
    definitions={key:{'title':title,'headers':headers,'rows':rows}
                 for key,(title,headers,rows) in original_tables().items()}
    protocols={name:read(path/'protocol.json') for name,path in [('main',MAIN),('architecture',ARCH),('langchain',CHAIN)]}
    receipts={name:read(path/'execution_receipt.json') for name,path in [('main',MAIN),('architecture',ARCH),('langchain',CHAIN)]}
    for name,path in [('main',MAIN),('architecture',ARCH),('langchain',CHAIN)]:
        assert read(path/'analysis/audit_receipt.json')['status']=='VERIFIED'
        for rel,digest in protocols[name]['code_sha256'].items():assert sha(ROOT/rel)==digest
        for dataset,info in protocols[name]['datasets'].items():
            for rel,digest in info['file_sha256'].items():assert sha(path/'data'/dataset/rel)==digest
    base=protocols['main'];arch=read(ARCH/'analysis/analysis_summary.json')
    fresh=read(CHAIN/'analysis/validation_summary.json')
    replay=read(CHAIN/'development/recorded_response_replay.json')
    hypotheses=read(ROOT/'study_results/conclusion_revision_20261013/evidence_map.json')
    for rel in ['analysis/overall_effects.json','analysis/shared_forecast_context.csv','analysis/seed_results.csv',
                'analysis/policy_means.csv','analysis/llm_diagnostics.csv','data/preparation_receipt.json']:
        USED.add(MAIN/rel)
    figures={f['key']:{**f,'absolute_path':BOOK/f['path']} for f in read(BOOK/'figure_register.json')['figures']}
    for f in figures.values():assert sha(f['absolute_path'])==f['sha256']
    text=(BOOK/'appendices.md').read_text();USED.add(BOOK/'appendices.md')
    text=text.replace('The source of these settings is the frozen protocol of the completed October 2026 study, together with its training and execution receipts.',
                      'The source of these settings is the frozen protocol of the completed main study, together with its forecasting and execution receipts. Appendices E–I extend the supporting material to the architecture comparison, LangChain validation, prompts, sample catalogues and a separately labelled service-target illustration.')
    text=text.replace('[[TABLE:technical_settings]]','[[TABLE:study_versions]]\n\n[[TABLE:technical_settings]]\n\n[[TABLE:model_environment]]',1)
    definitions['model_environment']['title']='Main-study model and computing environment; shared inference hardware retained for the extensions'
    text=text.replace('## A.1 Measures and aggregation','''## A.1 Inventory, data and simulation assumptions

The final study uses a consistent simulated inventory ledger. It does not inject inaccurate stock records, phantom inventory, unrecorded losses or delayed stock updates, and it does not test inventory-record repairs. Demand and supply uncertainty change the physical outcomes; the recorded stock follows the simulated physical balance. Source verification checks supplier terms against current authenticated fields and does not establish an ability to correct false inventory counts.

Starting stock is the rounded-up equivalent of three average days of sales from the past 56 days, with no opening pipeline. RetailNet opening stock is distributed across three initial age batches. New deliveries receive a hypothetical three-day life, and older usable batches are consumed first. Remaining units expire after the day's sales opportunity at the end of their usable life. M5 has no expiry within the weekly horizon.

Each day begins with due deliveries. The policy then reads current stock, forecasts and supplier terms, proposes an order and passes the release checks. Accepted quantities enter the delivery pipeline. Sales consume available stock, and unmet demand is lost rather than backlogged. Purchase spend records accepted quantities at the assumed unit cost of one; cancelled quantities add no stock or pipeline. Opening stock plus accepted purchases is reconciled against sales, expiry, closing stock and closing pipeline.

[[TABLE:simulation_settings]]

Released sales are a proxy for demand. The simulation multiplies them by a mean-one lognormal draw with log-scale standard deviation 0.10. This produces random variation and is not a fixed plus-or-minus 10% bound. Announced promotions multiply the demand proxy and available forecast by 1.3 from simulation day index 2, the third day. The same demand realisation and supply-opportunity draws are used for matched policies. Supplier disruption doubles carton sizes and minimum orders, adds one delivery day and reduces dispatch capacity from the third day, subject to the implemented minimum feasible capacity.

The LightGBM model uses 150 trees, maximum depth five, at most 20 leaves and one fitting thread. Each dataset and weekly origin has its own fit on earlier observations, using at most 365 days and a recursive 14-day forecast. Historical weekly sales changes supply the uncertainty estimate; test outcomes do not supply forecast inputs. Four primary fits serve every policy in the matched experiments.

## A.2 Measures and aggregation''',1)
    text=text.replace('## A.2 Shared forecast context','## A.3 Shared forecast context',1)
    text=text.replace('[[TABLE:evidence_locations]]','[[TABLE:current_evidence_locations]]',1)
    text=text.replace('The new dissertation package contains the Word and PDF documents, all figures, editable chapter sources, bibliography, build helpers and validation records. It is a document delivery package. The much larger complete experiment package remains separately linked in the repository. The supplied Gilbert thesis is not redistributed, and its content is not used as empirical evidence for retail replenishment. The rewrite brief records its title, author, page count and source hash solely to identify the structural reference.',
                      'The accompanying repository-ready code package contains only the relevant experiment code, tests, prepared panels, fitted forecasts, saved results, plots and logs. It contains 17,100 primary native model-call records across the three empirical stages. The full source-data releases, local model weights and environments are not embedded. The supplied Gilbert thesis guided presentation and is not redistributed or used as empirical evidence for retail replenishment.')
    text=text.replace('## C.1 Use of AI-assisted tools','''## C.1 Verification and scope

The independent audits verify the complete registered decision grids, frozen source and data hashes, raw model-call hashes, supplier-field and priority scores, released quantities, physical stock balances and the cost calculation. They reconstruct FIFO consumption and supplier deliveries separately from the evaluated simulator. A zero order remains distinct from a held action. The main and architecture studies recorded no released constraint violations or holds under their controlled primary conditions.

The portable code package was checked after extraction into a separate folder: all sixteen relevant tests, all three saved-result audits and all four plot-generation scripts passed. These checks read completed evidence and did not create new experimental model responses. Development records remain separate from primary results, including failed checks and subsequent corrections. Known historical weeks are not unseen holdouts, and thirty simulation seeds are not thirty independently trained language models.

## C.2 Reproduction commands

The following commands run from the root of the extracted code package after installing its pinned dependencies. The audit command writes derived CSV files and audit receipts in that copy. Verify archived bytes before running it. Fresh inference requires a running local model with the recorded digest, a separate output location and a newly frozen protocol; the archived scripts protect completed cases against overwriting.

[[CODE:reproduction]]

Official dataset locations are https://www.kaggle.com/competitions/m5-forecasting-accuracy/data and https://huggingface.co/datasets/Dingdong-Inc/FreshRetailNet-50K. The RetailNet source revision recorded for preparation is 08c1fab7f9257bc73679d415d65d644165d351d4, under the release's CC-BY-4.0 licence. Data-preparation receipts record the full-source file hashes. M5 uses published unit sales; RetailNet uses published normalised sale_amount values multiplied by 100. The final study does not recover latent demand using stockout flags.

## C.3 Use of AI-assisted tools''',1)
    definitions['study_versions']={'title':'Empirical stages and their separate evidence scopes',
        'headers':['Stage','Simulation seeds','Policy runs / portfolio days','Fresh primary calls','Live execution'],
        'rows':[['Main study: three-role agents and two controls','101–130','1,080 / 7,560','7,560','2.66 hours'],
                ['Architecture comparison: LangGraph implementation','201–230','1,080 / 7,560','9,240','4.01 hours'],
                ['LangChain implementation validation','233','36 / 252','300','8.00 minutes']]}
    definitions['simulation_settings']={'title':'Declared simulation assumptions; common across empirical stages',
        'headers':['Element','Setting and interpretation'],'rows':[
        ['Starting stock','ceil(3 × past-56-day mean); zero initial pipeline'],
        ['Synthetic suppliers','Three groups; ten series per group; current daily terms supplied before ordering'],
        ['Routine carton / minimum order','Cartons of 2, 4 and 6; minimum nonzero line twice its carton size'],
        ['Routine lead times','1, 2 and 2 days for the three supplier groups'],
        ['Daily budget and dispatch limits','Training-derived scale 2.8 × historical mean; announced daily variation of up to ±8%'],
        ['Delivery opportunities','2% cancellation; 10% draw for 80% fill; 10% draw for one extra delivery day'],
        ['Promotion','Announced 30% increase in demand proxy and forecast from the third day'],
        ['Supplier disruption','From the third day: carton/MOQ ×2, lead time +1 day, capacity ×0.7, with a minimum feasible limit'],
        ['Stock life','M5: no expiry within the horizon; RetailNet: hypothetical three-day FIFO'],
        ['Declared priorities','Balanced: safety factor 1.28; active promotion: service-first factor 1.65'],
        ['Hard constraints','Nonnegative quantities, carton multiples, minimum line, supplier capacity and daily budget'],
        ['Inventory-record errors','None injected; no inventory repair or sales-recovery mechanism']]}
    definitions['current_evidence_locations']={'title':'Evidence and code locations in the accompanying package',
        'headers':['Item','Location relative to the repository root'],'rows':[
        ['Main experiment','study_results/ollama_clean_study_20261009/'],
        ['Full architecture comparison','study_results/architecture_study_20261013/'],
        ['Fresh LangChain validation','study_results/langchain_extension_20261014/'],
        ['Within each study','protocol.json; data/; runs/; development/; logs/; analysis/; execution_receipt.json'],
        ['Baseline simulation and forecasting code','scripts/ollama_clean_study/'],
        ['Selective-delegation implementation','scripts/architecture_study/'],
        ['LangChain LCEL and StructuredTool implementation','scripts/langchain_extension/'],
        ['Relevant engineering tests','tests/test_ollama_clean_study.py; test_selective_architecture.py; test_langchain_extension.py'],
        ['Portable-package verification','tools/verify_bundle.py; tools/audit_saved_results.py; MANIFEST.json'],
        ['Package-level check records','packaging_checks/; PACKAGE_METADATA.json']]}
    codes={'reproduction':'.venv/bin/python -m pip install -r requirements.txt\n.venv/bin/python tools/verify_bundle.py\n.venv/bin/python -I -m pytest -q tests\n.venv/bin/python tools/audit_saved_results.py'}
    arch_diag={(r['dataset'],r['arm']):r for r in arch['diagnostics']}
    fresh_diag={(r['dataset'],r['arm']):r for r in fresh['diagnostics']}
    effects={(r['dataset'],r['metric']):r for r in arch['paired_effects']
             if r['scenario']=='overall' and r['comparison']=='fixed_agents'}
    definitions['architecture_diagnostics']={'title':'Full architecture comparison: decision reliability and workload',
        'headers':['Dataset / workflow','Calls','Priority accuracy','Active-promotion accuracy','Mean decision time'],'rows':[]}
    definitions['architecture_effects']={'title':'Selective minus fixed architecture: paired effects and 95% intervals',
        'headers':['Dataset','Fill rate (pp)','Cost change (%)','Expiry share (pp)'],'rows':[]}
    for dataset in ['M5','RetailNet']:
        for arm,label in [('fixed_agents','Fixed agent'),('selective_agents','Selective'),('structured_planner','Structured reference')]:
            row=arch_diag[(dataset,arm)]
            definitions['architecture_diagnostics']['rows'].append([dataset+' / '+label,str(row['fresh_model_calls']),
                f"{100*row['priority_accuracy']:.2f}%",f"{100*row['active_promotion_accuracy']:.2f}%",f"{row['mean_decision_seconds']:.3f} s"])
        definitions['architecture_effects']['rows'].append([dataset,interval(effects[(dataset,'fill_rate')]),
            interval(effects[(dataset,'cost_per_demand')]),interval(effects[(dataset,'expiry_share')])])
    definitions['architecture_hypotheses']={'title':'Criteria frozen before the full architecture evaluation',
        'headers':['Criterion','Assessment','M5 / RetailNet'],'rows':[
        ['A1','At least 50% fewer native model calls','Supported / Supported'],
        ['A2','Lower mean decision latency; upper paired difference bound below zero','Supported / Supported'],
        ['A3','Preserved or improved declared-policy accuracy; zero released violations','Supported / Supported'],
        ['A4','Lower fill-change bound > −0.10 pp; upper relative cost-change bound < +0.5%','Supported / Supported'],
        ['Exploratory attribution','Same planner; selective orders match the structured-information reference','100% / 100%']]}
    definitions['langchain_validation']={'title':'Fresh LangChain validation; descriptive one-seed results',
        'headers':['Dataset / workflow','Fill rate','Cost / demand','Priority accuracy','Calls / mean time'],'rows':[]}
    for dataset in ['M5','RetailNet']:
        for arm,label in [('fixed_agents','Fixed agent'),('selective_agents','Selective'),('structured_planner','Structured reference')]:
            row=fresh_diag[(dataset,arm)]
            definitions['langchain_validation']['rows'].append([dataset+' / '+label,f"{100*row['fill_rate']:.2f}%",
                f"{row['cost_per_demand']:.4f}",f"{100*row['priority_accuracy']:.2f}%",
                f"{row['fresh_model_calls']} / {row['mean_decision_seconds']:.3f} s"])
    extra_figures={}
    for source,prefix,keys in [(ARCH,'architecture_',['model_calls','priority_accuracy','inventory_effects','retailnet_expiry']),
                              (CHAIN,'langchain_',['agent_extension','fresh_langchain_validation'])]:
        for figure in read(source/'figure_register.json')['figures']:
            if figure['key'] not in keys:continue
            key=prefix+figure['key'];path=source/figure['path']
            extra_figures[key]={'key':key,'dataset':figure['scope'],'title':figure['title'],'caption':figure['caption'],
                                'absolute_path':path,'sha256':sha(path)}
    figures.update(extra_figures)
    text+='''

LangGraph coordinates steps through a graph with explicit state and conditional routes. LangChain is the library used for the later implementation; its LangChain Expression Language (LCEL) connects runnable steps and branches. A StructuredTool exposes a function with declared input fields. A JSON schema specifies the permitted structure and values of a response.

# Appendix E. Selective architecture and LangChain extension

## E.1 The extension of the earlier agent

The previous agent calls a supplier reader, a coordinator and a critic on each portfolio day. The selective extension retains supplier interpretation, source verification, the numerical planner and hard release checks. It changes delegation: complete declared tables are parsed directly, prose reaches the supplier LLM, and announced business priorities are applied through an explicit route. Failed or unknown conditions hold the order and request LLM review. A model approval cannot override a failed check.

The original fixed agent is rerun as the matched baseline for the architecture comparison. The model digest, prompts, historical panels, forecasts and numerical planner remain constant, while seeds 201–230 supply fresh simulation draws and native model responses. The 1,080 policy runs are a full architecture comparison executed through LangGraph 1.0.10. All three workflows share that framework, which helps isolate the delegation design rather than the mere introduction of a software library.

[[FIGURE:langchain_agent_extension]]

## E.2 Full architecture results

[[TABLE:architecture_diagnostics]]

[[TABLE:architecture_effects]]

Native model calls fall from 7,560 to 1,680 across the two datasets, a 77.78% reduction. Mean whole-decision time falls by 54.79% in M5 and 54.93% in RetailNet. Supplier-field accuracy and declared-priority accuracy are 100% in the selective workflow, with no released rule violations. These measures describe the controlled messages and declared event policies, not independent improvements in the pretrained model's reasoning.

The paired M5 intervals do not establish service or cost superiority. RetailNet shows a fill-rate gain of 0.236 percentage points and a 1.176% reduction in the operating-cost index. Its expiry share rises by 0.081 percentage points, concentrated in promotion. The selective orders match the structured-information reference throughout. The reference receives already prepared facts, and staff effort to prepare those facts is not measured.

[[FIGURE:architecture_model_calls]]

[[FIGURE:architecture_priority_accuracy]]

[[FIGURE:architecture_inventory_effects]]

[[FIGURE:architecture_retailnet_expiry]]

[[TABLE:architecture_hypotheses]]

Architecture criteria A1–A4 were frozen before these main runs. The performance-preservation tolerances are study choices, rather than retailer-approved service targets. Service and cost superiority are assessed separately. These criteria should not be confused with the dissertation's H1–H4 framework, which was articulated when the literature review was rewritten after the main study.

## E.3 LangChain implementation and separate validation

The extension is also implemented using LangChain core 1.6.10. Runnable sequences and conditional branches connect the roles, and a StructuredTool exposes the unchanged numerical planner. The earlier graph functions are reused as ordinary functions; the compiled LangGraph is not invoked by the LangChain chain. Native request formatting, prompts, generation settings, source checks and the release rule are retained. Conversation memory, external tracing, retries and response caches are disabled.

Two checks support that implementation. A recorded-response regression supplies the previously saved model outputs and reconstructs all 7,560 archived portfolio decisions. It reproduces the orders, priorities, role paths and inventory outcomes exactly, with zero new native model calls. This checks software consistency when model outputs are held constant.

Fresh LangChain validation then uses seed 233, both origins, all three scenarios and all thirty series in each dataset. It completes 36 matched policy runs, 252 portfolio days and 300 fresh native model calls. All three workflows use the LangChain chain in this comparison. The validation passes the independent ledger and model-provenance checks. It is one-seed implementation validation; it is not a new thirty-seed LangChain performance study and provides no new bootstrap interval estimate.

[[TABLE:langchain_validation]]

[[FIGURE:langchain_fresh_langchain_validation]]

Sixteen relevant engineering checks cover the original simulator, selective delegation and LangChain implementation. They include table parsing, source mismatches, ambiguous event flags, failed-check holds and the actual planning-tool interface. A complete valid table proceeds without an LLM supplier call; prose delegates to the reader. In development checks, a model approval cannot release an order whose source, route or quantity checks fail. The primary cases did not require unresolved exception handling, and no employee resolution was simulated.

# Appendix F. Agent instructions, message formats and response contracts

The following instructions and in-context examples are copied from the frozen main-study protocol and retained in the architecture and LangChain implementations. Each role uses a separate request and role-specific input. Supplier examples demonstrate field extraction; critic examples demonstrate pass/hold behaviour. The coordinator has no in-context examples. These examples are prompt content, not additional training or evaluation observations.
'''
    roles=[('supplier','Supplier-information agent'),('planner','Replenishment coordinator'),('critic','Independent review role')]
    for index,(role,label) in enumerate(roles,1):
        assert all(protocol['prompts'][role]==base['prompts'][role] for protocol in protocols.values())
        text+=f'\n## F.{index} {label}\n\n[[VERBATIM:prompt_{role}]]\n'
        codes['prompt_'+role]=base['prompts'][role]
        for example_index,example in enumerate(base['few_shot_examples'].get(role,[]),1):
            text+=f'\nExample {example_index}; message role: {example["role"]}.\n\n[[VERBATIM:example_{role}_{example_index}]]\n'
            codes[f'example_{role}_{example_index}']=example['content']
    text+='''

## F.4 Controlled supplier-message formats

The study rotates prose, a reordered memo and a declared table. The three samples below use the same invented agreement to show the formatting only. They are not observed retailer messages or additional experiment outcomes. A budget is appended to every format. The controlled wording supports exact field scoring and limits claims about arbitrary live communications.
'''
    spec=importlib.util.spec_from_file_location('appendix_source_core',ROOT/'scripts/ollama_clean_study/core.py')
    core=importlib.util.module_from_spec(spec);spec.loader.exec_module(core)
    USED.add(ROOT/'scripts/ollama_clean_study/core.py')
    rules={'rules':[{'supplier':'S1','pack':2,'moq':4,'lead_days':1,'capacity':100},
                    {'supplier':'S2','pack':4,'moq':8,'lead_days':2,'capacity':120},
                    {'supplier':'S3','pack':6,'moq':12,'lead_days':2,'capacity':150}],'budget':250}
    for style,label in enumerate(['Prose format','Reordered memo format','Declared table format']):
        text+=f'\n{label}.\n\n[[VERBATIM:message_{style}]]\n';codes['message_'+str(style)]=core.supplier_message(rules,0,style)
    text+='''

## F.5 Complete response schemas and verification

The native interface requests JSON conforming to the following role-specific schemas. A schema-valid response is still checked for correct source fields and supplier coverage. For supplier extraction, thirteen numerical slots are scored: four terms for each of three suppliers, plus the budget. Duplicate or missing suppliers fail exact verification even if the JSON parses. The coordinator must request the declared numerical tool and select one of the two named priorities. The critic must return approve or hold.
'''
    for role,label in roles:
        text+=f'\n{label}: JSON schema.\n\n[[CODE:schema_{role}]]\n'
        codes['schema_'+role]=json.dumps(core.SCHEMAS[role],indent=2,ensure_ascii=False)
    text+='''

The hard release condition requires verified current source fields, valid routing, solver feasibility and passing action checks. Failure holds the action; the system does not substitute a comparator's order. A permitted zero purchase is not a held decision. In the selective workflow, a recognised but malformed table or an unknown promotion flag also prevents automatic release. The deterministic checks retain authority even when the review model returns approve.

# Appendix G. Dataset windows and complete selected-series catalogues

## G.1 Data sources, selection and evaluation windows

Thirty store–product series are selected from each dataset using training-only volume strata and deterministic hash ordering, ten per stratum. M5 selection uses eligible FOODS series with selection-period average sales from 0.5 to 30. RetailNet selection uses series complete across the first 62 training dates and a selection-period normalised average from 0.02 to 1. The same selected identifiers are retained in both test windows and all three empirical stages.

[[TABLE:dataset_windows]]

RetailNet's earlier window is a rolling holdout within the published training release; the later window is the official evaluation release. M5 windows are taken from the end of the evaluation-history sales file. The selected histories were previously inspected, which limits claims of unseen-period confirmation. Stock, supplier groups, ordering terms and costs are constructed for the simulation rather than recovered from these sales files.

The catalogue means below are the values used for series selection. They differ from the past-56-day averages used to initialise stock. RetailNet catalogue means use published normalised sale_amount before multiplication by 100. Strata are coded 0, 1 and 2; the labels identify the selection bands and do not represent product categories.

## G.2 M5: all thirty selected series

[[TABLE:m5_catalogue]]

## G.3 RetailNet: all thirty selected series

[[TABLE:retailnet_catalogue]]
'''
    definitions['dataset_windows']={'title':'Selected-panel windows and quantity conventions',
        'headers':['Dataset','Series','Origin / dates','Quantity convention'],'rows':[]}
    for dataset in ['M5','RetailNet']:
        for origin in base['datasets'][dataset]['origins']:
            dates=base['datasets'][dataset]['evaluation_dates'][str(origin)]
            definitions['dataset_windows']['rows'].append([dataset,'30',f'{origin}: {dates[0]} to {dates[-1]}',
                'Published unit sales' if dataset=='M5' else 'Published normalised sales × 100'])
        cats=[read(MAIN/f'data/{dataset}/origin_{origin}/catalog.json') for origin in base['datasets'][dataset]['origins']]
        assert len(cats[0]['series_ids'])==30 and cats[0]['series_ids']==cats[1]['series_ids']
        for other in [ARCH,CHAIN]:
            for origin in base['datasets'][dataset]['origins']:
                assert read(other/f'data/{dataset}/origin_{origin}/catalog.json')['series_ids']==cats[0]['series_ids']
        key='m5_catalogue' if dataset=='M5' else 'retailnet_catalogue'
        definitions[key]={'title':dataset+': complete training-selected portfolio',
            'headers':['Series identifier','Store','Stratum','Selection-period mean'],
            'rows':[[row['series_id'],str(row['store_id']),str(row['stratum']),f"{row['mean']:.6f}"] for row in cats[0]['metadata']]}
    definitions['research_alignment']={'title':'Dissertation research questions and main-study hypothesis assessments',
        'headers':['Question / hypothesis','Comparison and meaning','Assessment'],'rows':[
        ['RQ1; H1a / H1b','Complete three-role workflow versus conventional ordering: fill rate / operating cost','Supported for both datasets within the simulated panels'],
        ['RQ2; H2a / H2b','LLM increment versus the same numerical planner with structured inputs','Not supported: M5 superiority inconclusive; RetailNet slightly adverse'],
        ['RQ3; H3','Supplier-fact extraction versus business-priority selection','Extraction more accurate in the controlled main-study tasks'],
        ['RQ4; H4','Difference in benefit between the two retail settings','Descriptive/exploratory; formal cross-dataset effect test not performed']]}
    definitions['source_identifiers']={'title':'Frozen protocols and archived evidence identifiers',
        'headers':['Stage','Protocol SHA-256'],'rows':[[name,sha(path/'protocol.json')]
            for name,path in [('Main study',MAIN),('Architecture study',ARCH),('LangChain validation',CHAIN)]]}
    text+='''

# Appendix H. Research-question alignment and provenance

## H.1 Research questions and hypotheses

The main dissertation asks whether the LLM-agent replenishment system improves service and simulated cost, which benefits belong to its language component, which tasks it performs reliably, and how the findings differ between M5 and RetailNet. The following table maps those questions to the reported assessments. It preserves the comparison with the same numerical planner, which helps attribute performance to the appropriate component.

[[TABLE:research_alignment]]

H1–H4 are the analytical framework introduced in the revised literature review, not a preregistered confirmatory test plan. The separate architecture criteria A1–A4 were frozen before the matched architecture evaluation. H4 needs an explicit cross-dataset comparison of effects. Different statistical verdicts within the two datasets do not establish a difference between them, and perishability alone cannot be identified as the cause of the descriptive gap.

## H.2 Model and evidence identifiers

The common pretrained model digest is shown below. It identifies the locally served Qwen model used by all three empirical stages. The tag name alone is insufficient if a future model release changes its underlying content.

[[CODE:model_digest]]

[[TABLE:source_identifiers]]

These protocol hashes identify the settings inspected for the appendix. Full source-file and prepared-data hashes are stored inside those protocols, and raw native response hashes remain with the decision records. The source package provides a complete manifest for file-level checking. Its archival comparison distinguishes the full LangGraph architecture experiment from the later LangChain implementation validation. The LangChain port does not turn the earlier full study into fresh LangChain inference.

# Appendix I. Illustrative service target, separate from empirical results

The later contribution chapter includes an optional illustration of a 3% relative improvement in demand satisfied. This is an assumed planning target. It was not achieved as a measured 3% improvement in the experiments and is excluded from their empirical effect estimates. For completeness, the calculation is reproduced below without changing the actual results.

The assumed rate equals the previous fixed agent's unrounded measured fill rate multiplied by 1.03. A 3% relative increase is different from three percentage points. It corresponds to about 2.72 additional units satisfied per 100 units of demand in M5 and 2.52 in RetailNet. Rounding the printed baseline first can slightly change the displayed last digit; the calculation uses the archived unrounded values.

[[TABLE:illustrative_target]]

No cost, expiry or profit improvement follows automatically from this assumption. Such effects depend on the inventory and supply needed to achieve the service target. The empirical tables in the preceding appendices retain the observed simulation results and their stated uncertainty.
'''
    codes['model_digest']=base['model']['digest']
    scenario=csv(ROOT/'study_results/langchain_contribution_scenario_20261014/illustrative_scenario_calculations.csv')
    assert ((scenario.hypothetical_fill_rate_percent-scenario.measured_previous_fill_rate_percent*1.03).abs()<1e-10).all()
    definitions['illustrative_target']={'title':'Observed fill rates and an explicitly assumed 3% relative target',
        'headers':['Dataset','Previous fixed: observed','Selective: observed','3% relative target: assumed'],
        'rows':[[row.dataset,f'{row.measured_previous_fill_rate_percent:.2f}%',f'{row.measured_selective_fill_rate_percent:.2f}%',
                 f'{row.hypothetical_fill_rate_percent:.2f}%'] for row in scenario.itertuples()]}
    return text.strip()+'\n',definitions,figures,codes


def build():
    OUT.mkdir(parents=True,exist_ok=True)
    manuscript,definitions,figures,codes=assemble()
    (OUT/'appendices.md').write_text(manuscript)
    selected_tables={key:definitions[key] for key in re.findall(r'\[\[TABLE:(\w+)\]\]',manuscript)}
    (OUT/'appendix_tables.json').write_text(json.dumps(selected_tables,indent=2,ensure_ascii=False)+'\n')
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
    section.header_distance=section.footer_distance=Inches(.45)
    for style in doc.styles:
        if style.type in [1,2]:style.font.name='Times New Roman'
    for fonts in doc.styles.element.iter(qn('w:rFonts')):
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():del fonts.attrib[attr]
    normal=doc.styles['Normal'];normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;normal.paragraph_format.widow_control=True
    for name,size in [('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
        style=doc.styles[name];style.font.size=Pt(size);style.font.bold=True
        style.font.color.rgb=RGBColor.from_string('20343D');style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing=1.15;style.paragraph_format.keep_with_next=True
        style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(7)
        if name=='Heading 1':style.paragraph_format.page_break_before=True
    for name,size,spacing in [('Caption',10.5,1.1),('Technical',10.5,1.05),('Verbatim',11,1.2)]:
        if name not in doc.styles:doc.styles.add_style(name,1)
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size)
        style.font.color.rgb=RGBColor.from_string('20343D');style.font.bold=False;style.font.italic=False
        style.paragraph_format.line_spacing=spacing;style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_after=Pt(4)
    header=section.header.paragraphs[0];header.text='APPENDICES | M5 AND FRESHRETAILNET-50K'
    header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:run.font.size=Pt(9)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Page ')
    for field,tail in [('PAGE',' of '),('NUMPAGES','')]:
        node=OxmlElement('w:fldSimple');node.set(qn('w:instr'),field);footer._p.append(node);footer.add_run(tail)
    for run in footer.runs:run.font.size=Pt(10)
    letter=None;table_counter=figure_counter=0;table_records=[];figure_records=[];verbatim_records=[]
    keep_explanation=False
    for line in manuscript.splitlines():
        if not line.strip():continue
        heading=re.match(r'^(#{1,3}) (.*)',line)
        if heading:
            if len(heading[1])==1:
                letter=re.match(r'Appendix ([A-I])\.',heading[2])[1];table_counter=figure_counter=0
            p=doc.add_heading(heading[2],level=len(heading[1]))
            if heading[2].startswith(('G.2 ','G.3 ')):p.paragraph_format.page_break_before=True
        elif line.startswith('[[TABLE:'):
            key=line[8:-2];data=definitions[key];table_counter+=1;label=f'{letter}.{table_counter}'
            p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True
            p.add_run(f"Table {label}. {data['title']}").bold=True
            table=doc.add_table(rows=0,cols=len(data['headers']));table.style='Table Grid'
            if key in ['m5_catalogue','retailnet_catalogue']:
                table.autofit=False
                for column,width in zip(table.columns,[2.85,.75,.75,1.85]):column.width=Inches(width)
            for index,row in enumerate([data['headers'],*data['rows']]):
                cells=table.add_row().cells;table.rows[index]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if index==0:
                    table.rows[index]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for cell,value in zip(cells,row):
                    cell.text=str(value)
                    for p in cell.paragraphs:
                        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1.1
                        p.paragraph_format.space_after=Pt(3)
                        for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(10.5);run.bold=index==0
            if len(data['rows'])<=6:
                for index,row in enumerate(table.rows):
                    for cell in row.cells:
                        for p in cell.paragraphs:p.paragraph_format.keep_with_next=index<len(table.rows)-1
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            table_records.append({'key':key,'number':label,**data})
        elif line.startswith('[[FIGURE:'):
            key=line[9:-2];figure=figures[key];figure_counter+=1;label=f'{letter}.{figure_counter}'
            USED.add(figure['absolute_path']);assert sha(figure['absolute_path'])==figure['sha256']
            p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True
            p.add_run(f"Figure {label}. {figure['dataset']}: {figure['title']}").bold=True
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(figure['absolute_path']),width=Inches(6.2))
            p=doc.add_paragraph(figure['caption'],style='Caption');p.paragraph_format.keep_together=True
            if key=='holds':p.paragraph_format.keep_with_next=True;keep_explanation=True
            figure_records.append({'key':key,'number':label,'path':str(figure['absolute_path'].relative_to(ROOT)),
                                   'sha256':figure['sha256'],'caption':figure['caption']})
        elif line.startswith('[[CODE:') or line.startswith('[[VERBATIM:'):
            key=line.split(':',1)[1][:-2];value=codes[key]
            p=doc.add_paragraph(value,style='Technical' if line.startswith('[[CODE:') else 'Verbatim')
            p.paragraph_format.keep_together=False
            verbatim_records.append({'key':key,'text':value})
        else:
            p=doc.add_paragraph(line)
            if keep_explanation:p.paragraph_format.keep_together=True;keep_explanation=False
    doc.core_properties.title='Full Dissertation Appendices: M5, FreshRetailNet and the LangChain Extension'
    doc.core_properties.subject='Appendices only; archived experiments and a separately labelled illustrative target'
    doc.core_properties.author='';doc.core_properties.modified=datetime.now(timezone.utc).replace(tzinfo=None)
    doc.save(DOCX)
    manifest={'status':'BUILT_PENDING_LAYOUT_CHECK','scope':'Appendices A–I only; no main chapters or bibliography',
              'appendices':9,'tables':table_records,'figures':figure_records,'verbatim_blocks':verbatim_records,
              'prose_words':len(re.findall(r"\b[\w’'-]+\b",'\n'.join(line for line in manuscript.splitlines() if not line.startswith(('#','[['))))),
              'new_experiment_runs':0,'new_model_calls':0,'source_files':{str(path.relative_to(ROOT)):sha(path) for path in sorted(USED)},
              'docx_sha256':sha(DOCX),'manuscript_sha256':sha(OUT/'appendices.md'),
              'format':{'body_font':'Times New Roman','body_points':12,'line_spacing':1.5,'margins_inches':1,'alignment':'justified'}}
    (OUT/'build_manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
    print(json.dumps({k:manifest[k] for k in ['scope','appendices','prose_words','new_experiment_runs','docx_sha256']}|
                     {'tables':len(table_records),'figures':len(figure_records),'verbatim_blocks':len(verbatim_records)}))


if __name__=='__main__':build()
