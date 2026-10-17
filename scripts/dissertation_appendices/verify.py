"""Verify the appendix-only DOCX and its internal layout-check PDF."""
from pathlib import Path
import hashlib,json,re
import fitz
from docx import Document
from PIL import Image,ImageDraw

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'study_results/full_appendices_20261017'
manifest=json.loads((OUT/'build_manifest.json').read_text())
docx=OUT/'Full_Dissertation_Appendices.docx'
pdf_path=docx.with_suffix('.pdf')
document=Document(docx)
blocks={block['key']:block['text'] for block in manifest['verbatim_blocks']}
expected=[]
for line in (OUT/'appendices.md').read_text().splitlines():
    if not line.strip():continue
    if line.startswith(('[[TABLE:','[[FIGURE:')):continue
    if line.startswith(('[[CODE:','[[VERBATIM:')):expected.append(blocks[line.split(':',1)[1][:-2]])
    else:expected.append(re.sub(r'^#{1,3} ','',line))
actual=[p.text for p in document.paragraphs if p.text.strip() and p.style.name!='Caption']
assert actual==expected,'Full appendix text differs from the manuscript'
headings=[p.text for p in document.paragraphs if p.style.name=='Heading 1']
assert len(headings)==9 and all(text.startswith('Appendix '+chr(65+index)+'.') for index,text in enumerate(headings))
assert len(document.tables)==len(manifest['tables'])==17
assert len(document.inline_shapes)==len(manifest['figures'])==14
for table,data in zip(document.tables,manifest['tables']):
    assert [[cell.text for cell in row.cells] for row in table.rows]==[[str(x) for x in row] for row in [data['headers'],*data['rows']]],data['key']
for shape,figure in zip(document.inline_shapes,manifest['figures']):
    relation=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    assert hashlib.sha256(document.part.related_parts[relation].blob).hexdigest()==figure['sha256'],figure['key']
for relative,digest in manifest['source_files'].items():
    assert hashlib.sha256((ROOT/relative).read_bytes()).hexdigest()==digest,relative
normal=document.styles['Normal']
assert normal.font.name=='Times New Roman' and normal.font.size.pt==12
assert normal.paragraph_format.line_spacing==1.5
assert all(abs(getattr(document.sections[0],key).inches-1)<.001 for key in ['top_margin','bottom_margin','left_margin','right_margin'])
catalogues={t['key']:t for t in manifest['tables'] if t['key'] in ['m5_catalogue','retailnet_catalogue']}
assert all(len(table['rows'])==30 for table in catalogues.values())
pdf=fitz.open(pdf_path)
texts=[re.sub(r'\s+',' ',page.get_text()).strip() for page in pdf]
assert all(len(text)>40 for text in texts),'Empty or almost empty PDF page'
all_text=' '.join(texts)
for term in ['Appendix '+chr(letter)+'.' for letter in range(65,74)]:assert term in all_text,term
for term in ['17,100','LangChain','LangGraph','one-seed','No inventory','3%','assumed']:
    assert term.lower() in all_text.lower(),term
for index,page in enumerate(pdf):
    for x0,y0,x1,y1,*_ in page.get_text('blocks'):
        assert x0>=18 and y0>=14 and x1<=page.rect.width-18 and y1<=page.rect.height-14,(index+1,(x0,y0,x1,y1))
sheet=Image.new('RGB',(1080,385*((len(pdf)+3)//4)),'white')
for index,page in enumerate(pdf):
    pix=page.get_pixmap(matrix=fitz.Matrix(.42,.42))
    tile=Image.new('RGB',(270,385),'#ECEFF1')
    image=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    tile.paste(image,((270-image.width)//2,8))
    ImageDraw.Draw(tile).text((10,364),f'Page {index+1}',fill='black')
    sheet.paste(tile,((index%4)*270,(index//4)*385))
sheet.save(OUT/'review_contact_sheet.png')
for term,name in [('Appendix A.','settings'),('Table E.2.','architecture'),('F.1 Supplier-information','prompts'),
                  ('G.2 M5:','m5_catalogue'),('G.3 RetailNet:','retailnet_catalogue'),('Appendix I.','illustration')]:
    index=next(i for i,text in enumerate(texts) if term in text)
    pdf[index].get_pixmap(matrix=fitz.Matrix(1.15,1.15)).save(OUT/f'review_{name}.png')
receipt={'status':'PASS','scope':'Appendices A–I only','pdf_pages':len(pdf),'appendices':9,'tables':17,'figures':14,
         'verbatim_blocks':len(blocks),'complete_manuscript_and_tables_verified':True,'embedded_image_bytes_verified':True,
         'source_hashes_verified':True,'selected_series':{'M5':30,'RetailNet':30},'format_and_page_bounds_verified':True,
         'new_experiment_runs':0,'new_model_calls':0,'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),
         'layout_pdf_sha256':hashlib.sha256(pdf_path.read_bytes()).hexdigest(),'visual_review':'Pending rendered-page inspection'}
(OUT/'validation_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n')
print(json.dumps(receipt))
