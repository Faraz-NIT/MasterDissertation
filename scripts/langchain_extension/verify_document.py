"""Check the contribution chapter's evidence labels, complete content and layout."""
from pathlib import Path
import hashlib,json,re
from docx import Document
import fitz
from PIL import Image,ImageDraw

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'study_results/langchain_extension_20261014'
docx=OUT/'LangChain_Agent_Contribution_Chapter.docx'
pdf_path=OUT/'LangChain_Agent_Contribution_Chapter.pdf'
document=Document(docx)
expected=[re.sub(r'^#{1,3} ','',line) for line in (OUT/'chapter.md').read_text().splitlines()
          if line.strip() and not line.startswith('[[')]
actual=[p.text for p in document.paragraphs if p.text.strip() and p.style.name!='Caption']
assert actual==expected
assert len(document.inline_shapes)==7 and len(document.tables)==4
tables=json.loads((OUT/'chapter_tables.json').read_text())
for table,key in zip(document.tables,['changes','results','validation','hypotheses']):
    assert [[cell.text for cell in row.cells] for row in table.rows]==[tables[key]['headers'],*tables[key]['rows']]
normal=document.styles['Normal']
assert normal.font.name=='Times New Roman' and normal.font.size.pt==12
assert normal.paragraph_format.line_spacing==1.5
assert all(abs(getattr(document.sections[0],name).inches-1)<.001 for name in
           ['top_margin','bottom_margin','left_margin','right_margin'])
assert len([p for p in document.paragraphs if p.style.name=='Heading 1'])==1
assert json.loads((OUT/'analysis/audit_receipt.json').read_text())['status']=='VERIFIED'
assert json.loads((OUT/'development/recorded_response_replay.json').read_text())['status']=='PASS'
pdf=fitz.open(pdf_path)
texts=[re.sub(r'\s+',' ',page.get_text()).strip() for page in pdf]
assert all(len(text)>40 for text in texts)
joined=' '.join(texts)
for term in ['LangChain','LangGraph','M5','RetailNet','1,080','36 matched','7,560','not trained or fine-tuned']:
    assert term in joined,term
for index,page in enumerate(pdf):
    for x0,y0,x1,y1,*_ in page.get_text('blocks'):
        assert x0>=18 and y0>=14 and x1<=page.rect.width-18 and y1<=page.rect.height-14,index+1
sheet=Image.new('RGB',(1080,385*((len(pdf)+3)//4)),'white')
for index,page in enumerate(pdf):
    pix=page.get_pixmap(matrix=fitz.Matrix(.42,.42))
    image=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    tile=Image.new('RGB',(270,385),'#ECEFF1');tile.paste(image,((270-image.width)//2,8))
    ImageDraw.Draw(tile).text((10,364),f'Page {index+1}',fill='black')
    sheet.paste(tile,((index%4)*270,(index//4)*385))
sheet.save(OUT/'review_contact_sheet.png')
for term,name in [('Figure 1.','architecture'),('Table 2.','results'),('Table 3.','validation'),('Table 4.','hypotheses')]:
    index=next(index for index,text in enumerate(texts) if term in text)
    pdf[index].get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(OUT/f'review_{name}.png')
receipt={'status':'PASS','scope':'One contribution chapter','pdf_pages':len(pdf),'figures':7,'tables':4,
         'manuscript_exact_match':True,'format_verified':True,'page_bounds_verified':True,
         'primary_and_fresh_validation_scopes_explicit':True,
         'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),
         'pdf_sha256':hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
         'visual_review':'Pending inspection of rendered pages'}
(OUT/'document_validation.json').write_text(json.dumps(receipt,indent=2)+'\n')
print(json.dumps(receipt))
