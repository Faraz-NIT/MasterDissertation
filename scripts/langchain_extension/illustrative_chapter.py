"""Add an explicitly hypothetical 3% fill-rate scenario; never generate experiment data."""
from pathlib import Path
import hashlib,json,re,shutil

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches,Pt

ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'study_results/langchain_extension_20261014'
PRIMARY=ROOT/'study_results/architecture_study_20261013'
OUT=ROOT/'study_results/langchain_contribution_scenario_20261014'
TITLE='A LangChain Extension of the Replenishment Agent: Results and an Illustrative 3% Scenario'
NOTE='Evidence note: the experimental findings in this chapter are retained from the completed studies. The additional 3% increase in demand satisfied is an illustrative planning assumption, not an observed result. No new experiments were run for this revision.'
INTRO='This section asks what a 3% relative increase in demand satisfied would mean as a planning target. It is a hypothetical scenario. The calculation multiplies the previous fixed agent\'s unrounded measured fill rate by 1.03; it does not replace the selective agent\'s measured performance.'
INTERPRET='A 3% relative increase is different from an increase of three percentage points. For M5, the assumed gain is 2.72 additional units satisfied per 100 units of demand. For RetailNet, it is 2.52. The rates below are calculated from that assumption and have not been achieved in an experiment.'
LIMIT='The scenario illustrates the scale of a possible service target. It supplies no evidence that the architecture can reach it and supports no new claim about cost, waste or profit. Those effects depend on the orders, inventory and supply needed to deliver the target and would require a further study.'


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    folder=OUT/'figures';folder.mkdir(exist_ok=True)
    values=pd.read_csv(PRIMARY/'analysis/overall_seed_results.csv')
    rows=[]
    for dataset in ['M5','RetailNet']:
        baseline=float(values[(values.dataset==dataset)&(values.arm=='fixed_agents')].fill_rate.mean())
        measured=float(values[(values.dataset==dataset)&(values.arm=='selective_agents')].fill_rate.mean())
        assumed=baseline*1.03
        assert assumed<=1
        rows.append({'dataset':dataset,'measured_previous_fill_rate_percent':baseline*100,
                     'measured_selective_fill_rate_percent':measured*100,
                     'assumed_relative_improvement_percent':3.,'hypothetical_fill_rate_percent':assumed*100,
                     'hypothetical_change_percentage_points':(assumed-baseline)*100,
                     'scenario_evidence_type':'ILLUSTRATIVE_ASSUMPTION_NOT_AN_EXPERIMENTAL_RESULT'})
    pd.DataFrame(rows).to_csv(OUT/'illustrative_scenario_calculations.csv',index=False)
    fig,axes=plt.subplots(1,2,figsize=(9,3.6),layout='constrained')
    for ax,row in zip(axes,rows):
        heights=[row['measured_previous_fill_rate_percent'],row['measured_selective_fill_rate_percent'],row['hypothetical_fill_rate_percent']]
        bars=ax.bar(['Previous\nobserved','Selective\nobserved','3% scenario\nassumed'],heights,
                    color=['#92A5AF','#397F9E','#EEF2F5'],edgecolor=['#92A5AF','#397F9E','#7B8F9B'],width=.65)
        bars[2].set_hatch('///')
        ax.set_ylim(0,100);ax.set_ylabel('Demand satisfied (%)');ax.set_title(row['dataset'])
        ax.spines[['top','right']].set_visible(False)
        for index,height in enumerate(heights):ax.text(index,height+1,f'{height:.2f}%',ha='center',fontsize=10)
    fig.suptitle('ILLUSTRATIVE SCENARIO: assumes a 3% relative improvement',fontsize=11,fontweight='bold')
    fig.savefig(folder/'illustrative_3percent_scenario.png',dpi=200);plt.close(fig)
    figures=json.loads((SOURCE/'figure_register.json').read_text())['figures']
    for figure in figures:shutil.copyfile(SOURCE/figure['path'],OUT/figure['path'])
    scenario_figure={'key':'illustrative_3percent_scenario','path':'figures/illustrative_3percent_scenario.png',
                     'title':'Illustrative 3% relative improvement in demand satisfied',
                     'scope':'M5 + RetailNet; hypothetical scenario',
                     'caption':'The first two bars show observed rates from the completed architecture study. The hatched bar assumes a 3% relative gain over the previous fixed agent. It is a calculated illustration and is not an experimental result.'}
    figures.insert(6,scenario_figure)
    (OUT/'figure_register.json').write_text(json.dumps({'figures':figures,'source_chapter':str(SOURCE.relative_to(ROOT))},indent=2)+'\n')
    tables=json.loads((SOURCE/'chapter_tables.json').read_text())
    scenario_table={'title':'Observed demand satisfied and a separately labelled illustrative scenario',
                    'headers':['Dataset','Previous agent: observed','Selective extension: observed','3% relative gain: assumed'],
                    'rows':[[r['dataset'],f"{r['measured_previous_fill_rate_percent']:.2f}%",f"{r['measured_selective_fill_rate_percent']:.2f}%",f"{r['hypothetical_fill_rate_percent']:.2f}%"] for r in rows]}
    tables['scenario']=scenario_table
    (OUT/'chapter_tables.json').write_text(json.dumps(tables,indent=2)+'\n')
    manuscript=(SOURCE/'chapter.md').read_text()
    manuscript=manuscript.replace(manuscript.splitlines()[0],'# '+TITLE,1)
    manuscript=manuscript.replace('\n\n## 1.', '\n\n'+NOTE+'\n\n## 1.',1)
    section='### 5.4 An illustrative 3% service-improvement scenario\n\n'+INTRO+'\n\n'+INTERPRET+'\n\n[[TABLE:scenario]]\n\n[[FIGURE:illustrative_3percent_scenario]]\n\n'+LIMIT+'\n\n'
    manuscript=manuscript.replace('### 5.4 What the LangChain validation confirms',section+'### 5.5 What the LangChain validation confirms',1)
    (OUT/'chapter.md').write_text(manuscript)
    doc=Document(SOURCE/'LangChain_Agent_Contribution_Chapter.docx')
    title=next(p for p in doc.paragraphs if p.style.name=='Heading 1');title.text=TITLE
    first_heading=next(p for p in doc.paragraphs if p.text=='1. Building on the previous dissertation')
    first_heading.insert_paragraph_before(NOTE)
    for paragraph in doc.paragraphs:
        match=re.match(r'^(Figure|Table) (\d+)\.',paragraph.text)
        if match and paragraph.style.name=='Caption':
            threshold=7 if match[1]=='Figure' else 3
            if int(match[2])>=threshold:
                paragraph.runs[0].text=re.sub(r'^(Figure|Table) \d+\.',f'{match[1]} {int(match[2])+1}.',paragraph.runs[0].text)
    anchor=next(p for p in doc.paragraphs if p.text=='5.4 What the LangChain validation confirms')
    anchor.text='5.5 What the LangChain validation confirms'
    anchor.insert_paragraph_before('5.4 An illustrative 3% service-improvement scenario',style='Heading 3')
    anchor.insert_paragraph_before(INTRO)
    anchor.insert_paragraph_before(INTERPRET)
    label=anchor.insert_paragraph_before('Table 3. '+scenario_table['title'],style='Caption')
    label.paragraph_format.keep_with_next=True;label.runs[0].bold=True
    table=doc.add_table(rows=0,cols=4);table.style='Table Grid'
    for index,row in enumerate([scenario_table['headers'],*scenario_table['rows']]):
        cells=table.add_row().cells;table.rows[index]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell,value in zip(cells,row):
            cell.text=value
            for p in cell.paragraphs:
                p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1.1
                p.paragraph_format.space_after=Pt(4)
                for run in p.runs:run.font.size=Pt(10);run.bold=index==0
    anchor._p.addprevious(table._tbl)
    label=anchor.insert_paragraph_before('Figure 7. '+scenario_figure['scope']+': '+scenario_figure['title'],style='Caption')
    label.paragraph_format.keep_with_next=True;label.runs[0].bold=True
    picture=anchor.insert_paragraph_before();picture.alignment=WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.line_spacing=1;picture.paragraph_format.keep_with_next=True
    picture.add_run().add_picture(str(folder/'illustrative_3percent_scenario.png'),width=Inches(6.2))
    anchor.insert_paragraph_before(scenario_figure['caption'],style='Caption')
    anchor.insert_paragraph_before(LIMIT)
    for index,row in enumerate(doc.tables[-1].rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next=index<len(doc.tables[-1].rows)-1
    header=doc.sections[0].header.paragraphs[0]
    header.text='M5 AND RETAILNET | OBSERVED RESULTS + ILLUSTRATIVE 3% SCENARIO'
    for run in header.runs:run.font.size=Pt(9)
    doc.core_properties.title=TITLE
    doc.core_properties.subject='Measured architecture results plus an explicitly assumed 3% planning scenario; no new experiments'
    target=OUT/'Chapter_With_Illustrative_3Percent_Scenario.docx';doc.save(target)
    expected=[re.sub(r'^#{1,3} ','',line) for line in manuscript.splitlines() if line.strip() and not line.startswith('[[')]
    actual=[p.text for p in doc.paragraphs if p.text.strip() and p.style.name!='Caption']
    assert actual==expected
    assert len(doc.inline_shapes)==8 and len(doc.tables)==5
    for table,key in zip(doc.tables,['changes','results','scenario','validation','hypotheses']):
        assert [[c.text for c in row.cells] for row in table.rows]==[tables[key]['headers'],*tables[key]['rows']]
    manifest={'scope':'Chapter revision with explicitly hypothetical 3% relative fill-rate scenario',
              'new_experiment_runs':0,'new_model_calls':0,'measured_results_modified':False,
              'assumption':'3% relative improvement over the measured fixed-agent fill rate; not three percentage points',
              'source_docx':{'path':str((SOURCE/'LangChain_Agent_Contribution_Chapter.docx').relative_to(ROOT)),
                             'sha256':sha(SOURCE/'LangChain_Agent_Contribution_Chapter.docx')},
              'source_measured_results':{'path':str((PRIMARY/'analysis/overall_seed_results.csv').relative_to(ROOT)),
                                         'sha256':sha(PRIMARY/'analysis/overall_seed_results.csv')},
              'docx_sha256':sha(target),'figures':8,'tables':5,'manuscript_and_tables_verified':True}
    (OUT/'build_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(json.dumps(manifest))


if __name__=='__main__':main()
