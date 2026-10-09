"""Readable results report and dataset-labelled figures from audited fresh runs."""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import subprocess
from datetime import datetime,timedelta
from zoneinfo import ZoneInfo

os.environ.setdefault('MPLCONFIGDIR','/tmp/ollama-clean-study-matplotlib')
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'study_results/ollama_clean_study_20261009'
ARMS=['conventional','structured_planner','llm_agents']
NAMES={'conventional':'Conventional ordering','structured_planner':'Structured-data planner','llm_agents':'Ollama agents'}
COLORS={'conventional':'#B9C1CA','structured_planner':'#6F9DBD','llm_agents':'#6BA98B'}
SCENARIOS=['routine','supplier_disruption','promotion']
LABELS={'routine':'Routine','supplier_disruption':'Supplier disruption','promotion':'Promotion'}
REGISTER=[]
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':14,'axes.spines.top':False,
    'axes.spines.right':False,'axes.titleweight':'bold','savefig.dpi':180,'figure.facecolor':'white'})


def sha(path):
    with Path(path).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()


def save(fig,name,title,caption,dataset):
    path=OUT/'figures'/f'{name}.png';path.parent.mkdir(exist_ok=True)
    fig.savefig(path,bbox_inches='tight',facecolor='white');plt.close(fig)
    REGISTER.append({'name':name,'title':title,'caption':caption,'dataset':dataset,
        'path':str(path.relative_to(OUT)),'sha256':sha(path)})


def bars(means,dataset,metric,label,scale,name,title):
    fig,ax=plt.subplots(figsize=(9.4,4.5));x=np.arange(3);width=.24
    for i,arm in enumerate(ARMS):
        frame=means[(means.dataset==dataset)&(means.arm==arm)].set_index('scenario').loc[SCENARIOS]
        center=frame[metric].to_numpy()*scale
        low=frame[metric+'_ci_low'].to_numpy()*scale;high=frame[metric+'_ci_high'].to_numpy()*scale
        ax.bar(x+(i-1)*width,center,width,color=COLORS[arm],label=NAMES[arm],
            yerr=np.maximum(0,np.array([center-low,high-center])),capsize=3,error_kw={'lw':1})
    ax.set_xticks(x,[LABELS[s] for s in SCENARIOS]);ax.set_ylabel(label);ax.set_title(dataset+' — '+title,pad=13)
    ax.grid(axis='y',alpha=.16);ax.set_axisbelow(True);ax.legend(loc='upper center',bbox_to_anchor=(.5,-.13),ncol=3,frameon=False)
    if metric in ('fill_rate','expiry_share'):ax.set_ylim(0,101 if metric=='fill_rate' else max(1,ax.get_ylim()[1]))
    save(fig,name,dataset+' — '+title,'Means across 30 paired simulation seeds, with both test weeks aggregated within seed. Error bars: 95% bootstrap intervals. '+('Higher is better.' if metric=='fill_rate' else 'Lower is better.'),dataset)


def figures():
    means=pd.read_csv(OUT/'analysis/policy_means.csv');days=pd.read_csv(OUT/'analysis/daily_results.csv')
    model=pd.read_csv(OUT/'analysis/model_calls.csv');effects=pd.read_csv(OUT/'analysis/paired_effects.csv')
    # Minimal architecture: three genuine LLM calls, bounded by shared tools.
    dot='''digraph NewStudy { graph [rankdir=TB,bgcolor="white",pad=.2,nodesep=.25,ranksep=.4]; node [shape=box,style="rounded,filled",fontname="Arial",fontsize=13,color="#BDC8D2",fillcolor="#F4F7FA",margin=".13,.12"]; edge [color="#8B9AA8",arrowsize=.7]; supplier [label="Supplier agent\\nOllama: reads messages",fillcolor="#F4EFF8"]; coordinator [label="Planning coordinator\\nOllama: selects tool and priority",fillcolor="#F4EFF8"]; tool [label="Planning tool\\nCalculates order quantities"]; review [label="Review agent\\nOllama: reviews checks",fillcolor="#F4EFF8"]; gate [label="Execution gate\\nChecks must pass"]; execution [label="Inventory simulator\\nService, cost and expiry",fillcolor="#EDF6F0"]; {rank=same;supplier;coordinator;tool;} {rank=same;execution;gate;review;} supplier -> coordinator [constraint=false]; coordinator -> tool [constraint=false]; tool -> review; review -> gate [constraint=false]; gate -> execution [constraint=false]; supplier -> execution [style=invis]; coordinator -> gate [style=invis]; }'''
    folder=OUT/'figures';folder.mkdir(exist_ok=True);(folder/'01_agent_workflow.dot').write_text(dot)
    subprocess.run(['dot','-Tpng','-Gdpi=180',str(folder/'01_agent_workflow.dot'),'-o',str(folder/'01_agent_workflow.png')],check=True,capture_output=True)
    subprocess.run(['dot','-Tsvg',str(folder/'01_agent_workflow.dot'),'-o',str(folder/'01_agent_workflow.svg')],check=True,capture_output=True)
    REGISTER.append({'name':'01_agent_workflow','title':'M5 + RetailNet — the new no-repair agent workflow','dataset':'M5 + RetailNet',
        'caption':'Three separate fresh Ollama calls per portfolio decision. Current contract fields are independently verified; numerical tools compute quantities and the execution gate controls permission. No repair mechanism is used.',
        'path':'figures/01_agent_workflow.png','sha256':sha(folder/'01_agent_workflow.png')})
    fig,axes=plt.subplots(2,1,figsize=(9.4,5.7))
    protocol=json.loads((OUT/'protocol.json').read_text())
    for ax,dataset,color in zip(axes,['M5','RetailNet'],['#24548A','#27775D']):
        origin=protocol['datasets'][dataset]['origins'][-1]
        with np.load(OUT/f'data/{dataset}/origin_{origin}/panel_and_forecast.npz') as d:sales=d['sales']
        start=origin-7
        ax.plot(np.arange(1,15),sales[:,start:start+14].sum(axis=0),color=color,marker='o',ms=4)
        forecast_parts=[]
        for fold in protocol['datasets'][dataset]['origins']:
            with np.load(OUT/f'data/{dataset}/origin_{fold}/panel_and_forecast.npz') as arrays:
                forecast_parts.append(arrays['forecast'][:,:7].sum(axis=0))
        ax.plot(np.arange(1,15),np.concatenate(forecast_parts),color='#818D98',ls='--',label='Shared LightGBM forecast')
        ax.legend(frameon=False,fontsize=12)
        ax.axvline(7.5,color='#AAB4BE',ls='--');ax.set_title(dataset+' — recorded sales in the two test weeks')
        ax.set_ylabel('Unit sales' if dataset=='M5' else 'Normalized sales × 100');ax.grid(alpha=.14)
        ax.set_xticks([1,4,7,8,11,14]);ax.set_xlabel('Historical day across the two weekly test windows')
    fig.tight_layout();save(fig,'02_dataset_profiles','M5 + RetailNet — test-week sales profiles','Solid lines: recorded sales across each 30-series panel. Dashed lines: the shared numerical LightGBM forecast, fitted afresh at each weekly origin. The axes use different quantity conventions; heights must not be interpreted as comparable commercial demand. The LLM does not produce these forecasts.','M5 + RetailNet')
    for dataset,prefix in [('M5','03'),('RetailNet','04')]:
        bars(means,dataset,'fill_rate','Demand fulfilled (%)',100,prefix+'_fill_rate','customer-demand fulfilment')
    for dataset,prefix in [('M5','05'),('RetailNet','06')]:
        bars(means,dataset,'cost_per_demand','Operating cost index / demand unit',1,prefix+'_operating_cost','inventory operating cost')
    for dataset,prefix in [('M5','07'),('RetailNet','08')]:
        fig,axes=plt.subplots(1,3,figsize=(11,3.5),sharey=True)
        for ax,scenario in zip(axes,SCENARIOS):
            frame=days[(days.dataset==dataset)&(days.scenario==scenario)]
            for arm in ARMS:
                s=frame[frame.arm==arm].groupby('day')[['sales','demand']].sum()
                ax.plot(s.index+1,100*s.sales/s.demand,label=NAMES[arm],color=COLORS[arm],lw=2,marker='o',ms=3)
            ax.set_title(LABELS[scenario],fontsize=14);ax.set_xticks([1,3,5,7]);ax.set_xlabel('Day within the week');ax.grid(alpha=.14);ax.set_ylim(0,103)
        axes[0].set_ylabel('Demand fulfilled (%)');fig.suptitle(dataset+' — service through each simulated week',y=1.03,fontweight='bold')
        axes[1].legend(loc='upper center',bbox_to_anchor=(.5,-.23),ncol=3,frameon=False,fontsize=13)
        save(fig,prefix+'_daily_service',dataset+' — daily service','Daily demand-weighted fulfilment, pooling both weekly windows and 30 seeds within each scenario. The daily plots are descriptive; days are not treated as independent replications.',dataset)
    bars(means,'RetailNet','expiry_share','Expired share of opening stock + receipts (%)',100,'09_retailnet_expiry','fresh-stock expiry')
    fig,axes=plt.subplots(1,2,figsize=(10,4.4))
    for ax,dataset in zip(axes,['M5','RetailNet']):
        frame=days[days.dataset==dataset].groupby('arm')[['holding_cost','shortage_cost','expiry_cost','demand']].sum().loc[ARMS]
        bottom=np.zeros(3)
        for component,color,label in [('holding_cost','#B9D2E1','Holding'),('shortage_cost','#E3B8A4','Unmet demand'),('expiry_cost','#AAC9B4','Expiry')]:
            values=frame[component].to_numpy()/frame.demand.to_numpy();ax.bar(np.arange(3),values,bottom=bottom,color=color,label=label);bottom+=values
        ax.set_title(dataset+'\nOperating-cost components');ax.set_xticks(np.arange(3),['Conventional','Structured\nplanner','Ollama\nagents']);ax.set_ylabel('Cost index / demand unit');ax.grid(axis='y',alpha=.15)
    handles,labels=axes[-1].get_legend_handles_labels()
    fig.legend(handles,labels,loc='upper center',bbox_to_anchor=(.5,-.02),ncol=3,frameon=False)
    fig.tight_layout();save(fig,'10_cost_components','M5 + RetailNet — operating-cost components','The primary index includes holding, unmet-demand penalties and expiry. Acquisition spend is reported separately. Ratios here use pooled demand; policy-mean tables use the mean of per-seed ratios.','M5 + RetailNet')
    accuracy=pd.read_csv(OUT/'analysis/source_accuracy_by_case.csv')
    fig,axes=plt.subplots(1,2,figsize=(10,4.3),sharey=True)
    for ax,dataset,color in zip(axes,['M5','RetailNet'],['#24548A','#27775D']):
        frame=accuracy[accuracy.dataset==dataset].groupby('style').agg(exact=('exact','mean'),fields=('correct_slots','mean'))
        x=np.arange(3);ax.bar(x-.17,frame.exact*100,.34,label='Complete extraction',color=color,alpha=.8)
        ax.bar(x+.17,frame.fields/13*100,.34,label='Individual fields',color='#C0CDD6')
        ax.set_xticks(x,['Prose','Memo','Table']);ax.set_title(dataset+' — supplier reading');ax.set_ylim(0,102);ax.grid(axis='y',alpha=.14)
    axes[0].set_ylabel('Correct (%)');axes[-1].legend(loc='lower left',frameon=False,fontsize=12);fig.tight_layout()
    save(fig,'11_source_accuracy','M5 + RetailNet — LLM supplier-reading accuracy','Exactly 13 numerical fields are scored per supplier call: four fields for each of three suppliers, plus the daily budget. Supplier identity/coverage must also be correct for a complete extraction. This is controlled language, not an unrestricted commercial document corpus.','M5 + RetailNet')
    fig,ax=plt.subplots(figsize=(8.8,4.2));x=np.arange(3);roles=['supplier','planner','critic']
    for i,dataset in enumerate(['M5','RetailNet']):
        frame=model[model.dataset==dataset];median=[frame[frame.role==r].seconds.median() for r in roles]
        high=[frame[frame.role==r].seconds.quantile(.95) for r in roles]
        ax.bar(x+(i-.5)*.34,median,.34,label=dataset,color=['#88ABC7','#95BCA8'][i],
            yerr=np.array([np.zeros(3),np.maximum(0,np.array(high)-median)]),capsize=4)
    ax.set_xticks(x,['Supplier reader','Planning coordinator','Review agent']);ax.set_ylabel('Seconds per Ollama call');ax.set_title('M5 + RetailNet — local model latency');ax.grid(axis='y',alpha=.14);ax.legend(frameon=False)
    save(fig,'12_model_latency','M5 + RetailNet — local model latency','Bars show median recorded call time; upper whiskers show the 95th percentile, not a confidence interval. The local model was warm following development. Hardware and integration costs were not monetised.','M5 + RetailNet')
    for metric,prefix,xlabel in [('fill_rate','13','LLM minus comparator service\n(percentage points)'),('cost_per_demand','14','LLM operating-cost change\nversus comparator (%)')]:
        fig,axes=plt.subplots(1,2,figsize=(11,5.3),sharey=True)
        for ax,comparison in zip(axes,['conventional','structured_planner']):
            subset=effects[(effects.metric==metric)&(effects.comparison==comparison)].set_index(['dataset','scenario'])
            labels=[]
            for y,(dataset,scenario) in enumerate([(d,s) for d in ['M5','RetailNet'] for s in SCENARIOS]):
                r=subset.loc[(dataset,scenario)]
                fields=('relative_change_pct','relative_ci_low','relative_ci_high') if metric=='cost_per_demand' else ('mean','ci_low','ci_high')
                center,lo,hi=[float(r[k]) for k in fields];color='#24548A' if dataset=='M5' else '#27775D'
                ax.plot([lo,hi],[y,y],color=color,lw=2);ax.scatter([center],[y],color=color,s=38,zorder=3)
                labels.append(dataset+' / '+LABELS[scenario])
            ax.axvline(0,color='#8D9BAC',ls='--');ax.set_yticks(np.arange(6),labels);ax.grid(axis='x',alpha=.14)
            ax.set_title('Versus '+NAMES[comparison],fontsize=14);ax.set_xlabel(xlabel,fontsize=13)
        axes[0].invert_yaxis()
        fig.suptitle('M5 + RetailNet — paired '+('service effects' if metric=='fill_rate' else 'operating-cost effects'),fontweight='bold');fig.tight_layout()
        friendly='service' if metric=='fill_rate' else 'operating-cost'
        cost_note='Cost changes are means of paired seed-relative percentages. ' if metric=='cost_per_demand' else ''
        save(fig,prefix+'_paired_effects','M5 + RetailNet — paired '+friendly+' effects','Thirty paired simulation seeds, with two weeks aggregated inside seed; 95% bootstrap intervals. '+cost_note+'The structured comparator uses the same numerical tool and current rules, isolating the model-assisted interpretation/routing layer.','M5 + RetailNet')
    fig,axes=plt.subplots(1,2,figsize=(10,4.2))
    for ax,dataset in zip(axes,['M5','RetailNet']):
        frame=means[means.dataset==dataset].set_index(['arm','scenario'])
        for i,arm in enumerate(ARMS):
            ax.bar(np.arange(3)+(i-1)*.24,[frame.loc[(arm,s),'held_days'] for s in SCENARIOS],.24,color=COLORS[arm],label=NAMES[arm])
        ax.set_xticks(np.arange(3),['Routine','Disruption','Promotion']);ax.set_title(dataset+' — held decisions');ax.set_ylabel('Mean held days per 14-day seed');ax.grid(axis='y',alpha=.14)
        ax.set_ylim(0,max(1,ax.get_ylim()[1]))
        if frame.held_days.max()==0:ax.text(.5,.45,'No held decisions',transform=ax.transAxes,ha='center',color='#516374',fontsize=16)
    axes[-1].legend(frameon=False,fontsize=11);fig.tight_layout();save(fig,'15_holds','M5 + RetailNet — held decisions','A hold is an unexecuted decision, not a successful zero-order recommendation. The same deterministic budget, capacity, pack and minimum-order checks govern all arms.','M5 + RetailNet')
    (OUT/'figure_register.json').write_text(json.dumps({'figures':REGISTER,'source_files':{p.name:sha(p) for p in (OUT/'analysis').glob('*.csv')}},indent=2)+'\n')


def narrative():
    from analyze import estimate
    seed=pd.read_csv(OUT/'analysis/seed_results.csv');diagnostics=pd.read_csv(OUT/'analysis/llm_diagnostics.csv')
    execution=json.loads((OUT/'execution_receipt.json').read_text());p=json.loads((OUT/'protocol.json').read_text())
    means=pd.read_csv(OUT/'analysis/policy_means.csv');matched=pd.read_csv(OUT/'analysis/matched_actions.csv')
    overall=seed.groupby(['dataset','arm','seed'])[['fill_rate','cost_per_demand','expiry_share']].mean().reset_index()
    overall_means=overall.groupby(['dataset','arm'])[['fill_rate','cost_per_demand','expiry_share']].mean()
    seed_count=int(seed['seed'].nunique())
    finish=datetime.fromisoformat(execution['completed_at_utc'])
    start=finish-timedelta(seconds=execution['runtime_seconds'])
    paris=ZoneInfo('Europe/Paris')
    overall_rows=[]
    for dataset in ['M5','RetailNet']:
        group=overall[overall.dataset==dataset];agent=group[group.arm=='llm_agents'].set_index('seed')
        for comparison in ['conventional','structured_planner']:
            baseline=group[group.arm==comparison].set_index('seed');ids=agent.index.intersection(baseline.index)
            fill=estimate(100*(agent.loc[ids,'fill_rate']-baseline.loc[ids,'fill_rate']))
            cost=estimate(100*(agent.loc[ids,'cost_per_demand']/baseline.loc[ids,'cost_per_demand']-1))
            overall_rows.append({'dataset':dataset,'comparison':comparison,'fill_effect_pp':fill,'cost_effect_pct':cost})
    (OUT/'analysis/overall_effects.json').write_text(json.dumps(overall_rows,indent=2)+'\n')
    lines=['# Fresh Ollama replenishment study: M5 and RetailNet',
        '',f"Execution status: **{execution['status']}**. This report concerns the redesigned study only. It does not combine earlier repairs or earlier experiment results.",
        '', '## What the study found', '',
        f"The study completed {execution['runs']:,} runs, {execution['portfolio_days']:,} portfolio decision days and {execution['fresh_primary_model_calls']:,} fresh local Ollama calls. Each portfolio contained 30 selected store–product series. The reported comparisons use {seed_count} complete, balanced simulation seeds across two seven-day historical windows in each dataset and three business situations.",
        '',f"Live evaluation ran from {start.astimezone(paris):%d %B %Y, %H:%M} to {finish.astimezone(paris):%d %B %Y, %H:%M} (Europe/Paris), taking {execution['runtime_seconds']/3600:.2f} hours. Development, fresh forecast fitting, auditing and report preparation are additional steps within the registered time budget. Exact UTC event times are preserved in the logs.",
        '']
    for metric,label,scale,precision in [('fill_rate','Customer demand fulfilled (%)',100,2),('cost_per_demand','Inventory operating-cost index per demand unit',1,3)]:
        lines += ['### '+label,'','| Dataset | Conventional ordering | Structured-data planner | Ollama agents |',
            '| --- | ---: | ---: | ---: |']
        for dataset in ['M5','RetailNet']:
            values=[overall_means.loc[(dataset,arm),metric]*scale for arm in ARMS]
            lines += ['| '+dataset+' | '+' | '.join(f'{value:.{precision}f}' for value in values)+' |']
        lines += ['','These summary values weight the three business situations equally. Figures below show each situation separately.','']
    for dataset in ['M5','RetailNet']:
        a=next(r for r in overall_rows if r['dataset']==dataset and r['comparison']=='conventional')
        b=next(r for r in overall_rows if r['dataset']==dataset and r['comparison']=='structured_planner')
        fill=a['fill_effect_pp']['mean'];cost=a['cost_effect_pct']['mean'];shared=matched[matched.dataset==dataset].orders_identical.mean()*100
        lines += [f"**{dataset}.** Compared with conventional ordering, the Ollama workflow changed demand fulfilment by {fill:+.2f} percentage points and the operating-cost index by {cost:+.2f}%. Against the structured-data planner, the changes were {b['fill_effect_pp']['mean']:+.2f} points and {b['cost_effect_pct']['mean']:+.2f}%. Its committed order vector matched that planner on {shared:.2f}% of portfolio days.",'']
    total_violations=int(means.actual_violations.sum())
    lines += [f"There were {total_violations:,} observed committed-action constraint violations across the recorded study. That finding describes the whole guarded workflow; it does not demonstrate that an unrestricted language model would be safe.", '',
        'The fair interpretation is to assess two questions separately. The comparison with conventional ordering measures the complete workflow, including numerical planning. The comparison with the structured-data planner asks whether adding the language-model roles changes the result once the same planning capability and current business rules are available.', '',
        '## A simple design', '',
        'Three actual language-model calls are made for each agent decision. The supplier agent reads a current supplier message. The planning coordinator requests the constrained-order tool and selects the permitted service priority. The review agent sees independent checks and either approves or holds the proposal. Numerical tools calculate quantities. A deterministic gate remains responsible for execution.', '',
        'No inventory repair, sales recovery, data-fault experiment or previous response cache is used. An incorrect extraction, unavailable model, infeasible plan or failed check produces a hold. The run does not quietly substitute a conventional-policy result for a failed agent decision.', '',
        'The conventional comparator uses an order-up-to rule and proportionate allocation under supplier and spending limits. The structured-data planner uses the same constrained-order tool, current rules, age projection and promotion-priority rule as the agents, without a language model. It is a strong, idealised comparator: its inputs are already structured. This study does not measure the staff time that would be needed to create those inputs in a real retailer.', '',
        'Routine trading uses unchanged rules. In the supplier-disruption situation, cartons and minimum orders double, delivery takes one extra day and capacity falls by 30% from the third simulated day. In the promotion situation, an announced 30% uplift affects demand and the forecast from the third day. Every policy receives the same announced information, initial stock and keyed demand/supply draws.', '',
        'Each test week is a separate simulation episode with the same opening-state rule. Inventory is not carried from the earlier window into the later one. Carton and minimum-order rules apply to individual order lines; dispatch capacity applies to each synthetic supplier group.', '',
        '## Data and model', '',
        'M5 uses 30 train-selected FOODS store–product pairs. RetailNet uses 30 train-selected store–product pairs from Dingdong-Inc/FreshRetailNet-50K. Both selections cover three historical sales-volume strata. These panels support the reported performance comparisons; they are not full-population evaluations of all M5 or all 50,000 RetailNet series.', '',
        'M5 is tested on 9–22 May 2016, divided into two seven-day windows. RetailNet uses 19 June–2 July 2024: a seven-day rolling holdout from the released training file and the official seven-day evaluation file. A fresh LightGBM forecasting model is fitted locally for each dataset and window using only earlier observations. This is one forecasting family, shared by every policy. Forecasting uses recursive predictions and no evaluation sales as input.', '',
        'The language model is pretrained Qwen2.5 1.5B, run through the local Ollama native API. It is not trained or fine-tuned on either retail dataset. The computer is Linux with a four-CPU allocation, 32 GiB memory and no GPU inference. Generation uses three CPU threads, a 4,096-token context, temperature zero and a fixed model seed. The protocol records the model digest and package versions.', '',
        '## How to read the results', '',
        'Fill rate is the share of simulated customer demand that stock could satisfy. Higher is better. The operating-cost index adds holding cost, a penalty for unmet demand and the cost of expired inventory. Lower is better. Purchase spend is recorded separately so acquisition and expiry are not double-counted. Costs, suppliers, inventories and supply events are experimental assumptions, not observed retailer profit.', '',
        'RetailNet quantities are published normalized sales multiplied by 100. M5 uses published unit sales. Absolute commercial volumes and profits cannot be compared across these datasets. The cross-dataset comparison uses fulfilment and within-dataset relative policy effects. Fresh-stock expiry assumes three usable days and FIFO; the source dataset does not provide an observed shelf-life field.', '',
        'Recorded sales can be limited by stockouts. This study deliberately uses released sales as a demand proxy, without estimating the purchases that were never recorded. Fulfilment therefore describes the declared simulation world, rather than the retailer\'s actual service level or recovered latent demand.', '',
        'The 95% intervals resample paired simulation seeds 10,000 times. Both weekly windows are aggregated inside each seed. The intervals reflect the selected simulated weeks and demand/supply draws. They do not turn 30 chosen series into a population estimate, and the repeated days are not treated as independent observations.', '',
        '## LLM reliability and workload', '']
    for row in diagnostics.itertuples():
        lines += [f"**{row.dataset}.** Complete supplier extraction was correct on {100*row.exact_extraction_rate:.2f}% of decisions; individual numerical fields were {100*row.field_accuracy:.2f}% correct. Service-priority selection was correct on {100*row.priority_accuracy:.2f}% of decisions. The workflow held {int(row.held_days):,} agent decisions. Mean recorded decision time was {row.mean_decision_seconds:.2f} seconds; the 95th percentile was {row.p95_decision_seconds:.2f} seconds. The review model gave {int(row.critic_false_approvals):,} approvals against failing independent checks and {int(row.critic_false_holds):,} holds against passing checks. The final gate still governs execution.",'']
    lines += ['Development was kept separate and its failures were preserved. The first development round exposed field-swapping and incorrect review decisions. Few-shot prompt examples were added before the design was frozen; there is no evaluation-time prompt tuning, response correction or model fine-tuning. Primary model requests and raw responses are all saved individually.', '',
        '## Comparing M5 and RetailNet', '',
        f"The Ollama workflow fulfilled {100*overall_means.loc[('M5','llm_agents'),'fill_rate']:.2f}% of simulated demand in M5 and {100*overall_means.loc[('RetailNet','llm_agents'),'fill_rate']:.2f}% in RetailNet, averaging the three situations equally. RetailNet also had {100*overall_means.loc[('RetailNet','llm_agents'),'expiry_share']:.2f}% expired inventory as a share of opening stock plus receipts. M5 had no expiry within the study horizon. These are two different retail settings, rather than a ranking of dataset quality or model intelligence.", '',
        'The difference matters for management. M5 tests keeping sufficient stock when delivery and ordering rules change. RetailNet adds a trade-off between availability and short-lived stock. A policy that orders more can protect customer service while increasing waste. The graphs therefore show fulfilment, operating cost and RetailNet expiry together. Absolute cost levels use synthetic quantity conventions and must not be treated as a comparison of retailer profitability.', '',
        '## Business conclusion', '']
    for dataset in ['M5','RetailNet']:
        control=next(row for row in overall_rows if row['dataset']==dataset and row['comparison']=='conventional')
        tool=next(row for row in overall_rows if row['dataset']==dataset and row['comparison']=='structured_planner')
        service=control['fill_effect_pp'];cost=control['cost_effect_pct']
        lines += [f"**{dataset}: the complete workflow.** The mean service difference was {service['mean']:+.2f} percentage points (95% interval {service['ci_low']:+.2f} to {service['ci_high']:+.2f}). The mean operating-cost change was {cost['mean']:+.2f}% (interval {cost['ci_low']:+.2f}% to {cost['ci_high']:+.2f}%). These are effects of the whole agent-and-tool workflow relative to conventional ordering.",'']
        if tool['fill_effect_pp']['ci_high']<0 or tool['cost_effect_pct']['ci_low']>0:
            interpretation='The agents fell short of the structured-data planner on at least one of these measures. Their role as a reader of business messages did not translate into superior numerical replenishment performance.'
        elif tool['fill_effect_pp']['ci_low']>0 or tool['cost_effect_pct']['ci_high']<0:
            interpretation='At least one measured outcome favoured the agents over the structured-data planner in this simulation. The paired graphs show the size and uncertainty of that contribution; it should not be generalised to other retailers.'
        else:
            interpretation='The study did not establish an improvement from adding the LLM roles to the structured-data planner. Any benefit over conventional ordering must also be credited to the numerical planning tool.'
        lines += [f"**{dataset}: the LLM contribution.** "+interpretation,'']
    lines += ['A local LLM can therefore be assessed as a bounded participant in replenishment. Its useful role is to interpret supplier messages and coordinate an approved planning process. This is different from showing that it is a better forecaster or quantity optimiser. The structured-data comparison is essential because a strong planning tool may explain much of any advantage over conventional ordering.', '',
        'The results support the measured workflow and identify where the LLM layer adds work, holds or interpretation capability. They do not establish production ROI, employee time savings or a universal model advantage. A retailer considering a pilot should keep the numerical planner and independent checks, test its own supplier messages, and measure the incremental value and review workload of the language-model layer.', '',
        '## Limits of the evidence', '',
        'The study covers two historical windows and selected panels, with controlled supplier prose, memos and tables. All inventory and supplier conditions are simulated. Historical evaluation periods were used in earlier work; fresh execution does not make them previously unseen research holdouts. No staff interviews, live purchasing or operational deployment took place. One small deterministic language model was tested, so conclusions do not automatically transfer to larger models or other retailers.', '',
        '## Figures', '']
    for index,figure in enumerate(REGISTER,1):
        lines += [f"### Figure {index}. {figure['title']}", '',f"![{figure['title']}]({figure['path']})",'',figure['caption'],'']
    lines += ['## Evidence files', '',
        'The protocol and frozen source identify the prospective design. The run folders contain supplier messages, current authenticated contracts, three model requests/responses per agent day, proposals, checks, executed quantities, and series-level outcomes. The analysis folder provides run, daily, simulation-seed and paired-effect tables plus the independent audit receipt. The original experiments remain separate.', '',
        'Dataset source: M5 competition retail sales; Dingdong-Inc/FreshRetailNet-50K at https://huggingface.co/datasets/Dingdong-Inc/FreshRetailNet-50K. Model source: https://huggingface.co/Qwen/Qwen2.5-1.5B-Instruct. The pinned dataset revision, license and file hashes are recorded in the data-preparation receipt.', '']
    text='\n'.join(lines);(OUT/'Fresh_Ollama_Study_Report.md').write_text(text)
    return text


def documents(text):
    from docx import Document
    from docx.shared import Inches,Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Image,PageBreak,KeepTogether,Table,TableStyle
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    import html,re
    document=Document();section=document.sections[0]
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
    document.styles['Normal'].font.name='Arial';document.styles['Normal'].font.size=Pt(11)
    document.styles['Normal'].paragraph_format.line_spacing=1.25
    styles=getSampleStyleSheet();styles['Normal'].fontSize=10.5;styles['Normal'].leading=15
    styles['Heading1'].fontSize=16;styles['Heading1'].leading=21;styles['Heading2'].fontSize=13;styles['Heading2'].leading=18
    story=[];web=['<!doctype html><html lang="en"><meta charset="utf-8"><title>Fresh Ollama study</title><style>body{font:17px/1.6 Arial,sans-serif;max-width:1000px;margin:40px auto;padding:0 20px;color:#263849}img{max-width:100%;height:auto}h1,h2,h3{line-height:1.3}figure{margin:30px 0}small{color:#546575}</style><body>']
    in_figures=False;lines=text.splitlines();position=0
    while position<len(lines):
        line=lines[position];position+=1
        if not line.strip():continue
        if line.startswith('| '):
            rows=[]
            while True:
                values=[value.strip() for value in line.strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',value) for value in values):rows.append(values)
                if position>=len(lines) or not lines[position].startswith('| '):break
                line=lines[position];position+=1
            table=document.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
            for cell,value in zip(table.rows[0].cells,rows[0]):cell.text=value
            for values in rows[1:]:
                for cell,value in zip(table.add_row().cells,values):cell.text=value
            pdfrows=[[Paragraph(html.escape(value),styles['Normal']) for value in values] for values in rows]
            pdftable=Table(pdfrows,colWidths=[73,135,135,142],repeatRows=1,hAlign='LEFT')
            pdftable.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EDF3F7')),
                ('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#BAC8D3')),
                ('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7)]))
            story.extend([pdftable,Spacer(1,9)])
            web.append('<table style="width:100%;border-collapse:collapse">')
            for index,values in enumerate(rows):
                tag='th' if index==0 else 'td'
                web.append('<tr>'+''.join(f'<{tag} style="text-align:left;padding:8px;border-bottom:1px solid #cbd6df">'+html.escape(value)+f'</{tag}>' for value in values)+'</tr>')
            web.append('</table>');continue
        if line.startswith('## Figures'):
            in_figures=True;document.add_page_break();story.append(PageBreak())
        if line.startswith('# '):
            value=line[2:];document.add_heading(value,0);story.append(Paragraph(html.escape(value),styles['Title']));story.append(Spacer(1,14));web.append('<h1>'+html.escape(value)+'</h1>')
        elif line.startswith('## '):
            value=line[3:];document.add_heading(value,1);story.append(Paragraph(html.escape(value),styles['Heading1']));web.append('<h2>'+html.escape(value)+'</h2>')
        elif line.startswith('### '):
            value=line[4:]
            if in_figures and not value.startswith('Figure 1.'):
                document.add_page_break();story.append(PageBreak())
            document.add_heading(value,2);story.append(Paragraph(html.escape(value),styles['Heading2']));web.append('<h3>'+html.escape(value)+'</h3>')
        elif line.startswith('!['):
            match=re.match(r'!\[(.*?)\]\((.*?)\)',line);path=OUT/match[2]
            document.add_picture(str(path),width=Inches(6.1))
            from PIL import Image as PILImage
            width,height=PILImage.open(path).size
            size=485;pdfimage=Image(str(path),width=size,height=size*height/width)
            if pdfimage.drawHeight>560:pdfimage.drawWidth*=560/pdfimage.drawHeight;pdfimage.drawHeight=560
            story.append(pdfimage);story.append(Spacer(1,8));web.append('<img src="'+html.escape(match[2])+'" alt="'+html.escape(match[1])+'">')
        else:
            cleaned=line.replace('**','');p=document.add_paragraph(cleaned)
            story.append(Paragraph(html.escape(cleaned),styles['Normal']));story.append(Spacer(1,7));web.append('<p>'+html.escape(cleaned)+'</p>')
    document.save(OUT/'Fresh_Ollama_Study_Report.docx')
    def footer(canvas,doc):
        canvas.setFont('Helvetica',8);canvas.setFillColorRGB(.4,.45,.5);canvas.drawString(55,30,'Fresh Ollama study | M5 and RetailNet');canvas.drawRightString(A4[0]-55,30,str(doc.page))
    SimpleDocTemplate(str(OUT/'Fresh_Ollama_Study_Report.pdf'),pagesize=A4,leftMargin=55,rightMargin=55,topMargin=55,bottomMargin=50).build(story,onFirstPage=footer,onLaterPages=footer)
    web.append('</body></html>');(OUT/'Fresh_Ollama_Study_Report.html').write_text('\n'.join(web))
    import fitz
    pdf=fitz.open(OUT/'Fresh_Ollama_Study_Report.pdf');issues=[]
    for n,page in enumerate(pdf,1):
        if not page.get_text().strip():issues.append({'page':n,'issue':'empty page'})
        for word in page.get_text('words'):
            if word[0]<0 or word[1]<0 or word[2]>page.rect.width+1 or word[3]>page.rect.height+1:issues.append({'page':n,'issue':'out-of-page text'})
    if issues:raise RuntimeError(json.dumps(issues))
    validation={'status':'VERIFIED','pdf_pages':len(pdf),'figures':len(REGISTER),'layout_issues':issues,
        'docx_sha256':sha(OUT/'Fresh_Ollama_Study_Report.docx'),'pdf_sha256':sha(OUT/'Fresh_Ollama_Study_Report.pdf'),
        'figure_inputs':{p.name:sha(p) for p in (OUT/'analysis').glob('*.csv')}}
    (OUT/'document_receipt.json').write_text(json.dumps(validation,indent=2)+'\n')
    (OUT/'Fresh_Ollama_Study_Report_rendered_text.txt').write_text('\n'.join(p.get_text() for p in pdf))
    print(json.dumps({'documents':'BUILT_AND_VERIFIED','pdf_pages':len(pdf),'figures':len(REGISTER)}))


if __name__=='__main__':
    receipt=json.loads((OUT/'analysis/audit_receipt.json').read_text())
    if receipt['status']!='VERIFIED':raise RuntimeError('Numerical audit must pass before reporting')
    figures();documents(narrative())
